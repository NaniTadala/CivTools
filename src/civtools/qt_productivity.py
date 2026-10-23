"""Local STAAD productivity suite, with reviewable command/file outputs."""
import csv
import io
import json
import math
from dataclasses import dataclass
from pathlib import Path

from PySide6.QtCore import Qt, QSize
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QFormLayout,
    QLineEdit, QPlainTextEdit, QTableWidget, QTableWidgetItem,
    QAbstractItemView, QHeaderView, QSpinBox, QFileDialog, QMessageBox, QGridLayout, QToolButton, QMainWindow,
    QStackedWidget, QFrame, QScrollArea, QButtonGroup, QMenu, QSizePolicy)

from .qt_pages import ToolPage, combo, number as numeric_field
from .qt_widgets import label, button, CheckBox
from .core.staad_productivity import (CATEGORIES, DOFS, ModelSnapshot, classify_members,
    group_commands, release_commands, renumber_plan, renumber_csv, combination_commands,
    read_combination_csv, tapered_commands, id_list)
from .core import staad_productivity_com as live


def number(*args):
    field = numeric_field(*args)
    # Keep native buttons: on some Windows Qt builds NoButtons + QSS collapses
    # the embedded line editor to a single pixel, hiding all numeric values.
    field.setButtonSymbols(field.ButtonSymbols.UpDownArrows)
    return field


from .qt_link_artwork import link_icon
from .core.link_generators import matrix_contributions

FEATURES = (
    ('select', 'Select'), ('renumber', 'Renum'), ('groups', 'Groups'),
    ('releases', 'Releases'), ('parameters', 'Dgn Prm'), ('tables', 'Tables'),
    ('combinations', 'Load Comb'), ('girders', 'Plate Grdrs'), ('rack', 'Pipe Rack'),
)
TITLES = dict(FEATURES)
WORKBENCH_TITLES = dict(FEATURES, select='Select members', renumber='Renumber members',
    groups='Member groups', releases='Member releases', parameters='Design parameters',
    tables='Results tables', combinations='Load combinations', girders='Plate girders', rack='Pipe rack',
    **{'peb' + key: 'PEB frame ' + key for key in 'ABCD'})
NAVIGATION = ('select', 'renumber', 'groups', 'releases', 'parameters',
              'combinations', 'tables', 'rack', 'pebA', 'girders')
NAV_LABELS = dict(WORKBENCH_TITLES, renumber='Renumber', groups='Groups',
                  releases='Releases', tables='Results', pebA='PEB frames')


@dataclass
class SelectionUpdate:
    snapshot: ModelSnapshot
    groups: frozenset
    result: object

    def __str__(self):
        return str(self.result)


class WorkflowWindow(QMainWindow):
    def __init__(self, palette, key):
        super().__init__()
        self.palette, self.key = palette, key
        from .qt_style import staad_palette
        self.setPalette(staad_palette())
        if key == 'select':
            self.page = WorkflowPage(palette.window, key, palette)
            title = TITLES[key]
        elif key in ('parameters', 'rack'):
            from .qt_productivity_migrated import ProductivityParametersPage, ProductivityRackPage
            self.page = (ProductivityParametersPage if key == 'parameters' else ProductivityRackPage)(palette.window, palette)
            title = 'Design Parameters' if key == 'parameters' else 'Pipe Rack'
        else:
            from .qt_productivity_tools import ProductivityToolPage, TITLES as TOOL_TITLES
            self.page = ProductivityToolPage(palette.window, key, palette)
            title = TOOL_TITLES[key]
        self.setWindowTitle(title + ' | CivTools')
        self.setWindowIcon(link_icon(key))
        self.setCentralWidget(self.page)
        self.page.set_workbench_mode()
        self.configure_compact_page()
        self.detached_toolbar = QWidget()
        tools = QHBoxLayout(self.detached_toolbar)
        tools.setContentsMargins(12, 8, 12, 0)
        tools.addWidget(button('Return to workbench', lambda: self.return_to_workbench(), True))
        tools.addStretch()
        pin = self.pin = CheckBox('Keep on top')
        def set_pinned(checked):
            self.setWindowFlag(Qt.WindowType.WindowStaysOnTopHint, checked)
            self.show()
        pin.toggled.connect(set_pinned)
        tools.addWidget(pin)
        self.setMenuWidget(self.detached_toolbar)
        self.detached_toolbar.hide()
        dimensions = {'select': (660, 320), 'renumber': (540, 240), 'groups': (540, 240),
                      'releases': (600, 260), 'parameters': (920, 760), 'rack': (760, 720), 'tables': (740, 560),
                      'combinations': (620, 320), 'girders': (620, 260)}
        self.resize(*dimensions.get(key, (820, 640)))
        self.setMinimumSize(0, 0)
        geometry = palette.window.settings.value('windows/productivity-v2/' + key)
        if geometry:
            self.restoreGeometry(geometry)

    def save_state(self):
        self.page.save_state()
        self.palette.window.settings.setValue('windows/productivity-v2/' + self.key, self.saveGeometry())

    def configure_compact_page(self):
        page = self.page
        if hasattr(page, 'model_caption'):
            page.model_caption.hide()
        actions = []
        while page.action_bar.count():
            item = page.action_bar.takeAt(0)
            if item.widget() and item.widget() not in (page.feedback, page.log_toggle):
                actions.append(item.widget())
        page.layout().removeItem(page.action_bar)
        page.action_bar = QGridLayout()
        page.action_bar.setSpacing(4)
        page.feedback.setWordWrap(True)
        page.feedback.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
        status = QHBoxLayout()
        status.addWidget(page.feedback, 1)
        status.addWidget(page.log_toggle)
        page.layout().addLayout(status)
        if self.key == 'parameters':
            menu = QMenu(page)
            for action in actions[1:]:
                action.hide()
                menu.addAction(action.text(), action.click)
            options = QToolButton()
            options.setText('Commands')
            options.setMenu(menu)
            options.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
            actions = [actions[0], options]
        for index, action in enumerate(actions):
            page.action_bar.addWidget(action, index // 2, index % 2)
        page.layout().addLayout(page.action_bar)
        if self.key == 'select':
            page.process_members_button.hide()
            page.feedback_text('Choose categories to select in STAAD.')
        elif self.key == 'releases':
            page.body.removeWidget(page.define_releases_button)
            page.action_bar.addWidget(page.define_releases_button, 0, 0)
            page.action_bar.addWidget(page.apply_release_button, 1, 0)

    def return_to_workbench(self):
        self.palette.window.navigate('productivity')
        self.palette.open_feature(self.key, detached=False)

    def closeEvent(self, event):
        self.save_state()
        event.ignore()
        self.hide()


class ProductivityPage(QWidget):
    """Shared engineering workbench with detachable, state-preserving tools."""
    def __init__(self, window):
        super().__init__()
        self.window = window
        self.children_windows = {}
        self.shared_snapshot = None
        self.snapshot_source = None
        self.buttons = {}
        self.shortcuts = []
        self.active_feature = None
        saved_variant = str(window.settings.value('productivity/peb_variant', 'A'))
        self.last_peb_feature = 'peb' + (saved_variant if saved_variant in ('A', 'B', 'C', 'D') else 'A')
        self.compact_window = None
        self.expanded = False
        self.collapsed = False
        try:
            favourites = json.loads(str(window.settings.value('productivity/favourites', '["select", "renumber", "groups", "releases"]')))
            self.favourites = [key for key in favourites if key in NAVIGATION][:4] if isinstance(favourites, list) else ['select', 'releases']
        except (ValueError, TypeError):
            self.favourites = ['select', 'releases']
        self.setObjectName('staadWorkbench')
        from .qt_style import STAAD_WORKBENCH_STYLE, staad_palette
        self.setPalette(staad_palette())
        self.setAutoFillBackground(True)
        self.setStyleSheet(STAAD_WORKBENCH_STYLE)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 6, 10, 8)
        layout.setSpacing(4)
        top = self.toolbar = QHBoxLayout()
        self.model_status = label('No model', 'modelStatus')
        self.model_status.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
        top.addStretch()
        self.tool_search = QLineEdit()
        self.tool_search.setPlaceholderText('Find tool…')
        self.tool_search.setToolTip('Find a tool · Ctrl+K')
        self.tool_search.setAccessibleName('Find a STAAD tool')
        self.tool_search.textChanged.connect(self.filter_tools)
        self.peb_variant = combo(['PEB A', 'PEB B', 'PEB C', 'PEB D'])
        self.peb_variant.setAccessibleName('PEB frame type')
        self.peb_variant.currentIndexChanged.connect(lambda index: self.open_feature('peb' + 'ABCD'[index]))
        self.peb_variant.hide()
        refresh = button('Refresh', self.refresh_model, True)
        refresh.setToolTip('Scan the active STAAD model and refresh cached geometry')
        top.addWidget(refresh)
        self.options_menu = QMenu(self)
        self.detach_button = self.options_menu.addAction('Detach tool', self.detach_active)
        self.options_menu.addAction('Add / remove favourite', self.toggle_favourite)
        options = QToolButton()
        options.setText('More')
        options.setAccessibleName('Workbench options')
        options.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        options.setMenu(self.options_menu)
        top.addWidget(options)
        layout.addLayout(top)
        self.task_area = QWidget()
        task_layout = self.task_layout = QVBoxLayout(self.task_area)
        task_layout.setContentsMargins(0, 0, 0, 0)
        task_layout.setSpacing(3)
        task_layout.addWidget(self.model_status)
        switcher = QHBoxLayout()
        self.tool_picker = combo(['PEB frames' if key == 'pebA' else WORKBENCH_TITLES[key]
                                 for key in NAVIGATION])
        self.tool_picker.setEditable(True)
        self.tool_picker.setInsertPolicy(self.tool_picker.InsertPolicy.NoInsert)
        self.tool_picker.setMinimumWidth(0)
        self.tool_picker.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Fixed)
        self.tool_picker.setAccessibleName('Find and switch STAAD tool')
        self.tool_picker.completer().setFilterMode(Qt.MatchFlag.MatchContains)
        self.tool_picker.completer().setCompletionMode(self.tool_picker.completer().CompletionMode.PopupCompletion)
        self.tool_picker.activated.connect(self.pick_tool)
        self.tool_picker.lineEdit().returnPressed.connect(self.open_picker_match)
        switcher.addWidget(self.tool_picker, 1)
        self.expand_button = button('Expand', self.toggle_expanded, True)
        switcher.addWidget(self.expand_button)
        self.fold_button = button('Fold', self.toggle_collapsed, True)
        top.addWidget(self.fold_button)
        task_layout.addLayout(switcher)
        content = QVBoxLayout()
        content.setSpacing(10)
        navigator = self.favourites_frame = QFrame()
        navigator.setObjectName('toolNavigator')
        nav = QHBoxLayout(navigator)
        nav.setContentsMargins(0, 0, 0, 0)
        nav.setSpacing(2)
        self.tool_search.hide()
        self.nav_group = QButtonGroup(self)
        self.nav_group.setExclusive(True)
        for key in NAVIGATION:
            action = button(NAV_LABELS[key], lambda k=key: self.open_section(k), True)
            action.setObjectName('favouriteTool')
            action.setText({'select': 'Select', 'renumber': 'Renumber', 'parameters': 'Parameters',
                            'combinations': 'Loads', 'girders': 'Girders'}.get(key, NAV_LABELS[key]))
            action.setCheckable(True)
            action.setAccessibleName('Open ' + WORKBENCH_TITLES[key])
            action.setToolTip(WORKBENCH_TITLES[key])
            self.nav_group.addButton(action)
            self.buttons[key] = action
            nav.addWidget(action)
            action.setVisible(key in self.favourites)
        self.no_tools = label('No matching tools.', 'muted', True)
        self.no_tools.hide()
        nav.addWidget(self.no_tools)
        content.addWidget(navigator)
        self.peb_family = QFrame()
        self.peb_family.setObjectName('pebFamily')
        variants = QHBoxLayout(self.peb_family)
        variants.setContentsMargins(4, 4, 4, 4)
        variants.setSpacing(4)
        self.peb_variant_group = QButtonGroup(self)
        self.peb_variant_group.setExclusive(True)
        self.peb_variant_buttons = {}
        descriptions = {'A': 'Gable frame', 'B': 'Gable frame with ridge columns',
                        'C': 'Gable frame with internal columns', 'D': 'Single-slope frame'}
        for variant in 'ABCD':
            action = button('Frame ' + variant, lambda v=variant: self.open_feature('peb' + v), True)
            action.setObjectName('pebVariant')
            action.setCheckable(True)
            action.setAccessibleName('PEB frame ' + variant + ': ' + descriptions[variant])
            action.setToolTip(descriptions[variant])
            self.peb_variant_group.addButton(action)
            self.peb_variant_buttons[variant] = action
            variants.addWidget(action, 1)
        self.peb_family.hide()
        content.addWidget(self.peb_family)
        self.tool_stack = QStackedWidget()
        self.tool_stack.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Expanding)
        self.detached_hint = QWidget()
        hint_layout = QVBoxLayout(self.detached_hint)
        hint_layout.addStretch()
        hint_layout.addWidget(label('This tool is open in a separate window.', 'sectionTitle'))
        hint_layout.addWidget(button('Return tool to workbench', self.dock_active, True))
        hint_layout.addStretch()
        self.tool_stack.addWidget(self.detached_hint)
        content.addWidget(self.tool_stack, 1)
        task_layout.addLayout(content, 1)
        layout.addWidget(self.task_area, 1)
        from PySide6.QtGui import QShortcut, QKeySequence
        for key, letter in [('select', 'S'), ('renumber', 'N'), ('groups', 'G'), ('releases', 'R'),
                            ('parameters', 'M'), ('tables', 'T'), ('combinations', 'L'), ('girders', 'E'),
                            ('rack', 'P'), ('pebA', 'A'), ('pebB', 'B'), ('pebC', 'C'), ('pebD', 'D')]:
            shortcut = QShortcut(QKeySequence('Alt+' + letter), self)
            shortcut.activated.connect(lambda feature=key: self.open_feature(feature))
            self.shortcuts.append(shortcut)
        search_shortcut = QShortcut(QKeySequence('Ctrl+K'), self)
        search_shortcut.activated.connect(self.focus_tool_picker)
        self.shortcuts.append(search_shortcut)
        self.tool_search.returnPressed.connect(self.open_search_result)
        self.open_feature('select')

    def resizeEvent(self, event):
        super().resizeEvent(event)
        compact = self.width() < 700
        if not self.collapsed:
            self.expanded = not compact
            self.expand_button.setText('Expand' if compact else 'Compact')
        self.favourites_frame.setVisible(compact)
        if compact:
            self.task_layout.insertWidget(0, self.model_status)
        else:
            self.toolbar.insertWidget(1, self.model_status, 1)

    def focus_tool_picker(self):
        self.tool_picker.setFocus()
        self.tool_picker.lineEdit().selectAll()

    def pick_tool(self, index):
        if 0 <= index < len(NAVIGATION):
            self.open_section(NAVIGATION[index])

    def open_section(self, key):
        return self.open_feature(self.last_peb_feature if key == 'pebA' else key)

    def open_picker_match(self):
        query = self.tool_picker.currentText().strip().casefold()
        if not query:
            return
        index = next((index for index in range(self.tool_picker.count())
                      if query in self.tool_picker.itemText(index).casefold()), None)
        if index is not None:
            self.pick_tool(index)

    def toggle_favourite(self):
        key = 'pebA' if self.active_feature.startswith('peb') else self.active_feature
        if key in self.favourites:
            self.favourites.remove(key)
        else:
            self.favourites = (self.favourites + [key])[-4:]
        for feature, action in self.buttons.items():
            action.setVisible(feature in self.favourites)
        self.window.settings.setValue('productivity/favourites', json.dumps(self.favourites))

    def toggle_expanded(self):
        if self.collapsed:
            self.toggle_collapsed()
        self.expanded = not self.expanded
        self.expand_button.setText('Compact' if self.expanded else 'Expand')
        self.parentWidget().window().resize(1000 if self.expanded else 420, 740 if self.expanded else 480)

    def toggle_collapsed(self):
        host = self.parentWidget().window()
        self.collapsed = not self.collapsed
        if self.collapsed:
            self.unfold_size = host.size()
        self.task_area.setVisible(not self.collapsed)
        self.fold_button.setText('Unfold' if self.collapsed else 'Fold')
        host.setMinimumSize(360, 50 if self.collapsed else 300)
        host.resize(QSize(host.width(), 54) if self.collapsed else self.unfold_size)

    def filter_tools(self, text):
        query = text.strip().casefold()
        for key, action in self.buttons.items():
            action.setVisible(query in (action.text() + ' ' + WORKBENCH_TITLES[key] + ' ' + TITLES.get(key, '')).casefold() if query else key in self.favourites)
        self.no_tools.setVisible(all(action.isHidden() for action in self.buttons.values()))

    def open_search_result(self):
        key = next((key for key, action in self.buttons.items() if not action.isHidden()), None)
        if key:
            self.open_section(key)

    def detach_active(self):
        if self.active_feature:
            self.open_feature(self.active_feature, detached=True)

    def dock_active(self):
        if self.active_feature:
            self.open_feature(self.active_feature, detached=False)

    def show_compact(self):
        self.compact_window = self.parentWidget().window()
        if self.collapsed:
            self.toggle_collapsed()
        self.expanded = False
        self.expand_button.setText('Expand')
        self.compact_window.resize(420, 480)
        self.compact_window.showNormal()
        self.compact_window.raise_()
        self.compact_window.activateWindow()

    def show_workbench(self):
        if not self.expanded:
            self.toggle_expanded()

    @property
    def busy(self):
        return any(child.page.busy for child in self.children_windows.values())

    def open_feature(self, key, detached=None):
        if key not in self.children_windows:
            self.children_windows[key] = WorkflowWindow(self, key)
            if self.shared_snapshot is not None and hasattr(self.children_windows[key].page, 'set_snapshot'):
                self.children_windows[key].page.set_snapshot(self.shared_snapshot, publish=False)
        child = self.children_windows[key]
        self.active_feature = key
        is_peb = key.startswith('peb')
        self.tool_picker.setCurrentText(self.tool_picker.itemText(NAVIGATION.index('pebA' if is_peb else key)))
        self.buttons['pebA' if is_peb else key].setChecked(True)
        self.peb_variant.hide()
        self.peb_family.setVisible(is_peb)
        if is_peb:
            self.last_peb_feature = key
            self.window.settings.setValue('productivity/peb_variant', key[-1])
            self.peb_variant_buttons[key[-1]].setChecked(True)
            self.peb_variant.blockSignals(True)
            self.peb_variant.setCurrentIndex('ABCD'.index(key[-1]))
            self.peb_variant.blockSignals(False)
        if detached is True:
            child.page.set_workbench_mode(embedded=False)
            self.tool_stack.removeWidget(child)
            child.setParent(None)
            child.setWindowFlags(Qt.WindowType.Window)
            child.setWindowFlag(Qt.WindowType.WindowStaysOnTopHint, child.pin.isChecked())
            from .qt_style import STAAD_WORKBENCH_STYLE
            child.setObjectName('staadWorkbench')
            child.setStyleSheet(STAAD_WORKBENCH_STYLE)
            child.detached_toolbar.show()
            child.resize(600, 480)
        elif detached is False or self.tool_stack.indexOf(child) < 0 and child.parent() is None and not child.isVisible():
            child.page.set_workbench_mode()
            child.detached_toolbar.hide()
            child.setParent(self.tool_stack)
            child.setWindowFlags(Qt.WindowType.Widget)
            self.tool_stack.addWidget(child)
        if self.tool_stack.indexOf(child) >= 0:
            self.tool_stack.setCurrentWidget(child)
            child.show()
        else:
            self.tool_stack.setCurrentWidget(self.detached_hint)
            child.showNormal()
            child.raise_()
            child.activateWindow()
        return child.page

    def refresh_model(self):
        self.children_windows['select'].page.scan()

    def share_snapshot(self, snapshot, source, owner=None):
        self.shared_snapshot, self.snapshot_source = snapshot, source
        self.model_status.setText(f'{Path(snapshot.path).name} · {snapshot.force_unit}, {snapshot.length_unit} · cached')
        self.model_status.setToolTip(f'{source}\n{snapshot.path}\n{len(snapshot.members):,} members\n'
                                    'Cached model data. Live operations recheck the active STAAD model.')
        for child in self.children_windows.values():
            if child.page is not owner and getattr(child.page, 'automatic', None) is not owner and hasattr(child.page, 'set_snapshot'):
                child.page.set_snapshot(snapshot, publish=False)

    def invalidate_snapshot(self, reason):
        self.shared_snapshot = None
        self.snapshot_source = None
        self.model_status.setText(reason)
        for child in self.children_windows.values():
            page = child.page
            page.snapshot, page.categories = None, {}
            if hasattr(page, 'members'):
                page.members.setRowCount(0)
            if page.feature == 'select':
                page.reset_selection_buttons()
            page.invalidate_mapping()
            if hasattr(page, 'model_caption'):
                page.model_caption.setText(reason)
            if hasattr(page, 'result_caption'):
                page.result_caption.setText('Scan the model again.')

    def restore_state(self, key):
        pass

    def save_state(self):
        for child in self.children_windows.values():
            child.save_state()
        # Application exit calls this after all workers finish. Closing just the
        # palette also hides its child windows while preserving their inputs.
        self.close_children()

    def close_children(self, *args):
        for child in self.children_windows.values():
            child.hide()
        if self.compact_window is not None:
            self.compact_window.hide()


class WorkflowPage(ToolPage):
    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.feature != 'select' or not hasattr(self, 'members'):
            return
        compact = self.width() < 700
        header = self.members.horizontalHeader()
        for column in (1, 2, 4):
            self.members.setColumnHidden(column, compact)
            header.setSectionResizeMode(column, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.members.setColumnWidth(0, 72)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(5, QHeaderView.ResizeMode.Fixed)
        self.members.setColumnWidth(5, 72)
        self.members.verticalHeader().setMinimumSectionSize(32)
        self.members.resizeRowsToContents()

    def run(self, operation, *, cad=False, cad_system="STAAD"):
        return super().run(operation, cad=cad, cad_system=cad_system)

    def __init__(self, window, feature, palette=None):
        super().__init__(window, WORKBENCH_TITLES[feature], 'STAAD tools | ' + TITLES[feature], 'Modeling', artwork='productivity')
        self.feature = feature
        self.palette = palette
        self.snapshot_source = "Offline snapshot"
        self.active_categories = None
        self.mutating = False
        self.snapshot, self.categories, self.mapping = None, {}, None
        self.output_extension = '.txt'
        self.state_widgets = {}
        self.header_art.hide()
        heading = self.layout().itemAt(0).layout().itemAt(0).layout()
        heading.itemAt(1).widget().hide()
        self.setObjectName('productivityWorkflow')
        self.setStyleSheet('QWidget#productivityWorkflow QAbstractSpinBox QLineEdit { border: none; padding: 0; min-height: 0; background: transparent; }')
        self.model_caption = label('Scan STAAD to begin.', 'muted', True)
        self.result_caption = label('', 'muted', True)
        if feature == 'select':
            self.selected_groups = set()
            self.body.setSpacing(6)
            self.body.setContentsMargins(0, 0, 0, 0)
            self.model_caption.setText('No model processed.')
            row = QHBoxLayout()
            row.addWidget(self.model_caption, 1)
            self.process_members_button = button('Scan active model', self.scan)
            row.addWidget(self.process_members_button)
            self.body.addLayout(row)
            self.build_selection_controls()
            filters = QHBoxLayout()
            self.search = QLineEdit()
            self.search.setPlaceholderText('Find member ID or category')
            self.category_filter = combo(['All categories', *CATEGORIES])
            self.search.textChanged.connect(self.filter_members)
            self.category_filter.currentTextChanged.connect(self.filter_members)
            filters.addWidget(self.search, 1)
            filters.addWidget(self.category_filter)
            details, self.member_details_toggle = self.section('Member details')
            details.addLayout(filters)
            self.members = QTableWidget(0, 6)
            self.members.setHorizontalHeaderLabels(['Member', 'Start node', 'End node', 'Category', 'Mid Y', 'Length'])
            self.members.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
            self.members.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            self.members.setAlternatingRowColors(True)
            self.members.verticalHeader().hide()
            self.members.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
            self.members.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
            self.members.setMinimumHeight(210)
            details.addWidget(self.members, 1)
            self.body.addStretch()
            self.feedback_text('Choose categories to select in STAAD.')
            return
        self.result_nodes, self.result_cases = QLineEdit(), QLineEdit()
        self.members = QTableWidget(0, 6)
        self.members.setHorizontalHeaderLabels(['Member', 'Start', 'End', 'Category', 'Mid Y', 'Length'])
        self.members.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.members.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.members.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.members.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.members.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.search = QLineEdit()
        self.search.setPlaceholderText('Find member ID or category')
        self.category_filter = combo(['All categories', *CATEGORIES])
        self.search.textChanged.connect(self.filter_members)
        self.category_filter.currentTextChanged.connect(self.filter_members)
        self.apply_button = button('Apply renumber plan', self.apply_renumber)
        self.apply_button.setEnabled(False)
        row = QHBoxLayout()
        row.addWidget(button('Process Members' if feature == 'parameters' else 'Scan STAAD', self.scan))
        self.body.addLayout(row)
        self.body.addWidget(self.model_caption)
        snapshot_options, _ = self.section('Snapshot options')
        row = QHBoxLayout()
        row.addWidget(button('Load snapshot', self.load_snapshot, True))
        row.addWidget(button('Save snapshot', self.save_snapshot, True))
        snapshot_options.addLayout(row)
        if feature in ('renumber', 'groups', 'releases'):
            member_layout = self.body
            if feature != 'releases':
                member_layout, _ = self.section('Member details')
            else:
                for column in (1, 2, 4, 5):
                    self.members.setColumnHidden(column, True)
            row = QHBoxLayout()
            row.addWidget(self.search, 1)
            row.addWidget(self.category_filter)
            member_layout.addLayout(row)
            self.members.setMinimumHeight(150)
            member_layout.addWidget(self.members)
            row = QHBoxLayout()
            row.addWidget(button('Select visible', self.select_visible, True))
            if feature == 'releases':
                row.addWidget(button('Select in STAAD', self.select_in_staad))
            member_layout.addLayout(row)
        if feature == 'renumber':
            self.renumber_start = QSpinBox()
            self.renumber_start.setRange(1, 1_000_000_000)
            self.state_widgets['renumber_start'] = self.renumber_start
            self.renumber_start.valueChanged.connect(self.invalidate_mapping)
            self.body.addWidget(label('Starting member number'))
            self.body.addWidget(self.renumber_start)
            self.body.addWidget(button('Preview renumbering', self.preview_renumber))
            self.body.addWidget(self.apply_button)
        elif feature == 'groups':
            self.group_prefix = QLineEdit('CIV')
            self.state_widgets['group_prefix'] = self.group_prefix
            self.body.addWidget(label('Group prefix'))
            self.body.addWidget(self.group_prefix)
            self.body.addWidget(button('Preview groups', self.preview_groups))
            self.body.addWidget(button('Apply groups to STAAD', self.apply_groups))
        elif feature == 'releases':
            self.release_end = combo(['START', 'END', 'BOTH'])
            self.state_widgets['release_end'] = self.release_end
            self.body.addWidget(self.release_end)
            self.release_dofs = {}
            row = QHBoxLayout()
            for dof in DOFS:
                field = CheckBox(dof)
                field.setChecked(dof == 'MZ')
                self.release_dofs[dof] = field
                self.state_widgets[dof] = field
                row.addWidget(field)
            self.body.addLayout(row)
            self.body.addWidget(button('Preview releases', self.preview_releases))
            self.body.addWidget(button('Apply releases to STAAD', self.apply_releases))
        elif feature == 'tables':
            self._results_tab()
            self.state_widgets['result_kind'] = self.result_kind
            self.state_widgets.update(self.pedestal_dimensions)
            self.state_widgets.update(nodes=self.result_nodes, cases=self.result_cases)
        elif feature == 'girders':
            self._girder_tab()
            self.state_widgets.update(self.girder_dimensions)
            self.state_widgets.update(girder_ids=self.girder_ids, unit=self.girder_unit)
        elif feature == 'combinations':
            self.combination_csv = QPlainTextEdit('combination,title,case,factor\n101,ULS DL LL,1,1.5\n101,ULS DL LL,2,1.5')
            self.combination_csv.setToolTip('Excel LOAD_COMB: case IDs on row 2, symbols on row 3, types on row 4; combination numbers in column A from row 5.')
            row = QHBoxLayout()
            row.addWidget(button('Save Excel template', self.save_combination_template, True))
            row.addWidget(button('Import CSV / Excel', self.import_combinations, True))
            self.body.addLayout(row)
            self.body.addWidget(self.combination_csv)
            self.body.addWidget(button('Preview combinations', self.preview_combinations))
            self.state_widgets['csv'] = self.combination_csv
        elif feature == 'parameters':
            note = label('LY/LZ: review actual restraints before use.', 'muted', True)
            note.setToolTip('Geometric length candidates assume each end joint is restrained. LX, UNL and LT depend on the design code and are omitted.')
            self.body.addWidget(note)
            self.design_actions = {}
            for key, title in [('columns', 'Columns LY / LZ candidates'), ('beams_dj', 'Beams DJ1 / DJ2'), ('beams_lengths', 'Beams LY / LZ candidates')]:
                action = button(title, lambda checked=False, mode=key: self.design_parameters(mode))
                self.body.addWidget(action)
                self.design_actions[key] = action
        preview_layout, self.preview_toggle = self.section('Preview')
        if feature in ('select', 'tables'):
            self.preview_toggle.hide()
        self.output_caption = label('', 'muted', True)
        preview_layout.addWidget(self.output_caption)
        self.output = QPlainTextEdit()
        self.output.setReadOnly(True)
        self.output.setMinimumHeight(130)
        self.output.setLineWrapMode(QPlainTextEdit.LineWrapMode.NoWrap)
        preview_layout.addWidget(self.output)
        row = QHBoxLayout()
        row.addWidget(button('Copy', self.copy_output, True))
        row.addWidget(button('Export preview', self.export_output))
        preview_layout.addLayout(row)
        self.body.addStretch()
        self.restore_inputs()

    def section(self, title):
        toggle = QToolButton()
        toggle.setText(title)
        toggle.setCheckable(True)
        toggle.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextBesideIcon)
        toggle.setArrowType(Qt.ArrowType.RightArrow)
        panel = QWidget()
        content = QVBoxLayout(panel)
        content.setContentsMargins(0, 4, 0, 4)
        panel.hide()
        toggle.toggled.connect(panel.setVisible)
        toggle.toggled.connect(lambda checked: toggle.setArrowType(
            Qt.ArrowType.DownArrow if checked else Qt.ArrowType.RightArrow))
        self.body.addWidget(toggle)
        self.body.addWidget(panel)
        return content, toggle

    def _tab(self, title):
        return self.form, self.body

    def design_parameters(self, mode):
        def start():
            model = self.require_model()
            operation = getattr(live, 'process_design_parameters', None)
            if operation is None:
                raise ValueError('This design-parameter operation is awaiting backend validation.')
            self.run(lambda report: operation(model, mode, report), cad=True)
        self.guard(start)

    def restore_inputs(self):
        try:
            data = json.loads(str(self.window.settings.value('forms/productivity/' + self.feature, '{}')))
            for name, field in self.state_widgets.items():
                if name not in data:
                    continue
                value = data[name]
                if isinstance(field, QPlainTextEdit): field.setPlainText(str(value))
                elif isinstance(field, QLineEdit): field.setText(str(value))
                elif isinstance(field, CheckBox): field.setChecked(value is True)
                elif hasattr(field, 'setCurrentText'): field.setCurrentText(str(value))
                elif math.isfinite(float(value)): field.setValue(float(value) if not isinstance(field, QSpinBox) else int(value))
        except (ValueError, TypeError, AttributeError):
            self.feedback_text('Saved inputs could not be restored.', 'warning')
        for field in self.state_widgets.values():
            signal = field.textChanged if isinstance(field, (QLineEdit, QPlainTextEdit)) else field.toggled if isinstance(field, CheckBox) else field.currentTextChanged if hasattr(field, 'currentTextChanged') else field.valueChanged
            signal.connect(self.save_state)

    def save_state(self, *args):
        data = {}
        for name, field in self.state_widgets.items():
            data[name] = field.toPlainText() if isinstance(field, QPlainTextEdit) else field.text() if isinstance(field, QLineEdit) else field.isChecked() if isinstance(field, CheckBox) else field.currentText() if hasattr(field, 'currentText') else field.value()
        self.window.settings.setValue('forms/productivity/' + self.feature, json.dumps(data))

    def save_combination_template(self):
        def save():
            from openpyxl import Workbook
            path, _ = QFileDialog.getSaveFileName(self, 'Save LOAD_COMB template', self.last_folder() + '/load-combinations.xlsx', 'Excel workbook (*.xlsx)')
            if not path: return
            wb = Workbook()
            try:
                ws = wb.active
                ws.title = 'LOAD_COMB'
                for row in [('Combination', 'Primary load cases'), (None, 1, 2), (None, 'DL', 'LL'), (None, 'Dead', 'Live'), (101, 1.5, 1.5), (102, 1., 1.)]: ws.append(row)
                ws.freeze_panes = 'B5'
                wb.save(path)
            finally: wb.close()
            self.remember_folder(path)
            self.feedback_text('Saved template: ' + path, 'success')
        self.guard(save)

    def import_combinations(self):
        def load():
            path, _ = QFileDialog.getOpenFileName(self, 'Import load combinations', self.last_folder(), 'Load combinations (*.csv *.xlsx)')
            if not path: return
            if Path(path).suffix.lower() == '.xlsx':
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True, data_only=False)
                try:
                    if 'LOAD_COMB' not in wb.sheetnames: raise ValueError('Workbook requires a LOAD_COMB worksheet.')
                    ws = wb['LOAD_COMB']
                    if ws.max_row > 204 or ws.max_column > 51: raise ValueError('Use at most 50 cases and 200 combinations.')
                    contributions = matrix_contributions(ws.iter_rows(values_only=True))
                    stream = io.StringIO()
                    writer = csv.writer(stream)
                    writer.writerow(['combination', 'title', 'case', 'factor'])
                    writer.writerows(contributions)
                    text = stream.getvalue()
                finally: wb.close()
            else:
                if Path(path).stat().st_size > 2_000_000: raise ValueError('CSV exceeds 2 MB.')
                text = Path(path).read_text(encoding='utf-8-sig')
            combination_commands(read_combination_csv(text))
            self.combination_csv.setPlainText(text)
            self.remember_folder(path)
        self.guard(load)

    def _results_tab(self):
        _, layout = self._tab('Results')
        form = QFormLayout()
        self.result_kind = combo(['Reactions', 'Displacements', 'Pedestal'])
        self.result_nodes = QLineEdit()
        self.result_nodes.setPlaceholderText('Blank: supports for reactions, all nodes for displacements')
        self.result_cases = QLineEdit()
        self.result_cases.setPlaceholderText('Blank: all scanned cases; or 1, 2, 10 TO 15')
        form.addRow('Results', self.result_kind)
        form.addRow('Node IDs', self.result_nodes)
        form.addRow('Load case IDs', self.result_cases)
        layout.addLayout(form)
        self.result_caption = label('No model scanned.', 'muted', True)
        layout.addWidget(self.result_caption)
        self.pedestal_panel = QWidget()
        form = QFormLayout(self.pedestal_panel)
        self.pedestal_dimensions = {}
        for name, title, default in [('height', 'Height', 1.5), ('width_x', 'Width X', .6),
                                     ('width_z', 'Width Z', .6), ('density', 'Weight density (force / lengthÂ³)', 24)]:
            field = number(default, .000001, 1e9)
            field.setDecimals(6)
            self.pedestal_dimensions[name] = field
            form.addRow(title, field)
        layout.addWidget(self.pedestal_panel)
        self.pedestal_panel.setToolTip('Use STAAD base units. Global Y is vertical. Transfers signed reactions downward, adds pedestal weight to Fy and retains torsion My. Negate reactions for foundation loads.')
        layout.addWidget(button('Export result workbook', self.export_results))
        layout.addStretch()
        self.result_kind.currentTextChanged.connect(lambda text: self.pedestal_panel.setVisible(text == 'Pedestal'))
        self.pedestal_panel.hide()

    def _girder_tab(self):
        _, layout = self._tab('Plate girders')
        form = QFormLayout()
        self.girder_ids = QLineEdit('1 TO 5')
        self.girder_unit = combo(['METER', 'INCHES'])
        form.addRow('Member IDs', self.girder_ids)
        form.addRow('Dimension units', self.girder_unit)
        self.girder_dimensions = {}
        for key, title, default in [('d1', 'Depth at start', .8), ('tw', 'Web thickness', .01),
            ('d2', 'Depth at end', .6), ('bt', 'Top flange width', .3), ('tt', 'Top flange thickness', .02),
            ('bb', 'Bottom flange width', .3), ('tb', 'Bottom flange thickness', .02)]:
            field = number(default, .000001, 100000)
            field.setDecimals(6)
            self.girder_dimensions[key] = field
            form.addRow(title, field)
        layout.addLayout(form)
        layout.addWidget(button('Preview plate-girder property', self.preview_girder))
        layout.addStretch()

    def guard(self, operation):
        try:
            operation()
        except Exception as exc:
            self.invalid(str(exc))

    def require_model(self):
        if self.snapshot is None:
            raise ValueError('Scan STAAD or load a model snapshot first.')
        return self.snapshot

    def scan(self):
        self.snapshot_source = 'STAAD scan (cached)'
        self.run(live.scan, cad=True)

    def on_success(self, result):
        super().on_success(result)
        if isinstance(result, SelectionUpdate):
            if self.snapshot is not result.snapshot:
                self.set_snapshot(result.snapshot)
            self.selected_groups = set(result.groups)
            self.sync_selection_buttons()
        elif isinstance(result, ModelSnapshot):
            self.set_snapshot(result)
        elif isinstance(result, str):
            self.show_output(result, 'Generated design parameters for review.')
        elif self.mapping is not None and getattr(result, 'output_path', ''):
            # A successful renumber changes every cached model reference.
            self.invalidate_mapping()

    def set_snapshot(self, snapshot, publish=True):
        changed = self.snapshot is None or self.snapshot.fingerprint != snapshot.fingerprint
        self.snapshot = snapshot
        self.categories = classify_members(snapshot, overrides=snapshot.classifications)
        self.invalidate_mapping()
        self.model_caption.setText(f'{Path(snapshot.path).name} · {len(snapshot.members):,} members')
        self.model_caption.setToolTip(str(snapshot) + '\n' + snapshot.path)
        self.result_caption.setText(f'Units: {snapshot.length_unit}, {snapshot.force_unit} · '
                                    f'{len(snapshot.load_cases):,} load cases')
        self.result_caption.setToolTip(f'Moments: {snapshot.force_unit} {snapshot.length_unit}; rotations: rad.\n'
                                       'Load cases: ' + ', '.join(map(str, snapshot.load_cases)))
        if self.feature == 'select':
            self.model_caption.setText(Path(snapshot.path).name)
            self.selection_ids = self.selection_group_ids(snapshot, self.categories, self.selection_categories)
            if changed:
                self.reset_selection_buttons()
            else:
                self.sync_selection_buttons()
            self.populate_members()
        else:
            self.populate_members()
        if publish and self.palette is not None:
            self.palette.share_snapshot(snapshot, self.snapshot_source, self)

    def populate_members(self):
        model = self.require_model()
        self.members.setUpdatesEnabled(False)
        try:
            self.members.setRowCount(len(model.members))
            for row, (mid, (a, b)) in enumerate(sorted(model.members.items())):
                values = (mid, a, b, self.categories[mid],
                          f'{(model.nodes[a][1] + model.nodes[b][1]) / 2:.5g}',
                          f'{math.dist(model.nodes[a], model.nodes[b]):.5g}')
                for col, value in enumerate(values):
                    self.members.setItem(row, col, QTableWidgetItem(str(value)))
        finally:
            self.members.setUpdatesEnabled(True)
        self.filter_members()

    def build_selection_controls(self):
        self.scope_summary = label('No selection · Scan the active model or choose a category to begin.', 'scopeSummary', True)
        self.body.addWidget(self.scope_summary)
        columns = QHBoxLayout()
        self.body.addLayout(columns)
        grids = []
        for title in ('MEMBER CATEGORIES', 'DIRECTION & SUPPORTS'):
            box = QVBoxLayout()
            box.setAlignment(Qt.AlignmentFlag.AlignTop)
            box.addWidget(label(title, 'navHeading'))
            grid = QVBoxLayout()
            grid.setAlignment(Qt.AlignmentFlag.AlignTop)
            grid.setSpacing(0)
            box.addLayout(grid)
            columns.addLayout(box, 1)
            grids.append(grid)
        self.category_buttons = {}
        self.selection_categories = dict([
                   ('Columns', ('Columns',)), ('Pri.Beams', ('Primary beams X', 'Primary beams Z')),
                   ('Sec.Beams', ('Secondary beams X', 'Secondary beams Z', 'Other horizontal')),
                   ('Vrt.Bracing', ('Vertical bracing',)), ('Plan.Bracing', ('Plan diagonals',)),
                   ('Sup.Nodes', None),
                   ('Pri.X', ('Primary beams X',)), ('Pri.Z', ('Primary beams Z',)),
                   ('Sec/Ter.X', ('Secondary beams X',)), ('Sec/Ter.Z', ('Secondary beams Z',)),
                   ('Rafters', ('Rafters',)), ('Raf.Bracing', ('Rafter bracing',))])
        self.selection_labels = dict(zip(self.selection_categories, (
            'Columns', 'Primary beams', 'Secondary beams', 'Vertical bracing', 'Plan bracing',
            'Support nodes', 'Primary beams · X', 'Primary beams · Z', 'Secondary beams · X',
            'Secondary beams · Z', 'Rafters', 'Rafter bracing')))
        for index, title in enumerate(self.selection_categories):
            action = CheckBox(self.selection_labels[title])
            action.clicked.connect(lambda checked=False, key=title: self.toggle_selection_group(key, checked))
            action.setAccessibleName(self.selection_labels[title])
            action.setToolTip('Click to select; click again to deselect. Combine groups by checking more buttons.')
            action.setStyleSheet('QPushButton#secondary:checked { background: #F2F7FA; border: 2px solid #7B94A3; }')
            self.category_buttons[title] = action
            grids[0 if index < 5 or index >= 10 else 1].addWidget(action)

    def reset_selection_buttons(self):
        self.selected_groups.clear()
        self.sync_selection_buttons()

    def sync_selection_buttons(self):
        for title, action in self.category_buttons.items():
            action.setChecked(title in self.selected_groups)
            action.setEnabled(self.snapshot is None or bool(self.selection_ids.get(title)))
            count = f' · {len(self.selection_ids.get(title, ())):,}' if self.snapshot else ''
            action.setText(self.selection_labels[title] + count)
        ids = set().union(*(self.selection_ids[key] for key in self.selected_groups
                            if self.selection_categories[key] is not None)) if self.selected_groups else set()
        nodes = len(self.selection_ids['Sup.Nodes']) if 'Sup.Nodes' in self.selected_groups else 0
        self.scope_summary.setText(f'{len(ids):,} members · {nodes:,} support nodes selected through CivTools')

    @staticmethod
    def selection_group_ids(model, classifications, definitions):
        by_role = {}
        for mid, role in classifications.items():
            by_role.setdefault(role, set()).add(mid)
        return {key: frozenset(model.supports) if roles is None else frozenset(
            mid for role in roles for mid in by_role.get(role, ()))
            for key, roles in definitions.items()}

    def toggle_selection_group(self, title, checked):
        if self.busy:
            self.sync_selection_buttons()
            return
        groups = set(self.selected_groups)
        groups.add(title) if checked else groups.discard(title)
        expected = self.snapshot
        if expected is None:
            self.snapshot_source = 'STAAD scan (cached)'
        definitions = dict(self.selection_categories)
        cached = dict(self.selection_ids) if expected is not None else None

        def select(report):
            model = expected if expected is not None else live.scan(report)
            indexed = cached if cached is not None else self.selection_group_ids(
                model, classify_members(model, overrides=model.classifications), definitions)
            ids = set().union(*(indexed[key] for key in groups if definitions[key] is not None))
            nodes = indexed['Sup.Nodes'] if 'Sup.Nodes' in groups else frozenset()
            if checked and not indexed[title]:
                raise ValueError(f'No {title} found in the active model.')
            result = live.set_selection(model, sorted(ids), sorted(nodes), report)
            return SelectionUpdate(model, frozenset(groups), result)

        self.run(select, cad=True)

    def filter_members(self, *args):
        term, category = self.search.text().strip().casefold(), self.category_filter.currentText()
        for row in range(self.members.rowCount()):
            role = self.members.item(row, 3).text()
            mid = int(self.members.item(row, 0).text())
            match = term in (str(mid) + ' ' + role).casefold()
            match = match and (role in self.active_categories if self.active_categories else category == 'All categories' or role == category)
            self.members.setRowHidden(row, not match)

    def selected_ids(self):
        self.require_model()
        ids = [int(self.members.item(index.row(), 0).text())
               for index in self.members.selectionModel().selectedRows()
               if not self.members.isRowHidden(index.row())]
        if not ids:
            raise ValueError('Select member rows first; use Select visible for a whole category.')
        return sorted(ids)

    def select_visible(self):
        from PySide6.QtCore import QItemSelection, QItemSelectionModel
        selection = QItemSelection()
        for row in range(self.members.rowCount()):
            if not self.members.isRowHidden(row):
                selection.select(self.members.model().index(row, 0), self.members.model().index(row, 5))
        self.members.selectionModel().select(selection, QItemSelectionModel.SelectionFlag.ClearAndSelect | QItemSelectionModel.SelectionFlag.Rows)

    def select_in_staad(self):
        def start():
            model, ids = self.require_model(), self.selected_ids()
            self.run(lambda report: live.select_members(model, ids, report), cad=True)
        self.guard(start)

    def select_supports(self):
        def start():
            model = self.require_model()
            self.run(lambda report: live.select_members(model, model.supports, report, nodes=True), cad=True)
        self.guard(start)

    def show_output(self, text, caption, extension='.txt'):
        if extension != '.csv':
            self.invalidate_mapping()
        self.output.setPlainText(text)
        self.preview_toggle.show()
        self.preview_toggle.setChecked(True)
        self.output_caption.setText(caption)
        self.output_extension = extension
        self.feedback_text(caption, 'success')

    def preview_groups(self):
        self.guard(lambda: self.show_output(group_commands(self.require_model(), self.categories, self.group_prefix.text().strip()),
                                            'Group commands for ' + self.snapshot.path))

    def preview_releases(self):
        self.guard(lambda: self.show_output(release_commands(self.selected_ids(), self.release_end.currentText(),
                      [k for k, field in self.release_dofs.items() if field.isChecked()]),
                      'Release commands in LOCAL member axes for ' + self.snapshot.path))

    def confirm_model_change(self, title, model, preview, effect):
        box = QMessageBox(self)
        box.setWindowTitle(title)
        box.setIcon(QMessageBox.Icon.Question)
        box.setText(title + '\n\nTarget: ' + model.path)
        box.setInformativeText(effect + '\n\nCivTools will save the current model and create a backup before applying these commands. Review the refreshed preview below; Details contains the complete commands.')
        box.setDetailedText(preview)
        box.setStandardButtons(QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        box.setDefaultButton(QMessageBox.StandardButton.No)
        return box.exec() == QMessageBox.StandardButton.Yes

    def run_model_change(self, operation):
        if self.busy:
            return
        self.mutating = True
        def execute(report):
            result = operation(report)
            report('model_changed', True)
            return result
        self.run(execute, cad=True)

    def apply_groups(self):
        def start():
            model = self.require_model()
            categories, prefix = dict(self.categories), self.group_prefix.text().strip()
            preview = group_commands(model, categories, prefix)
            self.show_output(preview, 'Fresh group preview for ' + model.path)
            if self.confirm_model_change('Apply groups to STAAD', model, preview,
                                         'Creates the previewed member groups. Existing group-name collisions are refused.'):
                self.run_model_change(lambda report: live.apply_groups(model, categories, prefix, report))
        self.guard(start)

    def apply_releases(self):
        def start():
            model, ids = self.require_model(), tuple(self.selected_ids())
            end = self.release_end.currentText()
            dofs = tuple(key for key, field in self.release_dofs.items() if field.isChecked())
            preview = release_commands(ids, end, dofs)
            self.show_output(preview, 'Fresh local-axis release preview for ' + model.path)
            if self.confirm_model_change('Apply releases to STAAD', model, preview,
                                         f'Replaces releases at selected ends ({end}) of {len(ids):,} selected members with the previewed local-axis releases.'):
                self.run_model_change(lambda report: live.apply_releases(model, ids, end, dofs, report))
        self.guard(start)

    def invalidate_mapping(self, *args):
        self.mapping = None
        if hasattr(self, 'apply_button'):
            self.apply_button.setEnabled(False)

    def preview_renumber(self):
        def generate():
            self.mapping = renumber_plan(self.require_model(), self.renumber_start.value(), self.categories)
            self.show_output(renumber_csv(self.mapping), 'Renumber plan: category, then Y/Z/X midpoint order. ' + self.snapshot.path, '.csv')
            self.apply_button.setEnabled(True)
        self.guard(generate)

    def apply_renumber(self):
        if self.mapping is None:
            self.invalid('Generate a fresh renumber plan first.')
            return
        model, mapping = self.snapshot, dict(self.mapping)
        answer = QMessageBox.question(self, 'Apply member renumbering',
            f'Save and renumber {len(mapping):,} members in:\n{model.path}\n\n'
            'CivTools will save the current model, create a backup and a numbering journal, '
            'then apply the previewed plan. Existing analysis results and saved selections may become invalid. '
            'Close the STAAD editor before continuing.', QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No)
        if answer == QMessageBox.StandardButton.Yes:
            def operation(report):
                result = live.apply_renumber(model, mapping, report)
                report('renumbered', True)
                return result
            self.run(operation, cad=True)

    def on_event(self, kind, value):
        if self.feature == 'select' and kind in ('target', 'stage', 'log'):
            return
        if kind == 'model_changed':
            if self.palette is not None:
                self.palette.invalidate_snapshot('Model changed. Refresh STAAD before continuing.')
            else:
                self.snapshot, self.categories = None, {}
                self.members.setRowCount(0)
                self.invalidate_mapping()
            self.mutating = False
        elif kind == 'renumbered':
            if self.palette is not None:
                self.palette.invalidate_snapshot('Model renumbered. Refresh STAAD before continuing.')
            self.snapshot = None
            self.categories = {}
            self.members.setRowCount(0)
            self.invalidate_mapping()
            self.model_caption.setText('Model renumbered. Scan again before using member tools.')
            self.result_caption.setText('Scan the renumbered model again.')
        else:
            super().on_event(kind, value)

    def on_failure(self, message):
        if self.feature == 'select':
            if 'Selection could not be restored' in message:
                self.reset_selection_buttons()
            else:
                self.sync_selection_buttons()
        if self.mutating or self.mapping is not None:
            if self.palette is not None:
                self.palette.invalidate_snapshot('Operation failed. Inspect STAAD and refresh before continuing.')
            # An interrupted renumber must never leave an applicable stale plan.
            self.snapshot = None
            self.categories = {}
            self.members.setRowCount(0)
            self.invalidate_mapping()
            self.model_caption.setText('Operation failed. Inspect STAAD and scan again before continuing.')
        self.mutating = False
        super().on_failure(message)
        if self.feature == 'select':
            self.log.hide()

    def on_finished(self):
        super().on_finished()
        if self.feature == 'select':
            self.runtime.hide()
            self.target.hide()
            self.log.hide()

    def export_results(self):
        def start():
            model = self.require_model()
            kind = self.result_kind.currentText()
            allowed = model.nodes if kind == 'Displacements' else model.supports
            nodes = id_list(self.result_nodes.text(), allowed) if self.result_nodes.text().strip() else sorted(allowed)
            cases = id_list(self.result_cases.text(), model.load_cases) if self.result_cases.text().strip() else list(model.load_cases)
            if not nodes or not cases:
                raise ValueError('The scanned model has no matching nodes or load cases.')
            pedestal = tuple(field.value() for field in self.pedestal_dimensions.values()) if kind == 'Pedestal' else None
            path, _ = QFileDialog.getSaveFileName(self, 'Export STAAD results', self.last_folder() + '/staad-results.xlsx', 'Excel workbook (*.xlsx)')
            if path:
                self.remember_folder(path)
                self.run(lambda report: live.export_results(model, kind, nodes, cases, path, report, pedestal), cad=True)
        self.guard(start)

    def preview_combinations(self):
        self.guard(lambda: self.show_output(combination_commands(read_combination_csv(self.combination_csv.toPlainText()),
                          self.snapshot.load_cases if self.snapshot else None,
                          self.snapshot.primary_cases if self.snapshot else None), 'Linear load-combination commands from your factors.'))

    def preview_girder(self):
        self.guard(lambda: self.show_output(tapered_commands(id_list(self.girder_ids.text(),
                         self.snapshot.members if self.snapshot else None),
                         [field.value() for field in self.girder_dimensions.values()], self.girder_unit.currentText()),
                         'Tapered-I property commands; check incidence and restore input units after insertion.'))

    def last_folder(self):
        return str(self.window.settings.value('folders/productivity', str(Path.home() / 'Documents')))

    def remember_folder(self, path):
        self.window.settings.setValue('folders/productivity', str(Path(path).parent))

    def copy_output(self):
        from PySide6.QtWidgets import QApplication
        QApplication.clipboard().setText(self.output.toPlainText())
        self.feedback_text('Copied preview.', 'success')

    def export_output(self):
        def save():
            if not self.output.toPlainText().strip():
                raise ValueError('Generate a preview first.')
            ext = self.output_extension
            path, _ = QFileDialog.getSaveFileName(self, 'Export preview', self.last_folder() + '/civtools-output' + ext,
                                                 f'Generated output (*{ext})')
            if path:
                if Path(path).suffix.lower() != ext:
                    raise ValueError(f'Use the {ext} extension for this output.')
                if self.snapshot and Path(path).resolve() == Path(self.snapshot.path).resolve():
                    raise ValueError('Choose a new output file; keep the source STAAD model.')
                Path(path).write_text(self.output.toPlainText(), encoding='utf-8')
                self.remember_folder(path)
                self.feedback_text('Exported ' + path, 'success')
        self.guard(save)

    def save_snapshot(self):
        def save():
            model = self.require_model()
            path, _ = QFileDialog.getSaveFileName(self, 'Save geometry snapshot', self.last_folder() + '/staad-snapshot.json', 'JSON (*.json)')
            if path:
                if Path(path).suffix.lower() != '.json':
                    raise ValueError('Save the snapshot as .json.')
                Path(path).write_text(model.to_json(), encoding='utf-8')
                self.remember_folder(path)
        self.guard(save)

    def load_snapshot(self):
        def load():
            path, _ = QFileDialog.getOpenFileName(self, 'Load geometry snapshot', self.last_folder(), 'JSON (*.json)')
            if path:
                if Path(path).stat().st_size > 16_000_000:
                    raise ValueError('Snapshot exceeds 16 MB.')
                self.snapshot_source = 'Offline snapshot'
                self.set_snapshot(ModelSnapshot.from_json(Path(path).read_text(encoding='utf-8')))
                self.remember_folder(path)
                self.feedback_text('Loaded offline geometry snapshot. Live actions recheck the active model.', 'success')
        self.guard(load)
