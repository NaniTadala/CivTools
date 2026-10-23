"""Independent, local file generators for the STAAD productivity windows.

The input fields follow the supplied dEZignLink guide and inspected client.
Generated files are reviewable exports; no function connects to a running model.
"""
from dataclasses import dataclass, replace
import csv
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from .geometry import RackPlan, MAX_BEAMS
from .link_generators import pipe_rack_plan
from .peb import PebConfig, build_peb_plan, geometry_commands
from .staad_productivity import combination_commands, finite, fmt, positive_id, continued


UNITS = ('METER', 'INCHES', 'FEET')


def _unit(unit):
    value = str(unit).strip().upper()
    aliases = {'METRIC (M)': 'METER', 'MKS': 'METER', 'ENGLISH (IN)': 'INCHES',
               'FPS': 'INCHES', 'M': 'METER', 'IN': 'INCHES', 'FT': 'FEET'}
    value = aliases.get(value, value)
    if value not in UNITS:
        raise ValueError('Choose METER, INCHES or FEET units.')
    return value


def _geometry_text(plan, unit):
    unit = _unit(unit)
    text = geometry_commands(plan, 'METER' if unit == 'METER' else 'INCHES')
    return text.replace('UNIT INCHES KIP', 'UNIT FEET KIP') if unit == 'FEET' else text


def _dimension(value, caption):
    value = finite(value, caption, True)
    if not 1e-6 <= value <= 1e9:
        raise ValueError(f'{caption} must be between 0.000001 and 1,000,000,000 in the selected units.')
    return value


@dataclass(frozen=True)
class PipeRackConfig:
    width: float = 6
    frames: int = 4
    frame_spacing: float = 6
    tier_spacings: str = '4,3'
    longitudinal_spacings: str = '4,3'
    unit: str = 'METER'
    support: str = 'Pinned'


def build_pipe_rack_plan(config):
    """Split column lines at both transverse and longitudinal tier heights."""
    _unit(config.unit)
    for caption, text in (('Tier spacing', config.tier_spacings),
                          ('Longitudinal spacing', config.longitudinal_spacings)):
        values = str(text).split(',')
        if len(values) > 500:
            raise ValueError('Use at most 500 spacings.')
        if sum(_dimension(value.strip(), caption) for value in values) > 1e9:
            raise ValueError('Total rack height must be at most 1,000,000,000 in the selected units.')
    _dimension(config.width, 'Rack width')
    _dimension(config.frame_spacing, 'Frame spacing')
    return pipe_rack_plan(config.width, config.frames, config.frame_spacing,
                          config.tier_spacings, config.longitudinal_spacings,
                          config.support)


def pipe_rack_commands(config):
    return _geometry_text(build_pipe_rack_plan(config), config.unit)


@dataclass(frozen=True)
class ProductivityPebConfig(PebConfig):
    add_ties: bool = True
    add_purlins: bool = True
    lateral_bracing: bool = False
    roof_slope: float | None = None


def build_productivity_peb_plan(config):
    """Build A gable, B ridge columns, C internal columns or D single slope."""
    unit = _unit(config.unit)
    numeric = {name: _dimension(getattr(config, name), name.replace('_', ' '))
               for name in ('width', 'eave_height', 'ridge_height', 'frame_spacing',
                            'tie_spacing', 'purlin_spacing', 'inner_left', 'inner_right')}
    if config.roof_slope is not None:
        slope = finite(config.roof_slope, 'Roof slope denominator', True)
        run = numeric['width'] if config.variant == 'D' else numeric['width'] / 2
        numeric['ridge_height'] = numeric['eave_height'] + run / slope
    normalized = replace(config, **numeric, unit='METER' if unit == 'METER' else 'INCHES')
    original = build_peb_plan(normalized)
    beams = [(a, b, kind) for a, b, kind in original.beams
             if (config.add_ties or kind != 'tie') and (config.add_purlins or kind != 'purlin')]
    edges = {tuple(sorted((a, b))) for a, b, _ in beams}
    lookup = {tuple(round(v, 10) for v in xyz): i for i, xyz in enumerate(original.nodes)}

    def node(x, y, z):
        return lookup[tuple(round(v, 10) for v in (x, y, z))]

    def beam(a, b, kind):
        edge = tuple(sorted((a, b)))
        if edge not in edges:
            if len(beams) >= MAX_BEAMS:
                raise ValueError('PEB exceeds the 30,000-member limit.')
            edges.add(edge)
            beams.append((a, b, kind))

    def height(x):
        rise = numeric['ridge_height'] - numeric['eave_height']
        if config.variant == 'D':
            return numeric['eave_height'] + rise * x / numeric['width']
        return numeric['ridge_height'] - rise * abs(x - numeric['width'] / 2) / (numeric['width'] / 2)

    xs = sorted({original.nodes[index][0] for index in original.foundations})
    if config.add_ties and not config.add_purlins:
        # The base planner deduplicates eave/ridge ties against roof purlins.
        # Restore these required ties when purlins have been switched off.
        for frame in range(1, int(config.frames)):
            for x in xs:
                beam(node(x, height(x), (frame - 1) * numeric['frame_spacing']),
                     node(x, height(x), frame * numeric['frame_spacing']), 'tie')
    if config.lateral_bracing:
        bays = sorted({0, int(config.frames) - 2})
        roof_xs = [0., numeric['width']] if config.variant == 'D' else [0., numeric['width'] / 2, numeric['width']]
        for bay in bays:
            za, zb = bay * numeric['frame_spacing'], (bay + 1) * numeric['frame_spacing']
            for x in (0., numeric['width']):
                beam(node(x, 0., za), node(x, height(x), zb), 'vertical_brace')
                beam(node(x, 0., zb), node(x, height(x), za), 'vertical_brace')
            for xa, xb in zip(roof_xs, roof_xs[1:]):
                beam(node(xa, height(xa), za), node(xb, height(xb), zb), 'roof_brace')
                beam(node(xa, height(xa), zb), node(xb, height(xb), za), 'roof_brace')
    return RackPlan(original.nodes, beams, original.foundations, original.support_type)


def peb_commands(config):
    return _geometry_text(build_productivity_peb_plan(config), config.unit)


@dataclass(frozen=True)
class PlateGirder:
    name: str
    depth_mm: float
    flange_width_mm: float
    flange_thickness_mm: float
    web_thickness_mm: float

    def properties(self):
        """STAAD WIDE FLANGE: AX D TW WF TF IZ IY IX AY AZ, in metre units.

        Rectangle area/inertias and thin-wall Saint-Venant torsion omit welds
        and fillets. Shear areas follow the inspected WIDE FLANGE catalogue.
        """
        if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_-]{0,49}', self.name):
            raise ValueError('Use a section name of at most 50 letters, digits, _ or -; start with a letter.')
        d, b, tf, tw = [finite(v, 'Girder dimension', True) / 1000 for v in
                        (self.depth_mm, self.flange_width_mm, self.flange_thickness_mm, self.web_thickness_mm)]
        if 2 * tf >= d or tw >= b:
            raise ValueError('Flanges must fit within the depth and the web within the flange width.')
        if max(d, b, tf, tw) > 100:
            raise ValueError('Girder dimensions must be at most 100,000 mm.')
        web = d - 2 * tf
        area = 2 * b * tf + web * tw
        iz = (b * d ** 3 - (b - tw) * web ** 3) / 12
        iy = (2 * tf * b ** 3 + web * tw ** 3) / 12
        ix = (2 * b * tf ** 3 + web * tw ** 3) / 3
        return area, d, tw, b, tf, iz, iy, ix, d * tw, 4 * b * tf / 3


def predefined_plate_girders():
    """Dimension families observed in the supplied local catalogue capture.

    Names and section properties are independently generated from dimensions.
    """
    result = []

    def add(d, widths, thicknesses):
        for b in widths:
            for tf, tw in thicknesses:
                result.append(PlateGirder(f'PG{d}X{b}X{tf}X{tw}', d, b, tf, tw))

    standard = [(25, 16), (20, 12), (16, 10)]
    for d in (500, 550, 600):
        add(d, (200, 150), standard)
    for d in (650, 700, 750):
        add(d, (250, 200), standard)
    for d in (800, 850):
        add(d, (400,), [(40, 20)])
        add(d, (250, 200), standard)
    for d in (900, 950):
        add(d, (300,), [(32, 20), *standard])
        add(d, (250,), standard)
    for d in (1000, 1100, 1200, 1300):
        add(d, (400, 350), [(32, 20), (25, 16), (20, 12), (16, 12)])
        add(d, (300,), [(25, 16), (20, 12), (16, 12)])
    for d in (1400, 1500):
        add(d, (500,), [(40, 25), (40, 20), (32, 20), (25, 16), (20, 12), (16, 12)])
        add(d, (400, 350), [(32, 20), (25, 16), (20, 12), (16, 12)])
        add(d, (300,), [(25, 16), (20, 12), (16, 12)])
    return result


def plate_girder_catalogue(sections=None):
    sections = predefined_plate_girders() if sections is None else list(sections)
    if not 1 <= len(sections) <= 1000:
        raise ValueError('Export between 1 and 1,000 plate girder sections.')
    names = [section.name.upper() for section in sections]
    if len(set(names)) != len(names):
        raise ValueError('Plate girder section names must be distinct.')
    lines = ['* CivTools plate girders: dimensions are in section names (mm).',
             '* Paste after MEMBER INCIDENCES; use an unused user table number.',
             '* Properties calculated from rectangular plates, omitting welds/fillets.',
             'START USER TABLE', 'TABLE 1', 'UNIT METER KN', 'WIDE FLANGE']
    for section in sections:
        lines.append(section.name)
        lines.extend(continued('', [fmt(value) for value in section.properties()]))
    return '\n'.join([*lines, 'END', ''])


@dataclass(frozen=True)
class LoadCase:
    number: int
    symbol: str
    load_type: str


@dataclass(frozen=True)
class LoadCombinationData:
    cases: tuple
    contributions: tuple

    @property
    def case_script(self):
        return '\n'.join(f'LOAD {c.number} LOADTYPE {c.load_type} TITLE {c.symbol}' for c in self.cases) + '\n'

    @property
    def combination_script(self):
        return _combination_commands(self.contributions)


# OpenSTAAD 0.2 documentation, GetLoadType (p.161), lists all 24 types.
LOAD_TYPES = ('None', 'Dead', 'Live', 'Roof Live', 'Wind', 'Seismic-H', 'Seismic-V',
              'Snow', 'Fluids', 'Soil', 'Temperature', 'Rain', 'Ponding', 'Dust',
              'Traffic', 'Imperfection', 'Accidental', 'Flood', 'Ice', 'Wind Ice',
              'Crane Hook', 'Mass', 'Gravity', 'Push')


def _load_type(value):
    text = str(value or 'None').strip()
    types = {re.sub(r'\s+', '', v).casefold(): v for v in LOAD_TYPES}
    types['temp'] = 'Temperature'
    types['seismic'] = 'Seismic-H'
    key = re.sub(r'\s+', '', text).casefold()
    if key not in types:
        raise ValueError(f'Unsupported load type {text!r}. Use a load type from the template.')
    return types[key]


def _symbol(value, number):
    text = str(value or number).strip()
    if len(text) > 40 or any(ord(ch) < 32 or ch in ';\n\r' for ch in text):
        raise ValueError('Load symbols must be at most 40 characters without control characters or semicolons.')
    return text


def load_combination_data(source):
    """Read LOAD_COMB matrix (IDs in row 2, symbols row 3, types row 4).

    CSV may contain that same matrix, or the existing long-form
    combination,title,case,factor format with optional symbol/load_type columns.
    """
    source = Path(source)
    if source.stat().st_size > 10_000_000:
        raise ValueError('Load combination input must be smaller than 10 MB.')
    if source.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            if 'LOAD_COMB' not in workbook.sheetnames:
                raise ValueError('Workbook requires a LOAD_COMB worksheet.')
            sheet = workbook['LOAD_COMB']
            if sheet.max_row > 204 or sheet.max_column > 51:
                raise ValueError('Use at most 50 primary cases and 200 combinations.')
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    elif source.suffix.lower() == '.csv':
        with source.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.reader(stream)
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > 10_001 or len(row) > 51:
                    raise ValueError('Load combination CSV is too large.')
        if rows and [v.strip().lower() for v in rows[0][:4]] == ['combination', 'title', 'case', 'factor']:
            return _long_form_data(rows)
    else:
        raise ValueError('Choose an XLSX, XLSM or CSV load combination file.')
    if len(rows) < 5:
        raise ValueError('LOAD_COMB requires IDs, symbols, load types and combination factors.')
    if len(rows) > 204 or max(map(len, rows), default=0) > 51:
        raise ValueError('Use at most 50 primary cases and 200 combinations.')
    cases = []
    columns = []
    for col, value in enumerate(rows[1][1:], 1):
        if value is None or not str(value).strip():
            continue
        number = positive_id(value)
        symbol = _symbol(rows[2][col] if col < len(rows[2]) else None, number)
        load_type = _load_type(rows[3][col] if col < len(rows[3]) else None)
        cases.append(LoadCase(number, symbol, load_type))
        columns.append(col)
    if len({c.number for c in cases}) != len(cases):
        raise ValueError('Primary load case IDs must be distinct.')
    contributions, numbers = [], set()
    for row in rows[4:]:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        number = positive_id(row[0])
        if number in numbers:
            raise ValueError(f'Combination ID {number} appears in more than one matrix row.')
        numbers.add(number)
        factors = []
        for col, case in zip(columns, cases):
            raw = row[col] if col < len(row) else None
            if raw is None or not str(raw).strip():
                continue
            factor = finite(raw, 'Load factor')
            if factor:
                factors.append((case, factor))
        if not factors:
            raise ValueError(f'Combination {number} has no nonzero factors.')
        title = ' + '.join(f'{fmt(factor)}{case.symbol}' for case, factor in factors)
        contributions.extend((number, title, case.number, factor) for case, factor in factors)
    return _validated_data(cases, contributions)


def _long_form_data(rows):
    headers = [v.strip().lower() for v in rows[0]]
    cases, contributions = {}, []
    for row in rows[1:]:
        if not any(str(v).strip() for v in row):
            continue
        values = dict(zip(headers, row))
        try:
            number = positive_id(values['case'])
            case = LoadCase(number, _symbol(values.get('symbol'), number), _load_type(values.get('load_type')))
            if number in cases and cases[number] != case:
                raise ValueError(f'Case {number} has inconsistent symbols or load types.')
            cases[number] = case
            contributions.append((positive_id(values['combination']), values['title'], number,
                                  finite(values['factor'], 'Load factor')))
        except KeyError as error:
            raise ValueError('Each CSV row requires combination, title, case and factor.') from error
    return _validated_data(list(cases.values()), contributions)


def _validated_data(cases, contributions):
    _combination_commands(contributions)
    if not cases or len(cases) > 50 or len({r[0] for r in contributions}) > 200:
        raise ValueError('Use 1–50 primary cases and 1–200 load combinations.')
    overlap = {c.number for c in cases} & {r[0] for r in contributions}
    if overlap:
        raise ValueError(f'Primary case IDs and combination IDs must differ: {sorted(overlap)}')
    return LoadCombinationData(tuple(cases), tuple(contributions))


def _combination_commands(contributions):
    # Preserve full expressions in report tables. Long expressions use a short
    # command title so each STAAD input line stays within the editor width.
    rows = [(number, title if re.fullmatch(r'[A-Za-z0-9 _().+/-]{1,48}', str(title))
             else f'Combination {number}', case, factor)
            for number, title, case, factor in contributions]
    return combination_commands(rows)


def load_combination_template(destination):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    destination = Path(destination)
    if destination.suffix.lower() != '.xlsx':
        raise ValueError('Save the load combination template as an .xlsx file.')
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = 'LOAD_COMB'
        for row in [('LOAD COMBINATIONS FOR STAAD.Pro',), ('LOAD COMB NO', 1, 2),
                    (None, 'DL', 'LL'), (None, 'Dead', 'Live'), (101, 1.5, 1.5), (102, 1, 1)]:
            sheet.append(row)
        sheet.freeze_panes = 'B5'
        sheet.column_dimensions['A'].width = 23
        for col in ('B', 'C'):
            sheet.column_dimensions[col].width = 16
        for row in sheet.iter_rows(min_row=1, max_row=4, max_col=3):
            for cell in row:
                cell.font = Font(name='Calibri', bold=True)
                cell.fill = PatternFill('solid', fgColor='DCE6F1')
                cell.alignment = Alignment(horizontal='center')
        validation = DataValidation(type='list', formula1='"' + ','.join(LOAD_TYPES) + '"')
        validation.errorTitle, validation.error = 'Invalid load type', 'Choose a load type from the list.'
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add('B4:AY4')
        workbook.save(destination)
    finally:
        workbook.close()
    return destination


_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
ET.register_namespace('w', _W)


def _element(parent, name, text=None, **attributes):
    item = ET.SubElement(parent, '{' + _W + '}' + name,
                         {'{' + _W + '}' + key: str(value) for key, value in attributes.items()})
    if text is not None:
        item.text = str(text)
    return item


def _paragraph(parent, text, heading=False, monospace=False):
    paragraph = _element(parent, 'p')
    if heading:
        _element(_element(paragraph, 'pPr'), 'pStyle', val='Heading1')
    run = _element(paragraph, 'r')
    if monospace:
        _element(_element(run, 'rPr'), 'rFonts', ascii='Consolas', hAnsi='Consolas')
    node = _element(run, 't', text)
    node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


def _table(parent, rows):
    table = _element(parent, 'tbl')
    props = _element(table, 'tblPr')
    _element(props, 'tblW', w='0', type='auto')
    borders = _element(props, 'tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        _element(borders, side, val='single', sz=4, color='B4B4B4')
    for index, row in enumerate(rows):
        tr = _element(table, 'tr')
        if index == 0:
            _element(_element(tr, 'trPr'), 'tblHeader')
        for value in row:
            cell = _element(tr, 'tc')
            if index == 0:
                _element(_element(cell, 'tcPr'), 'shd', fill='DCE6F1')
            _paragraph(cell, value)


def load_combination_report(source_xlsx_or_csv, destination_docx):
    """Write a genuine Word document containing both scripts and both tables."""
    destination = Path(destination_docx)
    if destination.suffix.lower() != '.docx':
        raise ValueError('Save the load combination report as a .docx file.')
    if Path(source_xlsx_or_csv).resolve() == destination.resolve():
        raise ValueError('Choose an output file different from the input.')
    data = load_combination_data(source_xlsx_or_csv)
    document = ET.Element('{' + _W + '}document')
    body = _element(document, 'body')
    _paragraph(body, 'Load combinations for STAAD.Pro', heading=True)
    _paragraph(body, 'Load case script', heading=True)
    _paragraph(body, 'Add the actual load items below each case header before analysis.')
    for line in data.case_script.splitlines():
        _paragraph(body, line, monospace=True)
    _paragraph(body, 'Load combination script', heading=True)
    for line in data.combination_script.splitlines():
        _paragraph(body, line, monospace=True)
    _paragraph(body, 'Load cases', heading=True)
    _table(body, [('Case', 'Symbol', 'Load type'), *[(c.number, c.symbol, c.load_type) for c in data.cases]])
    _paragraph(body, 'Load combinations', heading=True)
    combos = {}
    for number, title, case, factor in data.contributions:
        combos.setdefault((number, title), {})[case] = factor
    # Keep wide reports readable by splitting the matrix into 8-case blocks.
    for first in range(0, len(data.cases), 8):
        cases = data.cases[first:first + 8]
        _table(body, [('Combination', 'Expression', *[f'{c.number} ({c.symbol})' for c in cases]),
                      *[(number, title, *[fmt(factors.get(c.number, 0.)) for c in cases])
                        for (number, title), factors in combos.items()]])
    section = _element(body, 'sectPr')
    _element(section, 'pgSz', w=16838, h=11906, orient='landscape')
    _element(section, 'pgMar', top=720, right=720, bottom=720, left=720)
    contents = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    content_types = ('<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                     '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
                     '</Types>')
    relationships = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                     '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                     '</Relationships>')
    document_relationships = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                              '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                              '</Relationships>')
    styles = (f'<w:styles xmlns:w="{_W}"><w:docDefaults><w:rPrDefault><w:rPr>'
              '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="20"/>'
              '</w:rPr></w:rPrDefault></w:docDefaults>'
              '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="Heading 1"/>'
              '<w:rPr><w:b/><w:sz w:val="28"/></w:rPr></w:style></w:styles>')
    with ZipFile(destination, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', contents + content_types)
        archive.writestr('_rels/.rels', contents + relationships)
        archive.writestr('word/document.xml', ET.tostring(document, encoding='utf-8', xml_declaration=True))
        archive.writestr('word/_rels/document.xml.rels', contents + document_relationships)
        archive.writestr('word/styles.xml', contents + styles)
    return destination
