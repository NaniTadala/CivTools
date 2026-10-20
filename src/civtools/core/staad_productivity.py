"""Local STAAD productivity plans. Pure functions; no COM or network access."""
from dataclasses import dataclass, field
from collections import defaultdict
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import re

from .geometry import MAX_BEAMS, MAX_NODES

CATEGORIES = ('Columns', 'Primary beams X', 'Primary beams Z', 'Secondary beams X',
              'Secondary beams Z', 'Other horizontal', 'Plan diagonals',
              'Sloping members', 'Rafters', 'Vertical bracing', 'Rafter bracing')
DOFS = ('FX', 'FY', 'FZ', 'MX', 'MY', 'MZ')
GROUP_CODES = ('COL', 'PB_X', 'PB_Z', 'SB_X', 'SB_Z', 'HORIZ', 'PLAN_DIAG',
               'SLOPING', 'RAFTER', 'VBRACE', 'RBRACE')


def positive_id(value):
    if isinstance(value, bool) or not re.fullmatch(r'[1-9][0-9]*', str(value)):
        raise ValueError(f'Use a positive integer ID: {value!r}')
    result = int(value)
    if result > 2_000_000_000:
        raise ValueError('ID exceeds the supported STAAD range.')
    return result


def finite(value, caption='Value', positive=False):
    value = float(value)
    if not math.isfinite(value) or (positive and value <= 0):
        raise ValueError(f'{caption} must be finite' + (' and positive.' if positive else '.'))
    return value


def fmt(value):
    return f'{value:.10g}'


def id_list(text, allowed=None):
    """Strict IDs / ascending TO ranges. Never accept truncated or ambiguous lists."""
    tokens = re.split(r'[\s,]+', text.strip().upper()) if text.strip() else []
    result, i = set(), 0
    while i < len(tokens):
        start = positive_id(tokens[i])
        i += 1
        end = start
        if i < len(tokens) and tokens[i] == 'TO':
            if i + 1 >= len(tokens):
                raise ValueError('TO requires an ending ID.')
            end = positive_id(tokens[i + 1])
            i += 2
        if end < start or end - start > MAX_BEAMS:
            raise ValueError('Use an ascending range of at most 30,001 IDs.')
        result.update(range(start, end + 1))
        if len(result) > MAX_BEAMS:
            raise ValueError('Too many IDs.')
    if not result:
        raise ValueError('Choose at least one ID.')
    if allowed is not None and (unknown := result - set(allowed)):
        raise ValueError(f'Unknown IDs: {sorted(unknown)[:12]}')
    return sorted(result)


def continued(prefix, values, suffix='', width=78):
    """STAAD continuation lines including reserved space for the final hyphen."""
    lines, line = [], prefix
    for value in map(str, values):
        if len(line) + len(value) + 1 > width - 2:
            lines.append(line.rstrip() + ' -')
            line = ''
        line += (' ' if line else '') + value
    for value in suffix.split():
        if len(line) + len(value) + 1 > width - 2:
            lines.append(line.rstrip() + ' -')
            line = ''
        line += (' ' if line else '') + value
    lines.append(line)
    return lines


@dataclass
class ModelSnapshot:
    path: str
    base_unit: str
    nodes: dict
    members: dict
    supports: list = field(default_factory=list)
    load_cases: list = field(default_factory=list)
    source_digest: str = ''
    classifications: dict = field(default_factory=dict)
    primary_cases: list | None = None

    def __post_init__(self):
        self.nodes = {positive_id(k): tuple(finite(v) for v in xyz) for k, xyz in self.nodes.items()}
        self.members = {positive_id(k): tuple(positive_id(n) for n in ends) for k, ends in self.members.items()}
        self.supports = sorted({positive_id(n) for n in self.supports})
        self.load_cases = sorted({positive_id(n) for n in self.load_cases})
        if self.primary_cases is not None:
            self.primary_cases = sorted({positive_id(n) for n in self.primary_cases})
            if set(self.primary_cases) - set(self.load_cases):
                raise ValueError('Primary cases must belong to the scanned case list.')
        if not self.nodes or not self.members or len(self.nodes) > MAX_NODES or len(self.members) > MAX_BEAMS:
            raise ValueError('Scan a nonempty analytical model within 10,000 nodes / 30,000 members.')
        if self.base_unit not in ('Metric', 'English'):
            raise ValueError('STAAD base unit must be Metric or English.')
        if any(len(xyz) != 3 for xyz in self.nodes.values()):
            raise ValueError('Every node requires X, Y and Z coordinates.')
        for member, ends in self.members.items():
            if len(ends) != 2 or any(n not in self.nodes for n in ends):
                raise ValueError(f'Member {member} has invalid node references.')
            if math.dist(self.nodes[ends[0]], self.nodes[ends[1]]) <= 1e-9:
                raise ValueError(f'Member {member} has zero length.')
        if set(self.supports) - self.nodes.keys():
            raise ValueError('Support node is missing from geometry.')
        self.classifications = {positive_id(k): v for k, v in self.classifications.items()}
        if any(k not in self.members or v not in CATEGORIES for k, v in self.classifications.items()):
            raise ValueError('Snapshot classification has an unknown member or category.')

    @property
    def length_unit(self):
        return 'm' if self.base_unit == 'Metric' else 'in'

    @property
    def force_unit(self):
        return 'kN' if self.base_unit == 'Metric' else 'kip'

    @property
    def fingerprint(self):
        data = (str(Path(self.path).resolve()).casefold(), self.base_unit,
                sorted(self.nodes.items()), sorted(self.members.items()), self.supports, self.source_digest)
        return hashlib.sha256(json.dumps(data).encode()).hexdigest()

    def to_json(self):
        return json.dumps({'version': 1, 'path': self.path, 'base_unit': self.base_unit,
                           'nodes': self.nodes, 'members': self.members,
                           'supports': self.supports, 'load_cases': self.load_cases,
                           'source_digest': self.source_digest, 'classifications': self.classifications,
                           'primary_cases': self.primary_cases}, indent=2)

    @classmethod
    def from_json(cls, text):
        data = json.loads(text)
        if not isinstance(data, dict) or data.pop('version', None) != 1:
            raise ValueError('Unsupported model snapshot version.')
        if not isinstance(data.get('nodes'), dict) or not isinstance(data.get('members'), dict):
            raise ValueError('Snapshot requires node and member dictionaries.')
        return cls(**data)

    def __str__(self):
        return f'{len(self.members):,} members / {len(self.nodes):,} nodes scanned; base units {self.length_unit}, {self.force_unit}.'


def classify_members(model, tolerance=0.001, overrides=None):
    """Orientation rules and straight-chain connectivity, never inferred design restraint.

    A primary beam is an axis-aligned horizontal chain ending at two column nodes.
    Other axis-aligned beams are secondary. Sloping members remain unassigned roles.
    """
    tolerance = finite(tolerance, 'Orientation tolerance', True)
    if tolerance >= .1:
        raise ValueError('Orientation tolerance must be below 0.1.')
    categories, axes, column_nodes = {}, {}, set()
    adjacency = defaultdict(list)
    for mid, (a, b) in model.members.items():
        xyz1, xyz2 = model.nodes[a], model.nodes[b]
        dx, dy, dz = [abs(v - u) for u, v in zip(xyz1, xyz2)]
        length = math.dist(xyz1, xyz2)
        if math.hypot(dx, dz) / length <= tolerance:
            categories[mid] = 'Columns'
            column_nodes.update((a, b))
        elif dy / length <= tolerance:
            if dz / length <= tolerance:
                axes[mid] = 'X'
            elif dx / length <= tolerance:
                axes[mid] = 'Z'
            else:
                categories[mid] = 'Plan diagonals'
        else:
            categories[mid] = 'Sloping members'
        adjacency[a].append(mid)
        adjacency[b].append(mid)

    memo = {}
    def reaches_column(mid, node, axis):
        seen = {mid}
        visited = []
        while node not in column_nodes:
            anchor = (mid, node)
            if anchor in memo:
                reached = memo[anchor]
                break
            visited.append(anchor)
            onward = [m for m in adjacency[node] if m not in seen and axes.get(m) == axis]
            if len(onward) != 1:
                reached = False
                break
            mid = onward[0]
            seen.add(mid)
            a, b = model.members[mid]
            node = b if a == node else a
        else:
            reached = True
        for anchor in visited:
            memo[anchor] = reached
        return reached

    for mid, axis in axes.items():
        primary = all(reaches_column(mid, n, axis) for n in model.members[mid])
        categories[mid] = ('Primary' if primary else 'Secondary') + ' beams ' + axis
    for mid, category in (overrides or {}).items():
        mid = positive_id(mid)
        if mid not in model.members or category not in CATEGORIES:
            raise ValueError('Classification override has an unknown member or category.')
        categories[mid] = category
    return categories


def group_commands(model, categories, prefix='CIV', level_tolerance=.001):
    if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,11}', prefix):
        raise ValueError('Group prefix: 1–12 letters/digits/underscores, starting with a letter.')
    if set(categories) != set(model.members) or any(v not in CATEGORIES for v in categories.values()):
        raise ValueError('Classify the complete model before generating groups.')
    tolerance = finite(level_tolerance, 'Level tolerance', True)
    groups = defaultdict(list)
    for mid, category in sorted(categories.items()):
        groups[category].append(mid)
    lines = ['* Insert after geometry; merge with existing groups as needed.', 'START GROUP DEFINITION', 'MEMBER']
    for category, ids in groups.items():
        name = '_' + prefix + '_' + dict(zip(CATEGORIES, GROUP_CODES))[category]
        lines.append('* ' + category)
        lines.extend(continued(name, ids))
    levels = []
    for mid, (a, b) in sorted(model.members.items(), key=lambda item: model.nodes[item[1][0]][1]):
        ya, yb = model.nodes[a][1], model.nodes[b][1]
        if abs(ya - yb) > tolerance:
            continue
        y = (ya + yb) / 2
        if not levels or abs(levels[-1][0] - y) > tolerance:
            levels.append((y, []))
        levels[-1][1].append(mid)
    if model.supports:
        lines.append('JOINT')
        lines.extend(continued(f'_{prefix}_SUPPORTS', model.supports))
    if levels:
        lines.append('FLOOR')
        for i, (y, ids) in enumerate(levels, 1):
            lines.append(f'* Level {fmt(y)} {model.length_unit}')
            lines.extend(continued(f'_{prefix}_LEVEL_{i}', ids))
    return '\n'.join(lines + ['END GROUP DEFINITION', ''])


def release_commands(ids, end, dofs):
    ids = sorted({positive_id(i) for i in ids})
    if set(dofs) - set(DOFS):
        raise ValueError('Unknown release component.')
    dofs = [d for d in DOFS if d in dofs]
    if not ids or end not in ('START', 'END', 'BOTH') or not dofs:
        raise ValueError('Choose members, START/END/BOTH and at least one release component.')
    if set(dofs) == set(DOFS):
        raise ValueError('Releasing all six components disconnects a member end.')
    lines = ['* Member releases use LOCAL member axes.', '* Merge with existing releases; insert before load definitions.', 'MEMBER RELEASE']
    for location in ('START', 'END') if end == 'BOTH' else (end,):
        lines.extend(continued('', ids, location + ' ' + ' '.join(dofs)))
    return '\n'.join(lines) + '\n'


def renumber_plan(model, start=1, categories=None):
    start = positive_id(start)
    ranks = {name: i for i, name in enumerate(CATEGORIES)}
    def key(mid):
        a, b = (model.nodes[n] for n in model.members[mid])
        center = tuple((u + v) / 2 for u, v in zip(a, b))
        return (ranks.get((categories or {}).get(mid), 99), center[1], center[2], center[0], mid)
    ordered = sorted(model.members, key=key)
    positive_id(start + len(ordered) - 1)
    return {old: start + i for i, old in enumerate(ordered)}


def renumber_steps(mapping):
    """Use a disjoint temporary range so swaps and dense numbering cannot collide."""
    mapping = {positive_id(a): positive_id(b) for a, b in mapping.items()}
    if not mapping or len(set(mapping.values())) != len(mapping):
        raise ValueError('Renumbering requires unique target IDs.')
    changed = [(a, b) for a, b in mapping.items() if a != b]
    temp = max((*mapping, *mapping.values())) + 1
    if changed:
        positive_id(temp + len(changed) - 1)
    return ([(old, temp + i) for i, (old, _) in enumerate(changed)] +
            [(temp + i, new) for i, (_, new) in enumerate(changed)])


def renumber_csv(mapping):
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(['old_member', 'new_member'])
    writer.writerows(mapping.items())
    return stream.getvalue()


def combination_commands(rows, existing_cases=None, primary_cases=None):
    """Rows: combination ID, title, primary load ID, factor. Factors are user supplied."""
    combos = {}
    if len(rows) > 10000:
        raise ValueError('Limit combinations to 10,000 load contributions.')
    for row in rows:
        if len(row) != 4:
            raise ValueError('Each row requires combination, title, load case, factor.')
        cid, title, case, factor = row
        cid, case = positive_id(cid), positive_id(case)
        factor = finite(factor, 'Load factor')
        title = str(title).strip()
        if not re.fullmatch(r'[A-Za-z0-9 _().+/-]{1,48}', title):
            raise ValueError('Use a short plain title (letters, numbers, spaces, . + / - _ parentheses).')
        if factor == 0 or abs(factor) > 100:
            raise ValueError('Use nonzero load factors between -100 and 100.')
        if existing_cases is not None and (cid in existing_cases or case not in existing_cases):
            raise ValueError(f'Combination {cid} conflicts with a load case, or load {case} is unknown.')
        if primary_cases is not None and case not in primary_cases:
            raise ValueError(f'Load {case} is not a primary load case. Nested combinations are not supported.')
        old_title, factors = combos.setdefault(cid, (title, {}))
        if old_title != title or case in factors:
            raise ValueError('A combination must have one title and no repeated load-case rows.')
        factors[case] = factor
    if not combos:
        raise ValueError('Enter at least one combination.')
    if any(case in combos for _, factors in combos.values() for case in factors):
        raise ValueError('Nested combinations are not supported; reference primary load cases.')
    lines = ['* User-specified linear combinations; factors are not code-selected.', '* Insert after primary loads and before analysis commands.']
    for cid, (title, factors) in combos.items():
        lines.append(f'LOAD COMB {cid} {title}')
        lines.extend(continued('', [f'{case} {fmt(factor)}' for case, factor in factors.items()]))
    return '\n'.join(lines) + '\n'


def read_combination_csv(text):
    rows = list(csv.reader(io.StringIO(text)))
    rows = [r for r in rows if any(v.strip() for v in r)]
    if rows and rows[0][0].strip().lower() == 'combination':
        rows.pop(0)
    return rows


def tapered_commands(ids, dimensions, unit='METER'):
    if unit not in ('METER', 'INCHES') or len(dimensions) != 7:
        raise ValueError('Choose METER or INCHES and seven tapered-I dimensions.')
    d1, tw, d2, bt, tt, bb, tb = [finite(v, 'Section dimension', True) for v in dimensions]
    if min(d1, d2) <= tt + tb or tw >= min(bt, bb):
        raise ValueError('Depths must exceed flange thickness sum; web must fit within both flanges.')
    if d1 <= d2:
        raise ValueError('For tapered-I properties, start depth must exceed end depth. Reverse incidence if needed.')
    ids = sorted({positive_id(i) for i in ids})
    if not ids:
        raise ValueError('Choose members for the plate-girder property.')
    lines = ['* Insert in the member-property section; dimensions use the UNIT below.',
             '* Restore the previous UNIT before subsequent dimensional commands.',
             f'UNIT {unit} ' + ('KN' if unit == 'METER' else 'KIP'), 'MEMBER PROPERTY']
    lines.extend(continued('', ids, 'TAPERED ' + ' '.join(map(fmt, (d1, tw, d2, bt, tt, bb, tb)))))
    return '\n'.join(lines) + '\n'


def pedestal_forces(reaction, height, width_x, width_z, density):
    """GLOBAL Y-up: transfer a signed reaction vector down to pedestal bottom.

    Returns reactions, not the equal-and-opposite loads on a footing.
    M_bottom = M_top + (0,h,0) x F. Torsion My is retained.
    """
    fx, fy, fz, mx, my, mz = [finite(v) for v in reaction]
    h, x, z, gamma = [finite(v, 'Pedestal input', True) for v in (height, width_x, width_z, density)]
    weight = h * x * z * gamma
    return (fx, fy + weight, fz, mx + h * fz, my, mz - h * fx)


def write_result_workbook(path, model, kind, records, pedestal=None):
    """Create a complete workbook in a staging file; preserve an existing output on failure."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import os
    import tempfile
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Export results as .xlsx.')
    if kind not in ('Reactions', 'Displacements', 'Pedestal') or not records:
        raise ValueError('Choose a result type and fetch at least one result.')
    wb = Workbook()
    try:
        info = wb.active
        info.title = 'Read me'
        details = [('Model', model.path), ('Snapshot', model.fingerprint), ('Result', kind),
                   ('Axes', 'Global; Y is vertical'), ('Length', model.length_unit),
                   ('Force', model.force_unit), ('Moment', f'{model.force_unit} {model.length_unit}'),
                   ('Rotation', 'radians'), ('Scope', 'Selected nodes and load cases only; no design checks')]
        if kind == 'Pedestal':
            if pedestal is None or len(pedestal) != 4:
                raise ValueError('Pedestal height, width X/Z and weight density are required.')
            details += list(zip(('Height', 'Width X', 'Width Z', 'Weight density'), pedestal))
            details += [('Convention', 'Signed support reactions transferred down in global Y; My retained.'),
                        ('Foundation loads', 'Negate the reaction vector to obtain actions on the foundation.')]
        for row in details:
            info.append(row)
        sheet = wb.create_sheet('Results')
        length, force = model.length_unit, model.force_unit
        components = (['DX', 'DY', 'DZ', 'RX', 'RY', 'RZ'] if kind == 'Displacements' else list(DOFS))
        units = ([length] * 3 + ['rad'] * 3 if kind == 'Displacements' else [force] * 3 + [force + ' ' + length] * 3)
        sheet.append(['Node', 'Load case'] + [f'{c} ({u})' for c, u in zip(components, units)])
        envelope = defaultdict(list)
        for node, case, values in records:
            node, case = positive_id(node), positive_id(case)
            values = tuple(finite(v) for v in values)
            if len(values) != 6:
                raise ValueError('A result record requires six components.')
            if kind == 'Pedestal':
                values = pedestal_forces(values, *pedestal)
            sheet.append([node, case, *values])
            envelope[node].append((case, values))
        summary = wb.create_sheet('Envelopes')
        summary.append(['Node', 'Component', 'Unit', 'Minimum', 'Case at minimum', 'Maximum', 'Case at maximum'])
        for node, rows in sorted(envelope.items()):
            for i, (component, unit) in enumerate(zip(components, units)):
                lo, hi = min(rows, key=lambda r: r[1][i]), max(rows, key=lambda r: r[1][i])
                summary.append([node, component, unit, lo[1][i], lo[0], hi[1][i], hi[0]])
        for ws in wb:
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill('solid', fgColor='E8F1FB')
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(60, max(18, max(len(str(c.value or '')) for c in column) + 2))
            # Treat external model names as text, including formula-looking paths.
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
        descriptor, staging = tempfile.mkstemp(suffix='.xlsx', prefix='.civtools-', dir=path.parent)
        os.close(descriptor)
        try:
            wb.save(staging)
            os.replace(staging, path)
        finally:
            if os.path.exists(staging):
                os.unlink(staging)
    finally:
        wb.close()
