"""Local design candidates and analysis workbooks for the productivity palette.

The vendor guide pp. 14–17 describes these workflows. Restraints below are
geometry candidates, not verification of flange restraint, effective length
factors or code compliance. Only the two Indian steel-code dialects are exposed;
the retained IS800 response establishes LY/LZ/UNL and DJ1/DJ2, not an exhaustive
vendor code catalogue. Installed Bentley help D8.A.4 (GUID-C64E85CD-332E-44C1-
B783-A5B8D0050102) specifies LX for 2007; D8.B.4 (GUID-19E1572B-70F6-4F9D-
BF10-398359C19D80) specifies UNL for 1984. OpenSTAAD's code-number mapping is undocumented in the
supplied wrapper, so assignment uses a backed-up, validated STD insertion.
"""
from __future__ import annotations

from collections import defaultdict
import hashlib
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
from uuid import uuid4

from . import staad_productivity_com as live
from .results import OperationResult
from .staad_productivity import classify_members, finite, fmt, positive_id

DESIGN_CODES = ('IS 800:2007', 'IS 800:1984')
_CODE_COMMANDS = {'IS 800:2007': 'IS800 LSD', 'IS 800:1984': 'INDIAN'}
_MODES = {'columns': 'Column restraint lengths', 'beams_lengths': 'Beam restraint lengths',
          'beams_dj': 'Beam deflection joints'}
_MAX_PAIRS = 100_000


def _code(code):
    if code not in _CODE_COMMANDS:
        raise ValueError('Choose a supported Indian steel design code.')
    return _CODE_COMMANDS[code]


def _physical_lengths(model, service, mode):
    """Trace connected collinear segments between candidate directional restraints.

    A supported column system and its primary beams define column restraints.
    Beam vertical restraints are columns; transverse horizontal members provide
    lateral candidates. A release, branch or physical end terminates tracing.
    Only beta 0/90/180/270 is handled; other orientations require manual review.
    """
    categories = classify_members(model, overrides=model.classifications)
    members = {m.member_id: m for m in service.members}
    adjacency = defaultdict(list)
    for mid, ends in model.members.items():
        for node in ends:
            adjacency[node].append(mid)
    target = [m for m in members.values() if m.steel and
              (m.column_y if mode == 'columns' else (m.beam_x or m.beam_z))]
    primary = {mid for mid, category in categories.items() if category.startswith('Primary beams')}
    result = []
    released = {}
    support_nodes = set(model.supports)
    span_cache = {}

    def release_at(mid, node):
        key = (mid, int(model.members[mid][1] == node))
        if key not in released:
            released[key] = service.member_released_at_end(*key)
        return released[key]

    def span(member, axis):
        if (member.member_id, axis) in span_cache:
            return span_cache[member.member_id, axis]
        ends = model.members[member.member_id]
        total = member.length
        span_members = {member.member_id}
        for seed, away in ((ends[0], ends[1]), (ends[1], ends[0])):
            node, previous, incoming = seed, away, member.member_id
            seen = {member.member_id}
            while True:
                if node in support_nodes or release_at(incoming, node):
                    break
                restraints = []
                for other in adjacency[node]:
                    if other == incoming:
                        continue
                    a, b = model.members[other]
                    far = b if a == node else a
                    delta = [model.nodes[far][i] - model.nodes[node][i] for i in range(3)]
                    if mode == 'columns':
                        is_restraint = other in primary and abs(delta[axis]) > .001
                    else:
                        is_restraint = (abs(delta[axis]) > .001 and
                                        (axis != 1 or members.get(other) is not None
                                         and members[other].column_y))
                    if is_restraint and not release_at(other, node):
                        restraints.append(other)
                if restraints:
                    break
                options = []
                direction = [model.nodes[node][i] - model.nodes[previous][i] for i in range(3)]
                for other in adjacency[node]:
                    if other in seen or other not in members or not members[other].steel:
                        continue
                    a, b = model.members[other]
                    far = b if a == node else a
                    delta = [model.nodes[far][i] - model.nodes[node][i] for i in range(3)]
                    dot = sum(x*y for x, y in zip(direction, delta))
                    if dot > 0 and dot / (math.dist(model.nodes[node], model.nodes[previous]) *
                                          members[other].length) >= .999999:
                        options.append((other, far))
                if len(options) != 1:
                    break
                other, far = options[0]
                if release_at(other, node):
                    break
                seen.add(other)
                span_members.add(other)
                total += members[other].length
                previous, node, incoming = node, far, other
        converted = service.length_to_meters(total)
        for mid in span_members:
            span_cache[mid, axis] = converted
        return converted

    for member in sorted(target, key=lambda m: m.member_id):
        beta = finite(service.os.Property.GetBetaAngle(member.member_id), 'Member beta') % 360
        turn = round(beta / 90)
        if abs(beta - turn * 90) > .001:
            raise ValueError(f'Member {member.member_id} has beta {beta:g}; review its local axes manually.')
        # Local-y buckling means displacement along local z. The column mapping
        # agrees with the retained vendor response, including beta=90 swaps.
        axes = [2, 0] if mode == 'columns' else [2 if member.beam_x else 0, 1]
        if turn % 2:
            axes.reverse()
        ly, lz = (span(member, axis) for axis in axes)
        result.extend((parameter, value, member.member_id)
                      for parameter, value in (('LY', ly), ('LZ', lz), ('UNL', ly)))
    if not result:
        raise ValueError('No eligible steel members were found. Check assigned materials.')
    return result


def _saved_units(model):
    content = Path(model.path).read_bytes().decode('utf-8', 'replace')
    units = re.findall(r'^\s*UNIT(?:S)?\s+(\w+)\s+(\w+)\s*$', content, re.M | re.I)
    if not units:
        raise ValueError('Cannot determine original STD units; add a standalone UNIT command first.')
    return 'UNIT ' + ' '.join(units[-1]).upper()


def process_design_parameters(expected, mode, code, report):
    """Return a complete preview block; no live model changes are made."""
    command = _code(code)
    if mode not in _MODES:
        raise ValueError('Choose columns, beam lengths or deflection joints.')
    from .staad_parameters import StaadService, SteelGenerator, Parameter, MAX_LINE_DEFAULT
    staad = live.connect()
    current = live.require_snapshot(staad, expected, report)
    service = StaadService()
    service.load_model(.001, fallback=False)
    if Path(service.model_path).resolve() != Path(current.path).resolve():
        raise RuntimeError('Active model changed. Process Members again.')
    report('stage', _MODES[mode])
    if mode == 'beams_dj':
        body = SteelGenerator(service, MAX_LINE_DEFAULT).generate([
            Parameter(True, 'DESIGN JOINT', None, 'all', '')])
        if not re.search(r'^DJ[12] ', body, re.M):
            raise ValueError('No eligible multi-segment steel beam chains were found.')
    else:
        rows = _physical_lengths(current, service, mode)
        if code == 'IS 800:2007':
            rows = [('LX' if name == 'UNL' else name, value, mid) for name, value, mid in rows]
        body = '\n'.join(f'{name} {fmt(value)} MEMB {mid}' for name, value, mid in rows)
    live.require_snapshot(staad, current, report)
    return (f'* CivTools — {_MODES[mode]} — {code}\n'
            '* Geometry candidates: review local axes, releases and actual flange restraints.\n'
            '* Physical ends are candidate restraints; no effective-length factor is inferred.\n'
            '* LX (2007) / UNL (1984) uses lateral restraint candidates; review actual LTB restraint.\n'
            'UNIT METER KN\nPARAMETER 1\nCODE ' + command + '\n' + body.rstrip() + '\n' + _saved_units(current) + '\n')


def _validate_parameter_text(text, code, model):
    """Restrict Assign to the parameters generated here, with bounded valid IDs."""
    command, records, headers = _code(code), [], []
    assignment_count = 0
    pending = ''
    restored = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith('*'):
            continue
        pending += (' ' if pending else '') + line.rstrip('-').strip()
        if line.endswith('-'):
            continue
        line, pending = pending, ''
        if restored:
            raise ValueError('No parameters may follow the unit restoration command.')
        if records and line.startswith('UNIT '):
            if line != _saved_units(model):
                raise ValueError('Restore the original STD units at the end of the generated block.')
            restored = True
            continue
        if line in ('UNIT METER KN', 'PARAMETER 1', 'CODE ' + command):
            headers.append(line)
            continue
        match = re.fullmatch(r'(LY|LZ|LX|UNL|DJ1|DJ2) ([0-9]+(?:\.[0-9]*)?(?:[eE][+-]?[0-9]+)?) MEMB ([0-9 ]+)', line)
        if not match:
            raise ValueError('Assign accepts only the generated LY/LZ/LX/UNL/DJ1/DJ2 block.')
        name, value, identifiers = match.groups()
        if (code == 'IS 800:2007' and name == 'UNL') or (code == 'IS 800:1984' and name == 'LX'):
            raise ValueError('Use LX for IS 800:2007 and UNL for IS 800:1984.')
        value = finite(value, name, True)
        ids = [positive_id(item) for item in identifiers.split()]
        if not ids or set(ids) - model.members.keys():
            raise ValueError('Parameter block references unknown members.')
        if name.startswith('DJ') and (value != int(value) or int(value) not in model.nodes):
            raise ValueError('Deflection joints must reference existing node IDs.')
        records.append((name, value, ids))
        assignment_count += len(ids)
        if assignment_count > 150_000:
            raise ValueError('Parameter block exceeds the supported size.')
    if pending or not records or headers != ['UNIT METER KN', 'PARAMETER 1', 'CODE ' + command]:
        raise ValueError('Generate a complete parameter block before assigning.')
    return records


def apply_design_parameters(expected, text, code, report):
    """Apply after the UI confirms save/close/reopen and review of the preview.

    Keep all existing design blocks. Insert a distinct numbered block before
    FINISH, restore the preceding UNIT, and retain a unique original-file backup.
    No undocumented numeric design-code IDs are assumed.
    """
    _validate_parameter_text(text, code, expected)
    staad = live.connect()
    live.require_snapshot(staad, expected, report)
    source = Path(expected.path)
    if staad.SaveModel(True) is False:
        raise RuntimeError('STAAD refused to save the model. Assignment was cancelled.')
    saved = live.capture(staad, report)
    if (Path(saved.path).resolve() != source.resolve() or saved.base_unit != expected.base_unit or
            saved.nodes != expected.nodes or saved.members != expected.members or
            saved.supports != expected.supports or saved.load_cases != expected.load_cases):
        raise RuntimeError('Model geometry, supports, units or load cases changed while saving. Process Members again.')
    # SaveModel may normalize formatting. Use the fresh saved identity as the
    # baseline after proving the model itself still matches the reviewed scan.
    saved_digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if saved.source_digest != saved_digest:
        raise RuntimeError('Saved model changed during assignment preparation.')
    original = source.read_bytes()
    try:
        content = original.decode('utf-8-sig')
        encoding = 'utf-8-sig' if original.startswith(b'\xef\xbb\xbf') else 'utf-8'
    except UnicodeDecodeError:
        content, encoding = original.decode('cp1252'), 'cp1252'
    finishes = list(re.finditer(r'^\s*FINISH\s*$', content, re.M | re.I))
    if len(finishes) != 1:
        raise ValueError('Assign requires exactly one standalone FINISH command in the STD file.')
    before, after = content[:finishes[0].start()], content[finishes[0].start():]
    units = re.findall(r'^\s*UNIT(?:S)?\s+(\w+)\s+(\w+)\s*$', before, re.M | re.I)
    if not units:
        raise ValueError('Cannot determine the original STD input units; copy the preview manually.')
    numbers = [int(n) for n in re.findall(r'^\s*PARAMETER\s+(\d+)\s*$', before, re.M | re.I)]
    block = text.replace('PARAMETER 1\n', f'PARAMETER {max(numbers, default=0)+1}\n')
    # Comments are ASCII to preserve an existing Windows ANSI STD encoding.
    block = block.encode('ascii', 'replace').decode('ascii')
    changed = (before.rstrip() + '\n\n' + block.rstrip() + '\nUNIT ' +
               ' '.join(units[-1]) + '\n' + after.lstrip()).replace('\r\n', '\n')
    newline = '\r\n' if b'\r\n' in original else '\n'
    changed_bytes = changed.replace('\n', newline).encode(encoding)
    backup = source.with_name(source.stem + '.civtools-' + uuid4().hex + '.std')
    shutil.copy2(source, backup)
    report('log', 'Backup: ' + str(backup))
    descriptor, staging = tempfile.mkstemp(suffix='.std', prefix='.civtools-', dir=source.parent)
    os.close(descriptor)
    closed, written = False, False
    try:
        Path(staging).write_bytes(changed_bytes)
        if hashlib.sha256(source.read_bytes()).hexdigest() != saved_digest:
            raise RuntimeError('The saved model changed before assignment. Process Members again.')
        if Path(str(staad.GetSTAADFile())).resolve() != source.resolve():
            raise RuntimeError('Active STAAD model changed before assignment.')
        staad.CloseSTAADFile()
        try:
            active = str(staad.GetSTAADFile() or '')
        except Exception:
            active = ''  # The documented API may error when no model is open.
        if active:
            raise RuntimeError('STAAD did not close the model; no file changes were applied.')
        closed = True
        if source.read_bytes() != original:
            raise RuntimeError('Saved model changed while closing; assignment was cancelled.')
        os.replace(staging, source)
        written = True
        staad.OpenSTAADFile(str(source))
        if Path(str(staad.GetSTAADFile())).resolve() != source.resolve():
            raise RuntimeError('STAAD did not reopen the assigned model.')
    except Exception as exc:
        # Never overwrite a model opened independently while this job ran.
        if closed:
            try:
                active = str(staad.GetSTAADFile() or '')
            except Exception:
                active = ''
            if not active:
                if written:
                    source.write_bytes(original)
                try:
                    staad.OpenSTAADFile(str(source))
                except Exception:
                    pass
        raise RuntimeError(f'Assignment failed; inspect the active model. Backup: {backup}. {exc}') from exc
    finally:
        if Path(staging).exists():
            Path(staging).unlink()
    return OperationResult(f'Assigned reviewed parameters. Re-run analysis and design. Backup: {backup}', str(backup))


def load_case_inventory(expected, report):
    staad = live.connect()
    current = live.require_snapshot(staad, expected, report)
    primary = set(current.primary_cases or [])
    inventory = []
    for number in current.load_cases:
        title = str(staad.Load.GetLoadCaseTitle(number) or '')
        inventory.append((number, 'Primary' if number in primary else 'Combination', title))
    live.require_snapshot(staad, current, report)
    return inventory


def export_results(expected, kind, nodes, cases, path, report, pedestal=None):
    """Read fresh bounded OpenSTAAD output, then publish a complete workbook."""
    if kind not in ('Reactions', 'Displacements', 'Pedestal'):
        raise ValueError('Choose reactions, displacements or pedestal forces.')
    staad = live.connect()
    current = live.require_snapshot(staad, expected, report)
    nodes, cases = sorted({positive_id(n) for n in nodes}), sorted({positive_id(c) for c in cases})
    if not nodes or not cases or set(nodes) - current.nodes.keys() or set(cases) - set(current.load_cases):
        raise ValueError('Choose valid nodes and load cases from a fresh scan.')
    if kind != 'Displacements' and set(nodes) - set(current.supports):
        raise ValueError('Reaction and pedestal exports require support nodes.')
    if len(nodes) * len(cases) > _MAX_PAIRS:
        raise ValueError('Limit an export to 100,000 node/load-case pairs.')
    if not staad.Output.AreResultsAvailable():
        raise RuntimeError('Run analysis in STAAD.Pro before exporting results.')
    getter = staad.Output.GetNodeDisplacements if kind == 'Displacements' else staad.Output.GetSupportReactions
    records = []
    report('stage', 'Reading analysis results')
    for node in nodes:
        for case in cases:
            try:
                raw = getter(node, case)
                if not isinstance(raw, (tuple, list)) or len(raw) != 6:
                    raise ValueError('expected six output components')
                values = tuple(finite(v, 'STAAD result') for v in raw)
            except Exception as exc:
                raise RuntimeError(f'Cannot read {kind.lower()} for node {node}, case {case}: {exc}') from exc
            records.append((node, case, values))
            if len(records) % 50 == 0:
                report('progress', int(100 * len(records) / (len(nodes) * len(cases))))
    live.require_snapshot(staad, current, report)
    if not staad.Output.AreResultsAvailable():
        raise RuntimeError('Results became unavailable during export.')
    write_result_workbook(path, current, kind, records, pedestal)
    return OperationResult(f'Exported {len(records):,} rows with PivotTables and component summaries to {path}.', str(path))


def _add_result_pivot(workbook, source, title, pivot_name):
    """Build an OOXML PivotTable with saved numeric cache and populated display.

    Node is the row field. Each numeric output component has MIN and MAX data
    fields, with load case retained in the cache for subsequent Excel filtering.
    The worksheet range is explicit and bounded to the exported records.
    """
    from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource, CacheField, SharedItems
    from openpyxl.pivot.fields import Number, Index
    from openpyxl.pivot.record import Record, RecordList
    from openpyxl.pivot.table import (TableDefinition, Location, DataField, RowColField,
                                     RowColItem, PivotField, FieldItem, PivotTableStyle)
    from openpyxl.utils import get_column_letter
    rows = list(source.values)
    headings, values = rows[0], rows[1:]
    nodes = sorted({row[0] for row in values})
    indexes = {node: index for index, node in enumerate(nodes)}
    fields = []
    for column, name in enumerate(headings):
        column_values = [row[column] for row in values]
        shared = SharedItems(_fields=[Number(v=node) for node in nodes] if column == 0 else [],
                             containsNumber=True, containsInteger=all(float(v).is_integer() for v in column_values),
                             containsString=False, containsBlank=False, containsNonDate=True,
                             containsDate=False, containsMixedTypes=False, containsSemiMixedTypes=False,
                             minValue=min(column_values), maxValue=max(column_values))
        fields.append(CacheField(name=str(name), sharedItems=shared, numFmtId=0))
    cache = CacheDefinition(cacheSource=CacheSource(type='worksheet', worksheetSource=WorksheetSource(
                                sheet=source.title, ref=source.dimensions)),
                            cacheFields=fields, recordCount=len(values), saveData=True,
                            refreshOnLoad=True, enableRefresh=True, createdVersion=6,
                            refreshedVersion=6, minRefreshableVersion=3)
    cache.records = RecordList(r=[Record(_fields=[Index(v=indexes[row[0]])] +
                                        [Number(v=value) for value in row[1:]]) for row in values])
    data = [DataField(name=f'{operation.upper()} {headings[column]}', fld=column, subtotal=operation)
            for column in range(2, len(headings)) for operation in ('min', 'max')]
    pivot_fields = [PivotField(axis='axisRow', showAll=False, defaultSubtotal=False,
                               items=[FieldItem(x=index) for index in range(len(nodes))])]
    pivot_fields.extend(PivotField(dataField=column >= 2, showAll=False, defaultSubtotal=False)
                        for column in range(1, len(headings)))
    destination = workbook.create_sheet(title)
    # Save an immediately usable display. Excel's pivot engine can refresh it
    # from the embedded source after editing or changing fields in its UI.
    destination.append(['Node'] + [field.name for field in data])
    by_node = defaultdict(list)
    for row in values:
        by_node[row[0]].append(row)
    for node in nodes:
        destination.append([node] + [min(row[field.fld] for row in by_node[node])
                                     if field.subtotal == 'min' else max(row[field.fld] for row in by_node[node])
                                     for field in data])
    cache_id = sum(len(sheet._pivots) for sheet in workbook) + 1
    # openpyxl creates the cache-record relationship before copying the cache
    # identifier onto RecordList; preassign it so later caches link correctly.
    cache.records._id = cache_id
    pivot = TableDefinition(name=pivot_name, cacheId=cache_id, dataCaption='Values', rowHeaderCaption='Node',
                            location=Location(ref=f'A1:{get_column_letter(len(data)+1)}{len(nodes)+1}',
                                              firstHeaderRow=0, firstDataRow=1, firstDataCol=1),
                            pivotFields=pivot_fields, rowFields=[RowColField(x=0)],
                            rowItems=[RowColItem(x=[Index(v=index)]) for index in range(len(nodes))],
                            colFields=[RowColField(x=-2)],
                            colItems=[RowColItem(i=index, x=[Index(v=index)]) for index in range(len(data))],
                            dataFields=data, rowGrandTotals=False, colGrandTotals=False,
                            compact=False, compactData=False, outline=False,
                            createdVersion=6, updatedVersion=6, minRefreshableVersion=3,
                            pivotTableStyleInfo=PivotTableStyle(name='PivotStyleMedium9',
                                                               showRowHeaders=True, showColHeaders=True,
                                                               showRowStripes=True, showColStripes=False,
                                                               showLastColumn=True))
    pivot.cache = cache
    destination.add_pivot(pivot)
    return destination


def write_result_workbook(path, model, kind, records, pedestal=None):
    """Atomic XLSX export with signed envelopes and editable pedestal formulas.

    Reaction and per-level displacement pivots retain populated numeric caches
    and request refresh on opening in Excel. Ordinary computed summaries are
    available immediately. Pedestal formulas request spreadsheet recalculation.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or kind not in ('Reactions', 'Displacements', 'Pedestal'):
        raise ValueError('Choose a supported result type and an .xlsx output file.')
    if not records or len(records) > _MAX_PAIRS:
        raise ValueError('Export between one and 100,000 result rows.')
    validated = []
    for node, case, values in records:
        node, case = positive_id(node), positive_id(case)
        if node not in model.nodes or case not in model.load_cases or len(values) != 6:
            raise ValueError('Result records must match the scanned nodes and load cases.')
        validated.append((node, case, tuple(finite(v, 'Result') for v in values)))
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    try:
        info = wb.active
        info.title = 'Read me'
        length, force = model.length_unit, model.force_unit
        details = [('Model', model.path), ('Snapshot', model.fingerprint), ('Result', kind),
                   ('Axes', 'Global; Y is vertical'), ('Length', length), ('Force', force),
                   ('Moment', f'{force} {length}'), ('Rotation', 'radians'),
                   ('Summaries', 'Excel PivotTables refresh on opening in Excel; computed min/max summaries are also included.'),
                   ('Scope', 'Exported nodes and load cases only. No design checks.'),
                   ('Recalculation', 'Open in Excel/LibreOffice to calculate workbook formulas.')]
        for row in details:
            info.append(row)
        # All external strings stay text, including formula-looking model paths.
        for row in info:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        components = ['DX', 'DY', 'DZ', 'RX', 'RY', 'RZ'] if kind == 'Displacements' else ['FX', 'FY', 'FZ', 'MX', 'MY', 'MZ']
        units = [length]*3 + ['rad']*3 if kind == 'Displacements' else [force]*3 + [f'{force} {length}']*3
        headings = ['Node', 'Load case'] + [f'{name} ({unit})' for name, unit in zip(components, units)]
        raw = wb.create_sheet('Results')
        resultant_headings = [f'Resultant ({length})', f'Horizontal resultant XZ ({length})']
        raw.append(headings + (resultant_headings if kind == 'Displacements' else []))
        envelopes = defaultdict(list)
        levels = defaultdict(list)
        for node, case, values in validated:
            row = [node, case, *values]
            if kind == 'Displacements':
                row.append(math.sqrt(sum(v*v for v in values[:3])))
                row.append(math.hypot(values[0], values[2]))
                levels[round(model.nodes[node][1], 6)].append(row)
            raw.append(row)
            envelopes[node].append((case, values))
        summary = wb.create_sheet('Envelopes')
        summary.append(['Node', 'Component', 'Unit', 'Minimum', 'Case at minimum', 'Maximum', 'Case at maximum'])
        for node, rows in sorted(envelopes.items()):
            for i, (component, unit) in enumerate(zip(components, units)):
                low, high = min(rows, key=lambda r: r[1][i]), max(rows, key=lambda r: r[1][i])
                summary.append([node, component, unit, low[1][i], low[0], high[1][i], high[0]])
        if kind == 'Displacements':
            level_summary = wb.create_sheet('Level summary')
            level_summary.append([f'Level Y ({length})', 'Component', 'Unit', 'Minimum', 'Node at minimum', 'Case at minimum',
                                  'Maximum', 'Node at maximum', 'Case at maximum'])
            if len(levels) > 250:
                raise ValueError('Limit a per-level displacement workbook to 250 distinct levels.')
            for index, (elevation, rows) in enumerate(sorted(levels.items()), 1):
                sheet = wb.create_sheet(f'Level {index} Y={elevation:g}'[:31])
                sheet.append(headings + resultant_headings)
                for row in rows:
                    sheet.append(row)
                for column, (component, unit) in enumerate(zip(components + ['Resultant', 'Horizontal resultant XZ'], units + [length, length]), 2):
                    low, high = min(rows, key=lambda r: r[column]), max(rows, key=lambda r: r[column])
                    level_summary.append([elevation, component, unit, low[column], low[0], low[1], high[column], high[0], high[1]])
                _add_result_pivot(wb, sheet, f'Pivot level {index}', f'CivDisplacements{index}')
        else:
            _add_result_pivot(wb, raw, 'Reaction pivot', 'CivSupportReactions')
        if kind == 'Pedestal':
            if pedestal is None:
                pedestal = (1., 1., 1., 24.) if model.base_unit == 'Metric' else (39.37007874,)*3 + (.0000885181526,)
            if len(pedestal) != 4:
                raise ValueError('Provide pedestal height, width X, width Z and weight density.')
            pedestal = tuple(finite(v, 'Pedestal input', True) for v in pedestal)
            inputs = wb.create_sheet('Pedestal inputs')
            inputs.append(['Input', 'Value', 'Unit'])
            for name, value, unit in zip(['Height', 'Width X', 'Width Z', 'Weight density'], pedestal,
                                         [length, length, length, f'{force}/{length}^3']):
                inputs.append([name, value, unit])
                inputs.cell(inputs.max_row, 2).fill = PatternFill('solid', fgColor='FFF2CC')
            inputs.append(['Self-weight', '=B2*B3*B4*B5', force])
            from openpyxl.worksheet.datavalidation import DataValidation
            validation = DataValidation(type='decimal', operator='greaterThan', formula1=0)
            validation.error, validation.errorTitle = 'Enter a positive number.', 'Pedestal input'
            validation.showErrorMessage, validation.errorStyle = True, 'stop'
            inputs.add_data_validation(validation)
            validation.add('B2:B5')
            final = wb.create_sheet('Pedestal forces')
            final.append(headings)
            for row in range(2, len(validated)+2):
                final.append([f"='Results'!A{row}", f"='Results'!B{row}", f"='Results'!C{row}",
                              f"='Results'!D{row}+'Pedestal inputs'!$B$6", f"='Results'!E{row}",
                              f"='Results'!F{row}+'Pedestal inputs'!$B$2*'Results'!E{row}",
                              f"='Results'!G{row}", f"='Results'!H{row}-'Pedestal inputs'!$B$2*'Results'!C{row}"])
            pedestal_summary = wb.create_sheet('Pedestal envelopes')
            pedestal_summary.append(['Component', 'Unit', 'Minimum', 'Maximum'])
            for col, component, unit in zip('CDEFGH', components, units):
                address = f"'Pedestal forces'!{col}2:{col}{len(validated)+1}"
                pedestal_summary.append([component, unit, f'=MIN({address})', f'=MAX({address})'])
            info.append(['Pedestal convention', 'Signed reactions: Fyb=Fy+weight; Mxb=Mx+h*Fz; Myb=My; Mzb=Mz-h*Fx.'])
            info.append(['Foundation actions', 'Negate all six final reactions to obtain actions on the foundation.'])
            info.append(['Pedestal inputs', 'Edit yellow cells. Density is weight per unit volume, not mass density.'])
        for index, sheet in enumerate(wb, 1):
            sheet.freeze_panes = 'A2'
            if not sheet._pivots:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill('solid', fgColor='E8F1FB')
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(65, max(18, max(len(str(cell.value or '')) for cell in column) + 2))
            if sheet.title not in ('Read me', 'Pedestal inputs') and sheet.max_row > 1 and not sheet._pivots:
                table = Table(displayName=f'CivResults{index}', ref=sheet.dimensions)
                table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
                sheet.add_table(table)
        descriptor, staging = tempfile.mkstemp(suffix='.xlsx', prefix='.civtools-', dir=path.parent)
        os.close(descriptor)
        try:
            wb.save(staging)
            os.replace(staging, path)
        finally:
            if Path(staging).exists():
                Path(staging).unlink()
    finally:
        wb.close()
