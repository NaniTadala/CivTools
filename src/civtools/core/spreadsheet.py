"""One workbook interpretation for the sample table and AutoCAD mapping."""
import math
import openpyxl

MAX_MAPPING_ROWS = 100000


def read_mapping(path, sheet_name=None):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        name = sheet_name or workbook.sheetnames[0]
        if name not in workbook.sheetnames:
            raise ValueError(f"Worksheet '{name}' is no longer available.")
        sheet = workbook[name]
        rows, sample, valid, skipped, invalid = [], [], 0, 0, 0
        for number, values in enumerate(sheet.iter_rows(min_row=2, max_col=5, values_only=True), 2):
            if number > MAX_MAPPING_ROWS + 1:
                raise ValueError(f"Use at most {MAX_MAPPING_ROWS:,} spreadsheet rows per mapping job.")
            status, data = "Valid", None
            if all(value is None for value in values):
                skipped += 1
                status = "Blank row"
            else:
                try:
                    coords = tuple(float(value) for value in values[:4])
                    if not all(math.isfinite(value) for value in coords):
                        raise ValueError("Use finite numbers in A–D")
                    if coords[2] <= 0 or coords[3] <= 0:
                        raise ValueError("Width and height must be positive")
                    data = coords + ("" if values[4] is None else str(values[4]),)
                    valid += 1
                except (ValueError, TypeError) as error:
                    invalid += 1
                    status = str(error) if isinstance(error, ValueError) and "could not convert" not in str(error) else "Enter numbers in columns A–D"
            rows.append((number, data, status))
            if len(sample) < 8:
                sample.append((number, values, status))
        return {"sheet_names": workbook.sheetnames, "sheet": name, "rows": rows, "sample": sample,
                "valid": valid, "skipped": skipped, "invalid": invalid}
    finally:
        workbook.close()


def inspect_mapping(path, sheet_name, report):
    result = read_mapping(path, sheet_name)
    del result["rows"]
    return result
