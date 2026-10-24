"""Capture empty, prepared, compact and settings states using synthetic reports."""
import os
from pathlib import Path
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))


def main():
    os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
    output = ROOT/'artifacts'/'publisher-polish'
    output.mkdir(parents=True,exist_ok=True)
    os.environ['CIVTOOLS_CONFIG_DIR'] = str(output/'config')
    from PySide6.QtWidgets import QApplication
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from civtools.qt_app import CivToolsApp
    app=QApplication([])
    window=CivToolsApp()
    window.navigate('reports')
    page=window.pages['reports']
    if '--reuse' in sys.argv and page.prepared:
        utility=window.utility_windows['reports']
        page.combined_toggle.setChecked('--combined' in sys.argv)
        page.tree.setCurrentItem(page.tree.topLevelItem(0).child(0))
        for width,height,name in [(1000,700,'prepared'),(780,540,'compact')]:
            utility.resize(width,height)
            app.processEvents()
            deadline=time.monotonic()+20
            while time.monotonic()<deadline:
                app.processEvents()
                time.sleep(.02)
                if not page.jobs and not page.thumbnail_timer.isActive() and not page.preview_resize_timer.isActive() and not page.combined_timer.isActive():
                    break
            print(name,page.preview.size(),page.preview_scroll.viewport().size(),page.preview_scroll.horizontalScrollBar().maximum())
            if '--combined' in sys.argv:
                name='combined-'+name
            utility.grab().save(str(output/(name+'.png')))
        utility.resize(1000,700)
        page.settings_dialog.show()
        for index,name in [(0,'settings'),(1,'export-options'),(2,'advanced')]:
            page.editor.setCurrentIndex(index)
            app.processEvents()
            page.settings_dialog.grab().save(str(output/(name+'.png')))
        page.settings_dialog.hide()
        window.close()
        return
    page.records.clear()
    page.prepared.clear()
    page.rebuild()
    utility=window.utility_windows['reports']
    utility.resize(1000,700)
    app.processEvents()
    utility.grab().save(str(output/'empty.png'))
    samples=output/'samples'
    samples.mkdir(exist_ok=True)
    files=[]
    for name in ('Foundation design','Base plate check','Steel connection'):
        book=Workbook()
        sheet=book.active
        sheet.title='Calculation'
        sheet.merge_cells('A1:D1')
        sheet['A1']=name.upper()
        sheet['A1'].font=Font(name='Calibri',size=17,bold=True,color='193E58')
        sheet.merge_cells('A2:D2')
        sheet['A2']='SAMPLE REPORT  /  PAGE LAYOUT PREVIEW'
        sheet['A2'].font=Font(size=9,color='7A8C99')
        for row,heading in [(5,'01   DESIGN BASIS'),(15,'02   DESIGN ACTIONS'),(25,'03   VERIFICATION')]:
            sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=4)
            cell=sheet.cell(row,1,heading)
            cell.font=Font(bold=True,size=11,color='193E58')
            cell.fill=PatternFill('solid',fgColor='EAF1F6')
            sheet.row_dimensions[row].height=25
        values=[('Concrete grade','M30',''),('Steel grade','Fe 500',''),('Foundation width',2400,'mm'),('Foundation length',3000,'mm'),
                ('Allowable bearing pressure',180,'kN/m²')]
        for row,(label,value,unit) in enumerate(values,7):
            sheet.cell(row,1,label)
            sheet.cell(row,3,value)
            sheet.cell(row,4,unit)
        for row,label,value,unit in [(17,'Axial load',650,'kN'),(18,'Shear force',42,'kN'),(19,'Bending moment',85,'kN·m'),
                                     (27,'Bearing pressure',128.4,'kN/m²'),(28,'Utilisation',.71,''),(30,'Result','PASS','')]:
            sheet.cell(row,1,label)
            sheet.cell(row,3,value)
            sheet.cell(row,4,unit)
        sheet['C28'].number_format='0%'
        sheet['C30'].font=Font(bold=True,color='258568')
        for col,width in [('A',28),('B',9),('C',16),('D',14)]:
            sheet.column_dimensions[col].width=width
        for row in range(1,33):
            if row not in (1,5,15,25):
                sheet.row_dimensions[row].height=21
        wide=book.create_sheet('Load summary')
        wide.append(['LOAD CASE','Fx','Fy','Fz','Mx','My','Mz','Resultant','Factor','Status'])
        for cell in wide[1]:
            cell.font=Font(bold=True,color='193E58')
        for i in range(1,20):
            wide.append([f'Combination {i}',40,650,12,85,15,20,652,1.5,'PASS'])
        for col in 'ABCDEFGHIJ':
            wide.column_dimensions[col].width=14
        path=samples/(name+'.xlsx')
        book.save(path)
        files.append(str(path))
    def wait(predicate,timeout=180):
        deadline=time.monotonic()+timeout
        while not predicate() and time.monotonic()<deadline:
            app.processEvents()
            time.sleep(.01)
        if not predicate():
            raise RuntimeError(page.feedback.text())
    page.output.setText(str(output/'PDFs'))
    page.intake(files)
    wait(lambda:not page.busy)
    page.prepare()
    wait(lambda:not page.busy)
    first=page.tree.topLevelItem(0)
    page.tree.setCurrentItem(first.child(0))
    wait(lambda:page.preview.pixmap() is not None and not page.preview.pixmap().isNull())
    wait(lambda:not page.jobs and not page.thumbnail_timer.isActive() and not page.preview_resize_timer.isActive())
    app.processEvents()
    utility.grab().save(str(output/'prepared.png'))
    utility.resize(780,540)
    app.processEvents()
    wait(lambda:not page.jobs and not page.preview_resize_timer.isActive())
    utility.grab().save(str(output/'compact.png'))
    utility.resize(1000,700)
    page.settings_dialog.show()
    app.processEvents()
    page.settings_dialog.grab().save(str(output/'settings.png'))
    page.settings_dialog.hide()
    wait(lambda:not page.jobs and not page.preview_resize_timer.isActive())
    print('Screenshots:',output)
    print('Statuses:',[(Path(k).name,v['status']) for k,v in page.records.items()])
    window.close()


if __name__=='__main__':
    import multiprocessing
    multiprocessing.freeze_support()
    main()
