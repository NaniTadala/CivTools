import hashlib
from pathlib import Path
import re
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from civtools.core.staad_productivity import ModelSnapshot
from civtools.core import productivity_design_results as tools


def model(path='example.std', unit='Metric'):
    return ModelSnapshot(str(path), unit,
        {1: (0, 0, 0), 2: (0, 3, 0), 3: (0, 6, 0),
         4: (5, 3, 0), 5: (5, 0, 0), 6: (0, 6, 5), 7: (0, 0, 5)},
        {10: (1, 2), 11: (2, 3), 20: (2, 4), 30: (4, 5), 40: (3, 6), 50: (6, 7)},
        [1, 5, 7], [1, 2, 101], primary_cases=[1, 2])


class DesignCandidatesTests(unittest.TestCase):
    def service(self, expected, beta=0):
        from civtools.core.staad_parameters import Member
        import math
        members = []
        for mid, (a, b) in expected.members.items():
            dx, dy, dz = [abs(x-y) for x, y in zip(expected.nodes[a], expected.nodes[b])]
            members.append(Member(mid, a, b, math.dist(expected.nodes[a], expected.nodes[b]),
                                  'I', 'STEEL', False, True, dx == dz == 0,
                                  dy == dz == 0, dx == dy == 0))
        staad = MagicMock()
        staad.Property.GetBetaAngle.return_value = beta
        return SimpleNamespace(members=members, os=staad, member_released_at_end=lambda *args: False,
                               length_to_meters=lambda value: value)

    def test_split_column_uses_directional_physical_restraints_and_beta(self):
        expected = model()
        service = self.service(expected)
        rows = {(name, mid): value for name, value, mid in tools._physical_lengths(expected, service, 'columns')}
        self.assertEqual(rows['LY', 10], 6)
        self.assertEqual(rows['LZ', 10], 3)
        self.assertEqual(rows['UNL', 10], 6)
        service.os.Property.GetBetaAngle.return_value = 90
        swapped = {(name, mid): value for name, value, mid in tools._physical_lengths(expected, service, 'columns')}
        self.assertEqual(swapped['LY', 10], 3)
        self.assertEqual(swapped['LZ', 10], 6)

    def test_nonorthogonal_beta_is_rejected(self):
        expected = model()
        with self.assertRaisesRegex(ValueError, 'local axes'):
            tools._physical_lengths(expected, self.service(expected, beta=30), 'columns')

    def test_release_terminates_physical_column_trace(self):
        expected = model()
        service = self.service(expected)
        service.member_released_at_end = lambda mid, end: mid == 10 and end == 1
        rows = {(name, mid): value for name, value, mid in tools._physical_lengths(expected, service, 'columns')}
        self.assertEqual(rows['LY', 10], 3)

    def block(self):
        return '* reviewed\nUNIT METER KN\nPARAMETER 1\nCODE IS800 LSD\nLY 6 MEMB 10 11\nDJ1 1 MEMB 10\n'

    def test_assign_rejects_injected_commands_wrong_codes_and_unknown_joints(self):
        for text in (self.block()+'DELETE MEMBER 10\n', self.block().replace('IS800 LSD', 'INDIAN'),
                     self.block().replace('DJ1 1 ', 'DJ1 999 '), self.block()+'UNL 3 MEMB 10\n'):
            with self.subTest(text=text), patch.object(tools.live, 'connect') as connection:
                with self.assertRaises(ValueError):
                    tools.apply_design_parameters(model(), text, 'IS 800:2007', lambda *args: None)
                connection.assert_not_called()

    def fake_staad(self, path):
        staad = MagicMock()
        active = {'path': str(path)}
        staad.GetSTAADFile.side_effect = lambda: active['path']
        staad.CloseSTAADFile.side_effect = lambda: active.update(path='')
        staad.OpenSTAADFile.side_effect = lambda name: active.update(path=name)
        return staad, active

    def test_assign_preserves_existing_blocks_units_and_original_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'model.std'
            original = b'STAAD SPACE\r\nUNIT FEET KIP\r\nPERFORM ANALYSIS\r\nPARAMETER 7\r\nCODE INDIAN\r\nLY 1 MEMB 10\r\nFINISH\r\n'
            path.write_bytes(original)
            expected = model(path)
            expected.source_digest = hashlib.sha256(original).hexdigest()
            staad, _ = self.fake_staad(path)
            with patch.object(tools.live, 'connect', return_value=staad), patch.object(tools.live, 'require_snapshot', return_value=expected), patch.object(tools.live, 'capture', return_value=expected):
                result = tools.apply_design_parameters(expected, self.block(), 'IS 800:2007', lambda *args: None)
            self.assertEqual(Path(result.output_path).read_bytes(), original)
            content = path.read_text()
            self.assertIn('PARAMETER 7\nCODE INDIAN', content)
            self.assertIn('PARAMETER 8\nCODE IS800 LSD', content)
            self.assertTrue(content.endswith('UNIT FEET KIP\nFINISH\n'))
            staad.CloseSTAADFile.assert_called_once()
            staad.OpenSTAADFile.assert_called_once_with(str(path))

    def test_assignment_reopen_failure_restores_original_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'model.std'
            original = b'STAAD SPACE\nUNIT METER KN\nPERFORM ANALYSIS\nFINISH\n'
            path.write_bytes(original)
            expected = model(path)
            expected.source_digest = hashlib.sha256(original).hexdigest()
            staad, active = self.fake_staad(path)
            attempts = []
            def reopen(name):
                attempts.append(name)
                if len(attempts) == 1:
                    raise RuntimeError('Cannot read generated model')
                active['path'] = name
            staad.OpenSTAADFile.side_effect = reopen
            with patch.object(tools.live, 'connect', return_value=staad), patch.object(tools.live, 'require_snapshot', return_value=expected), patch.object(tools.live, 'capture', return_value=expected):
                with self.assertRaisesRegex(RuntimeError, 'Backup:'):
                    tools.apply_design_parameters(expected, self.block(), 'IS 800:2007', lambda *args: None)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(active['path'], str(path))

    def test_code_dialect_changes_ltb_parameter_and_preview_restores_input_units(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'model.std'
            path.write_text('STAAD SPACE\nUNIT FEET KIP\nFINISH\n')
            expected = model(path)
            service = self.service(expected)
            service.model_path = str(path)
            service.load_model = MagicMock()
            with patch.object(tools.live, 'connect'), patch.object(tools.live, 'require_snapshot', return_value=expected), patch('civtools.core.staad_parameters.StaadService', return_value=service):
                modern = tools.process_design_parameters(expected, 'columns', 'IS 800:2007', lambda *args: None)
                legacy = tools.process_design_parameters(expected, 'columns', 'IS 800:1984', lambda *args: None)
            self.assertRegex(modern, r'(?m)^LX 6 MEMB 10$')
            self.assertNotRegex(modern, r'(?m)^UNL ')
            self.assertRegex(legacy, r'(?m)^UNL 6 MEMB 10$')
            self.assertNotRegex(legacy, r'(?m)^LX ')
            self.assertTrue(modern.endswith('UNIT FEET KIP\n'))
            tools._validate_parameter_text(modern, 'IS 800:2007', expected)

    def test_assignment_accepts_only_serialization_changes_after_save(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'model.std'
            original = b'STAAD SPACE\nUNIT METER KN\nFINISH\n'
            path.write_bytes(original)
            expected = model(path)
            expected.source_digest = hashlib.sha256(original).hexdigest()
            saved = model(path)
            normalized = original.replace(b'FINISH', b'* saved formatting\nFINISH')
            saved.source_digest = hashlib.sha256(normalized).hexdigest()
            staad, _ = self.fake_staad(path)
            staad.SaveModel.side_effect = lambda silent: path.write_bytes(normalized)
            with patch.object(tools.live, 'connect', return_value=staad), patch.object(tools.live, 'require_snapshot', return_value=expected), patch.object(tools.live, 'capture', return_value=saved):
                result = tools.apply_design_parameters(expected, self.block(), 'IS 800:2007', lambda *args: None)
            self.assertEqual(Path(result.output_path).read_bytes(), normalized)
            self.assertIn('* saved formatting', path.read_text())

    def test_explicit_save_failure_cancels_assignment_before_closing_or_backup(self):
        staad = MagicMock()
        staad.SaveModel.return_value = False
        with patch.object(tools.live, 'connect', return_value=staad), patch.object(tools.live, 'require_snapshot'), patch.object(tools.live, 'capture') as capture:
            with self.assertRaisesRegex(RuntimeError, 'refused to save'):
                tools.apply_design_parameters(model(), self.block(), 'IS 800:2007', lambda *args: None)
            capture.assert_not_called()
            staad.CloseSTAADFile.assert_not_called()


class ResultsWorkbookTests(unittest.TestCase):
    def test_displacements_group_levels_and_report_actual_resultant_maximum(self):
        from openpyxl import load_workbook
        expected = model()
        records = [(2, 1, (3, 4, 0, 0, 0, 0)), (4, 2, (-6, 8, 0, .1, .2, .3)),
                   (3, 1, (0, 0, 2, 0, 0, 0))]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'displacements.xlsx'
            tools.write_result_workbook(path, expected, 'Displacements', records)
            wb = load_workbook(path, data_only=False)
            try:
                level_names = [name for name in wb.sheetnames if name.startswith('Level ') and name != 'Level summary']
                self.assertEqual(len(level_names), 2)
                self.assertEqual(wb['Results']['I3'].value, 10)
                summaries = list(wb['Level summary'].values)
                resultant = next(row for row in summaries[1:] if row[0] == 3 and row[1] == 'Resultant')
                self.assertEqual(resultant[6:], (10, 4, 2))
                self.assertIn('PivotTables refresh', wb['Read me']['B9'].value)
                self.assertEqual(wb['Results']['J3'].value, 6)
                horizontal = next(row for row in summaries[1:] if row[0] == 3 and row[1] == 'Horizontal resultant XZ')
                self.assertEqual(horizontal[6:], (6, 4, 2))
            finally:
                wb.close()

    def test_pedestal_inputs_update_signed_moments_weight_and_torsion_through_formulas(self):
        from openpyxl import load_workbook
        expected = model(unit='English')
        expected.path = '=HYPERLINK("https://example.invalid")'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'pedestal.xlsx'
            tools.write_result_workbook(path, expected, 'Pedestal', [(1, 1, (2, -4, 3, -5, 7, 11))], (2, 3, 4, .1))
            wb = load_workbook(path, data_only=False)
            try:
                inputs, final = wb['Pedestal inputs'], wb['Pedestal forces']
                self.assertEqual(inputs['B6'].value, '=B2*B3*B4*B5')
                self.assertEqual(inputs['C5'].value, 'kip/in^3')
                self.assertEqual(final['F2'].value, "='Results'!F2+'Pedestal inputs'!$B$2*'Results'!E2")
                self.assertEqual(final['G2'].value, "='Results'!G2")
                self.assertEqual(final['H2'].value, "='Results'!H2-'Pedestal inputs'!$B$2*'Results'!C2")
                self.assertEqual(wb['Read me']['B1'].data_type, 's')
                self.assertTrue(wb.calculation.fullCalcOnLoad)
                self.assertEqual(wb['Pedestal envelopes']['C2'].value, "=MIN('Pedestal forces'!C2:C2)")
            finally:
                wb.close()

    def test_failed_workbook_validation_preserves_existing_output(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'existing.xlsx'
            path.write_bytes(b'original')
            with self.assertRaisesRegex(ValueError, 'positive'):
                tools.write_result_workbook(path, model(), 'Pedestal', [(1, 1, (0,)*6)], (0, 1, 1, 24))
            self.assertEqual(path.read_bytes(), b'original')

    def test_actual_com_failure_identifies_node_and_case_without_publishing(self):
        staad = MagicMock()
        staad.Output.AreResultsAvailable.return_value = True
        staad.Output.GetSupportReactions.side_effect = RuntimeError('OpenSTAAD general error')
        expected = model()
        with patch.object(tools.live, 'connect', return_value=staad), patch.object(tools.live, 'require_snapshot', return_value=expected), patch.object(tools, 'write_result_workbook') as publish:
            with self.assertRaisesRegex(RuntimeError, 'node 1, case 101'):
                tools.export_results(expected, 'Reactions', [1], [101], 'out.xlsx', lambda *args: None)
            publish.assert_not_called()

    def test_load_inventory_preserves_primary_combination_and_case_titles(self):
        staad = MagicMock()
        staad.Load.GetLoadCaseTitle.side_effect = lambda case: f'Case {case}'
        expected = model()
        with patch.object(tools.live, 'connect', return_value=staad), patch.object(tools.live, 'require_snapshot', return_value=expected):
            self.assertEqual(tools.load_case_inventory(expected, lambda *args: None),
                             [(1, 'Primary', 'Case 1'), (2, 'Primary', 'Case 2'), (101, 'Combination', 'Case 101')])

    def test_refreshable_pivots_have_numeric_caches_unique_ids_and_bounded_sources(self):
        from openpyxl import load_workbook
        from xml.etree import ElementTree as ET
        from zipfile import ZipFile
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        records = [(2, 1, (3, 4, 0, 0, 0, 0)), (4, 2, (-6, 8, 0, .1, .2, .3)),
                   (3, 1, (0, 0, 2, 0, 0, 0))]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'pivots.xlsx'
            tools.write_result_workbook(path, model(), 'Displacements', records)
            with ZipFile(path) as archive:
                names = archive.namelist()
                definitions = sorted(name for name in names if re.fullmatch(r'xl/pivotTables/pivotTable\d+\.xml', name))
                caches = sorted(name for name in names if re.fullmatch(r'xl/pivotCache/pivotCacheDefinition\d+\.xml', name))
                cache_records = sorted(name for name in names if re.fullmatch(r'xl/pivotCache/pivotCacheRecords\d+\.xml', name))
                self.assertEqual((len(definitions), len(caches), len(cache_records)), (2, 2, 2))
                self.assertEqual({ET.fromstring(archive.read(name)).attrib['cacheId'] for name in definitions}, {'1', '2'})
                for name in definitions:
                    pivot = ET.fromstring(archive.read(name))
                    data = pivot.find('s:dataFields', ns)
                    self.assertEqual(len(data), 16)
                    self.assertIn(('MAX Horizontal resultant XZ (m)', 'max'), {(field.attrib['name'], field.attrib['subtotal']) for field in data})
                    self.assertIn('xl/pivotTables/_rels/' + name.rsplit('/', 1)[1] + '.rels', names)
                for name in caches:
                    cache = ET.fromstring(archive.read(name))
                    self.assertEqual(cache.attrib['refreshOnLoad'], '1')
                    self.assertEqual(cache.attrib['enableRefresh'], '1')
                    source = cache.find('s:cacheSource/s:worksheetSource', ns)
                    self.assertTrue(source.attrib['sheet'].startswith('Level '))
                    self.assertRegex(source.attrib['ref'], r'^A1:J[23]$')
                    for field in cache.find('s:cacheFields', ns):
                        items = field.find('s:sharedItems', ns)
                        self.assertEqual(items.attrib['containsNumber'], '1')
                        self.assertEqual(items.attrib['containsString'], '0')
                self.assertIn('xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels', names)
            wb = load_workbook(path)
            try:
                pivots = [pivot for sheet in wb for pivot in sheet._pivots]
                self.assertEqual(len(pivots), 2)
                self.assertEqual({pivot.cacheId for pivot in pivots}, {1, 2})
                self.assertEqual(sum(len(pivot.cache.records.r) for pivot in pivots), 3)
                self.assertTrue(all(pivot.cache.refreshOnLoad for pivot in pivots))
            finally:
                wb.close()

    def test_support_pivot_aggregates_signed_component_minima_and_maxima(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'support.xlsx'
            tools.write_result_workbook(path, model(), 'Reactions',
                [(1, 1, (-10, 5, 4, -3, 8, 2)), (1, 2, (12, -6, -1, 9, -7, 4))])
            wb = load_workbook(path)
            try:
                pivot_sheet = wb['Reaction pivot']
                self.assertEqual(pivot_sheet['B2'].value, -10)
                self.assertEqual(pivot_sheet['C2'].value, 12)
                self.assertEqual(len(pivot_sheet._pivots[0].dataFields), 12)
                self.assertEqual(pivot_sheet._pivots[0].cache.cacheSource.worksheetSource.ref, 'A1:H3')
                self.assertEqual(pivot_sheet._pivots[0].cache.records.count, 2)
                self.assertIn('Envelopes', wb.sheetnames)
            finally:
                wb.close()


if __name__ == '__main__':
    unittest.main()
