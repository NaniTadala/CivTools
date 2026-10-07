import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import openpyxl
import pymupdf
from civtools.jobs import run_quantities, scan_pdfs, render_pdf


class WorkbookJobTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "mapping.xlsx"
        self.destination = self.root / "result.xlsx"
        workbook = openpyxl.Workbook()
        mapping = workbook.active
        mapping.title = "Drawing_Structure_Mapping"
        mapping.append(["Drawing", "Structure"])
        mapping.append(["a.dxf", "Rack A"])
        mapping.append(["b.dxf", "Rack B"])
        extracted = workbook.create_sheet("DXF_Extracted_Data")
        extracted.append(["File Name", "Extraction Status"])
        for name, quantity in (("a.dxf", 5), ("b.dxf", 9)):
            extracted.append([name, "Table #1 from Model Space"])
            extracted.append(["Description", "Quantity"])
            extracted.append(["M25", quantity])
        workbook.save(self.source)
        workbook.close()
        self.original = self.source.read_bytes()

    def tearDown(self):
        self.temp.cleanup()

    def test_extraction_preserves_other_drawings_and_does_not_log_false_removals(self):
        rows = [["a.dxf", "Table #1 from Model Space", "filename_header"],
                ["Description", "Quantity", "is_header"], ["M25", 7]]
        with patch("civtools.core.quantities.extract_table_data_from_dxf", return_value=(rows, True)):
            result = run_quantities(str(self.source), str(self.destination), ["a.dxf"], False, lambda *args: None)
        self.assertIn("completed", result)
        self.assertEqual(self.source.read_bytes(), self.original)
        workbook = openpyxl.load_workbook(self.destination)
        try:
            extracted = list(workbook["DXF_Extracted_Data"].values)
            self.assertIn(("b.dxf", "Table #1 from Model Space"), extracted)
            self.assertIn(("M25", 9), extracted)
            self.assertIn(("M25", 7), extracted)
            log_values = list(workbook["Update_Log"].values)
            self.assertFalse(any(row[0] == "b.dxf" for row in log_values))
        finally:
            workbook.close()

    def test_consolidation_produces_totals_and_drawing_sheet(self):
        run_quantities(str(self.source), str(self.destination), [], True, lambda *args: None)
        workbook = openpyxl.load_workbook(self.destination)
        try:
            self.assertIn("Consolidated_Materials", workbook.sheetnames)
            self.assertIn("Materials_by_Drawing", workbook.sheetnames)
            summary = list(workbook["Consolidated_Materials"].values)
            self.assertIn(("M25", 5, 9), summary)
        finally:
            workbook.close()
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_failure_keeps_source_and_previous_output_intact(self):
        self.destination.write_bytes(b"previous output")
        with patch("civtools.core.quantities.extract_table_data_from_dxf", side_effect=RuntimeError("Bad drawing")):
            with self.assertRaisesRegex(RuntimeError, "Bad drawing"):
                run_quantities(str(self.source), str(self.destination), ["a.dxf"], False, lambda *args: None)
        self.assertEqual(self.destination.read_bytes(), b"previous output")
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertFalse(list(self.root.glob(".civtools-*")))


class DocumentJobTests(unittest.TestCase):
    def test_metadata_cache_skips_pdf_reopen_and_invalidates_changed_files(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / "standard.pdf"
            cache = root / "cache.json"
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.save(path)
            documents = scan_pdfs(root, lambda *args: None, cache)
            self.assertEqual(documents[0][3], 1)
            with patch("pymupdf.open", side_effect=AssertionError("Cache should skip reopening")):
                self.assertEqual(scan_pdfs(root, lambda *args: None, cache), documents)
            self.assertTrue(render_pdf(str(path), lambda *args: None).startswith(b"\x89PNG"))
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.new_page()
                pdf.save(root / "replacement.pdf")
            (root / "replacement.pdf").replace(path)
            self.assertEqual(scan_pdfs(root, lambda *args: None, cache)[0][3], 2)


if __name__ == "__main__":
    unittest.main()
