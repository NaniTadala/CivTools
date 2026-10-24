import os
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0,str(Path(__file__).resolve().parents[1] / 'src'))
from civtools.core.report_models import choose_layout, discover, fingerprint, resolve_settings, versioned_path
from civtools.core.report_service import pdf_task, publish


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.config = tempfile.TemporaryDirectory()
        self.previous_config = os.environ.get('CIVTOOLS_CONFIG_DIR')
        os.environ['CIVTOOLS_CONFIG_DIR'] = self.config.name

    def tearDown(self):
        if self.previous_config is None:
            os.environ.pop('CIVTOOLS_CONFIG_DIR',None)
        else:
            os.environ['CIVTOOLS_CONFIG_DIR'] = self.previous_config
        self.config.cleanup()

    def test_layout_and_override_precedence(self):
        settings = resolve_settings({'minimum_font':9}, {'paper':'A3'}, {'paper':'A4'})
        self.assertEqual(settings['paper'],'A4')
        self.assertEqual(choose_layout(700,11,'Portrait',settings)['orientation'],'Landscape')
        self.assertEqual(choose_layout(350,11,'Portrait',settings)['orientation'],'Portrait')
        self.assertTrue(choose_layout(2000,11,'Portrait',settings)['needs_review'])
        settings['allow_a3'] = True
        self.assertEqual(choose_layout(1100,11,'Portrait',settings)['paper'],'A3')

    def test_discovery_excludes_output_locks_and_duplicates(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root/'nested').mkdir()
            (root/'output').mkdir()
            for relative in ['a.xlsx','~$a.xlsx','nested/b.XLS','output/old.xlsx','notes.txt']:
                (root/relative).touch()
            found, ignored = discover([root,root/'a.xlsx'],excluded=[root/'output'])
            self.assertEqual(len(found),2)
            self.assertEqual(len(ignored),1)
            self.assertEqual(len(discover([root],recursive=False,excluded=[root/'output'])[0]),1)

    def test_font_inspection_uses_uniform_ranges_and_preserves_populated_counts(self):
        from types import SimpleNamespace
        from unittest.mock import Mock
        from civtools.core.report_excel_worker import sample_body_font
        bounds=Mock()
        bounds.Rows.Count=35
        bounds.Columns.Count=25
        sample=Mock()
        sheet=Mock()
        sheet.Range.return_value=sample
        sample.Value2=tuple(tuple(1 for _ in range(25)) for _ in range(35))
        sample.Font.Size=11
        self.assertEqual(sample_body_font(sheet,bounds),11)
        sample.Cells.assert_not_called()
        sample.Rows.assert_not_called()
        sample.Value2=(('Heading',None,None),('Load',42,None),('Result',10,None))
        sample.Font.Size=None
        sample.Rows.side_effect=lambda row:SimpleNamespace(Font=SimpleNamespace(Size=None if row==1 else 11))
        sample.Cells.return_value.Font.Size=18
        self.assertEqual(sample_body_font(sheet,bounds),11)
        sample.Cells.assert_called_once_with(1,1)

    def test_prepared_cache_requires_current_bytes_settings_order_and_files(self):
        from civtools.core.report_models import ENGINE_VERSION
        from civtools.core.report_service import can_reuse_prepared
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            source=root/'source.xlsx'
            source.write_bytes(b'original')
            artifact=root/'prepared.pdf'
            artifact.touch()
            settings={name:resolve_settings() for name in ('A','B')}
            prepared=dict(engine_version=ENGINE_VERSION,fingerprint=fingerprint(source),artifact=str(artifact),settings=settings,
                rendered=[dict(name=name,path=str(artifact)) for name in settings])
            record=dict(source=str(source))
            self.assertTrue(can_reuse_prepared(record,prepared,settings))
            self.assertFalse(can_reuse_prepared(record,prepared,dict(reversed(list(settings.items())))))
            changed=dict(settings,A=resolve_settings({'orientation':'Landscape'}))
            self.assertFalse(can_reuse_prepared(record,prepared,changed))
            source.write_bytes(b'changed')
            self.assertFalse(can_reuse_prepared(record,prepared,settings))
            source.write_bytes(b'original')
            artifact.unlink()
            self.assertFalse(can_reuse_prepared(record,prepared,settings))

    def test_pdf_assembly_preview_and_exact_publication(self):
        import pymupdf as fitz
        class Pdf:
            def call(self, operation, report, action, payload):
                return pdf_task(action,payload)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            entries=[]
            for index,(w,h) in enumerate([(595,842),(842,595)]):
                path = root/f'{index}.pdf'
                with fitz.open() as doc:
                    page = doc.new_page(width=w,height=h)
                    page.insert_text((40,60),f'Calculation {index}',fontsize=11)
                    doc.save(path)
                entries.append(dict(name=f'Sheet {index}',path=str(path)))
            assembled = root/'prepared.pdf'
            mapping = pdf_task('assemble',dict(entries=entries,output=str(assembled),contents=True,cover='Design calculations',numbering=True))
            self.assertEqual(mapping[0]['first'],3)
            preview=pdf_task('preview',dict(path=str(assembled),page=2))
            self.assertTrue(preview['image'].startswith(b'\x89PNG'))
            self.assertEqual(pdf_task('page_sizes',dict(path=str(assembled)))[2:],[(595,842),(842,595)])
            with fitz.open(assembled) as doc:
                self.assertGreater(doc[3].rect.width,doc[3].rect.height)
                self.assertEqual(len(doc.get_toc()),2)
                self.assertEqual(len(doc[1].get_links()),2)
            source=root/'source.xlsx'
            source.write_bytes(b'original')
            item=dict(source=str(source),relative='nested/report.xlsx',fingerprint=fingerprint(source),artifact=str(assembled))
            first=publish([item],root/'out',Pdf(),lambda *args:None)
            second=publish([item],root/'out',Pdf(),lambda *args:None)
            self.assertNotEqual(first,second)
            self.assertEqual(fingerprint(first[0]),fingerprint(assembled))
            source.write_bytes(b'changed')
            with self.assertRaisesRegex(ValueError,'changed'):
                publish([item],root/'out',Pdf(),lambda *args:None)

    def test_sheet_cache_renders_only_changed_layouts_and_rejects_changed_sources(self):
        import pymupdf as fitz
        from civtools.core.report_service import prepare_report
        class Pdf:
            def call(self,operation,report,action,payload):
                return pdf_task(action,payload)
        class Excel:
            def __init__(self):
                self.calls=[]
            def call(self,request):
                self.calls.append(request)
                rendered=[]
                for name,layout in request['selected'].items():
                    path=Path(request['staging'])/(name+'.pdf')
                    w,h=(842,595) if layout['orientation']=='Landscape' else (595,842)
                    with fitz.open() as doc:
                        doc.new_page(width=w,height=h).insert_text((40,60),f'{name} version {len(self.calls)}')
                        doc.save(path)
                    rendered.append(dict(name=name,path=str(path),issues=[],plan={'reason':'Fixture','orientation':layout['orientation']}))
                return dict(source=request['source'],fingerprint=fingerprint(request['source']),issues=[],rendered=rendered)
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            source=root/'fixture.xlsx'
            source.write_bytes(b'original saved workbook')
            record=dict(source=str(source),relative=source.name)
            settings={name:resolve_settings({'orientation':'Portrait'}) for name in ('A','B','C')}
            excel=Excel()
            def prepare(label,layout,cached=None,force=False):
                return prepare_report(record,layout,resolve_settings(),root/label,excel,Pdf(),lambda *args:None,cached=cached,force=force)
            first=prepare('first',settings)
            old_b=first['rendered'][1]['path']
            old_b_hash=fingerprint(old_b)
            edited=dict(A=resolve_settings({'orientation':'Landscape'}),B=settings['B'])
            second=prepare('edited',edited,first)
            self.assertEqual(list(excel.calls[-1]['selected']),['A'])
            self.assertEqual(second['rendered'][1]['path'],old_b)
            self.assertEqual(fingerprint(old_b),old_b_hash)
            with fitz.open(second['artifact']) as doc:
                self.assertGreater(doc[0].rect.width,doc[0].rect.height)
                self.assertLess(doc[1].rect.width,doc[1].rect.height)
                self.assertIn('B version 1',doc[1].get_text())
            # Excluded, unchanged sheets remain available after a layout edit.
            restored=prepare('restored',dict(edited,C=settings['C']),second)
            self.assertEqual(len(excel.calls),2)
            self.assertEqual([s['name'] for s in restored['rendered']],['A','B','C'])
            self.assertEqual(prepare('unchanged',dict(edited,C=settings['C']),restored)['artifact'],restored['artifact'])
            source.write_bytes(b'changed saved workbook')
            changed=prepare('changed-source',dict(edited,C=settings['C']),restored)
            self.assertEqual(list(excel.calls[-1]['selected']),['A','B','C'])
            Path(changed['rendered'][1]['path']).unlink()
            repaired=prepare('missing-page',dict(edited,C=settings['C']),changed)
            self.assertEqual(list(excel.calls[-1]['selected']),['B'])
            prepare('forced',dict(edited,C=settings['C']),repaired,force=True)
            self.assertEqual(list(excel.calls[-1]['selected']),['A','B','C'])
            recalculated=prepare('recalculated',dict(edited,C=resolve_settings({'recalculate':True})),repaired)
            self.assertEqual(list(excel.calls[-1]['selected']),['A','B','C'])
            prepare('saved-values',dict(edited,C=settings['C']),recalculated)
            self.assertEqual(list(excel.calls[-1]['selected']),['A','B','C'])

    def test_ui_overrides_invalidation_and_session_restore(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import QSettings, QThreadPool
        from civtools.qt_report_publisher import ReportPublisherPage
        app=QApplication.instance() or QApplication([])
        class Host:
            def update_status(self):
                pass
        with tempfile.TemporaryDirectory() as temp:
            host=Host()
            host.settings=QSettings(str(Path(temp)/'settings.ini'),QSettings.Format.IniFormat)
            host.pool=QThreadPool()
            page=ReportPublisherPage(host)
            page.restore_state('reports')
            page.records={'source.xlsx':dict(source='source.xlsx',relative='source.xlsx',status='Ready',overrides={},included=True,
                sheets=[dict(name='Calculations',included=True,visible=True,empty=False,overrides={})])}
            page.rebuild()
            page.tree.setCurrentItem(page.tree.topLevelItem(0).child(0))
            page.orientation.setCurrentText('Landscape')
            page.apply_selected()
            self.assertEqual(page.effective('source.xlsx','Calculations')['orientation'],'Landscape')
            self.assertEqual(page.records['source.xlsx']['status'],'Not prepared')
            page.batch['paper']='A3'
            self.assertEqual(page.effective('source.xlsx','Calculations')['paper'],'A3')
            page.save_state()
            restored=ReportPublisherPage(host)
            restored.restore_state('reports')
            self.assertEqual(restored.effective('source.xlsx','Calculations')['orientation'],'Landscape')
            page.deleteLater()
            restored.deleteLater()
            app.processEvents()

    def test_settings_cancel_and_export_options_preserve_prepared_report(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        from PySide6.QtWidgets import QApplication, QPushButton
        from PySide6.QtCore import QSettings, QThreadPool, QTimer
        from civtools.qt_report_publisher import ReportPublisherPage
        from civtools.qt_report_workspace import show_settings
        app=QApplication.instance() or QApplication([])
        class Host:
            def update_status(self):
                pass
        with tempfile.TemporaryDirectory() as temp:
            host=Host()
            host.settings=QSettings(str(Path(temp)/'settings.ini'),QSettings.Format.IniFormat)
            host.pool=QThreadPool()
            page=ReportPublisherPage(host)
            page.restore_state('reports')
            page.records={'source.xlsx':dict(source='source.xlsx',relative='source.xlsx',status='Ready',overrides={},included=True,
                sheets=[dict(name='Calculations',included=True,visible=True,empty=False,overrides={})])}
            page.rebuild()
            page.prepared={'source.xlsx':{'artifact':'unchanged.pdf'}}
            def cancel():
                page.orientation.setCurrentText('Landscape')
                page.combined.setChecked(True)
                page.settings_dialog.reject()
            QTimer.singleShot(0,cancel)
            show_settings(page)
            self.assertFalse(page.combined.isChecked())
            self.assertEqual(page.batch['orientation'],'Auto')
            self.assertIn('source.xlsx',page.prepared)
            def apply_export_only():
                page.combined.setChecked(True)
                next(b for b in page.settings_dialog.findChildren(QPushButton) if b.text()=='Apply').click()
            QTimer.singleShot(0,apply_export_only)
            show_settings(page)
            self.assertTrue(page.combined.isChecked())
            self.assertIn('source.xlsx',page.prepared)
            self.assertTrue(page.stop_button.isHidden())
            self.assertTrue(page.progress.isHidden())
            self.assertTrue(page.settings_dialog.isHidden())
            page.combined_timer.stop()
            page.preview_resize_timer.stop()
            page.thumbnail_timer.stop()
            page.deleteLater()
            app.processEvents()

    def test_completed_reports_are_reviewable_and_saved_before_batch_finishes(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        import json
        import threading
        import time
        import pymupdf as fitz
        from unittest.mock import patch
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import QSettings,QThreadPool,Qt
        from civtools.qt_report_publisher import ReportPublisherPage
        app=QApplication.instance() or QApplication([])
        class Host:
            def update_status(self):
                pass
        class Pdf:
            lock=threading.Lock()
            def call(self,operation,report,action,payload):
                with self.lock:
                    return pdf_task(action,payload)
        second_started=threading.Event()
        last_started=threading.Event()
        release_second=threading.Event()
        release_last=threading.Event()
        class Excel:
            def __init__(self,report,stop_now):
                self.stop_now=stop_now
            def __enter__(self):
                return self
            def __exit__(self,*args):
                pass
            def call(self,request):
                name=Path(request['source']).stem
                if name in ('B','C'):
                    started,release=(second_started,release_second) if name=='B' else (last_started,release_last)
                    started.set()
                    if not release.wait(15):
                        raise TimeoutError('Fixture gate timed out')
                    if self.stop_now.is_set():
                        raise InterruptedError('Stopped fixture')
                rendered=[]
                for sheet in request['selected']:
                    path=Path(request['staging'])/(sheet+'.pdf')
                    with fitz.open() as doc:
                        for number in (1,2):
                            doc.new_page().insert_text((40,60),f'{name} {sheet} page {number}')
                        doc.save(path)
                    rendered.append(dict(name=sheet,path=str(path),issues=[],plan={'reason':'Fixture','orientation':'Portrait'}))
                return dict(source=request['source'],fingerprint=fingerprint(request['source']),issues=[],rendered=rendered)
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            host=Host()
            host.settings=QSettings(str(root/'state.ini'),QSettings.Format.IniFormat)
            host.pool=QThreadPool()
            host.pool.setMaxThreadCount(3)
            host.pdf=Pdf()
            page=ReportPublisherPage(host)
            page.restore_state('reports')
            for name in ('A','B','C'):
                source=str(root/(name+'.xlsx'))
                Path(source).write_bytes(name.encode())
                page.records[source]=dict(source=source,relative=name+'.xlsx',status='Not prepared',included=True,overrides={},
                    sheets=[dict(name='Calculation',included=True,visible=True,empty=False,overrides={})])
            page.rebuild()
            page.resize(1000,700)
            page.show()
            page.combined_toggle.setChecked(True)
            def wait(predicate):
                deadline=time.monotonic()+10
                while time.monotonic()<deadline:
                    app.processEvents()
                    if predicate():
                        return
                    time.sleep(.01)
                self.fail(page.feedback.toolTip())
            a,b,c=list(page.records)
            try:
                with patch('civtools.qt_report_publisher.ExcelSession',Excel):
                    page.prepare()
                    wait(lambda:second_started.is_set() and a in page.prepared and page.preview.pixmap() is not None)
                    self.assertTrue(page.busy)
                    self.assertTrue(page.tree.isEnabled())
                    self.assertFalse(page.tree.topLevelItem(0).flags() & Qt.ItemFlag.ItemIsUserCheckable)
                    self.assertFalse(page.export_button.isEnabled())
                    self.assertIsNone(page.combined_preview)
                    self.assertIn('individual preview',page.preview_name.full_text)
                    manifest=Path(str(host.settings.value('reports/last_manifest')))
                    self.assertEqual(list(json.loads(manifest.read_text())['prepared']),[a])
                    page.tree.setCurrentItem(page.tree.topLevelItem(0).child(0))
                    wait(lambda:page.preview_path==page.prepared[a]['rendered'][0]['path'] and page.preview.pixmap() is not None)
                    page.page.setValue(2)
                    wait(lambda:page.preview.pixmap() is not None)
                    path=page.preview_path
                    page.tree.topLevelItem(1).setExpanded(False)
                    release_second.set()
                    wait(lambda:last_started.is_set() and b in page.prepared)
                    self.assertEqual(page.preview_path,path)
                    self.assertEqual(page.page.value(),2)
                    self.assertFalse(page.tree.topLevelItem(1).isExpanded())
                    self.assertEqual(page.key(page.tree.currentItem()),(a,'Calculation'))
                    # A later cancellation must not lose completed reports.
                    page.stop_now.set()
                    release_last.set()
                    wait(lambda:not page.busy and not page.jobs and not page.combined_timer.isActive())
                self.assertEqual(set(page.prepared),{a,b})
                self.assertEqual(page.records[c]['status'],'Failed')
                self.assertEqual(page.preview_path,path)
                self.assertEqual(page.page.value(),2)
                self.assertTrue(page.tree.topLevelItem(0).flags() & Qt.ItemFlag.ItemIsUserCheckable)
                page.output.setText(str(root/'Published'))
                page.export()
                self.assertIn('incomplete',page.feedback.full_text)
                restored=ReportPublisherPage(host)
                restored.restore_state('reports')
                self.assertEqual(set(restored.prepared),{a,b})
                # Retain an edited workbook's sheet cache across a restart.
                page.invalidate(a)
                page.save_state()
                cached=ReportPublisherPage(host)
                cached.restore_state('reports')
                self.assertNotIn(a,cached.prepared)
                self.assertIn(a,cached.sheet_cache)
                for other in (restored,cached):
                    other.combined_timer.stop()
                    other.preview_resize_timer.stop()
                    other.thumbnail_timer.stop()
                    other.preview.timer.stop()
                    other.deleteLater()
            finally:
                release_second.set()
                release_last.set()
                wait(lambda:not page.busy and not page.jobs)
                host.pool.waitForDone()
                page.close()
                page.deleteLater()
                app.processEvents()

    def test_remove_all_restores_start_screen_and_discards_late_combined_preview(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        import json
        import threading
        import time
        import pymupdf as fitz
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import QSettings,QThreadPool
        from civtools.qt_report_publisher import ReportPublisherPage
        app=QApplication.instance() or QApplication([])
        started=threading.Event()
        release=threading.Event()
        class Host:
            def update_status(self):
                pass
        class Pdf:
            def call(self,operation,report,action,payload):
                if action=='assemble':
                    started.set()
                    if not release.wait(10):
                        raise TimeoutError('Fixture gate timed out')
                return pdf_task(action,payload)
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            host=Host()
            host.settings=QSettings(str(root/'settings.ini'),QSettings.Format.IniFormat)
            host.pool=QThreadPool()
            host.pool.setMaxThreadCount(3)
            host.pdf=Pdf()
            page=ReportPublisherPage(host)
            page.restore_state('reports')
            artifact=root/'issued.pdf'
            with fitz.open() as doc:
                doc.new_page().insert_text((40,60),'Issued calculation')
                doc.save(artifact)
            pdf_hash=fingerprint(artifact)
            for name in ('A','B'):
                source=str(root/(name+'.xlsx'))
                Path(source).write_bytes(name.encode())
                page.records[source]=dict(source=source,relative=name+'.xlsx',status='Ready',included=True,overrides={},
                    sheets=[dict(name='Calculation',included=True,visible=True,empty=False,overrides={})])
                page.prepared[source]=dict(source=source,relative=name+'.xlsx',fingerprint=fingerprint(source),artifact=str(artifact),issues=[],
                    rendered=[dict(name='Calculation',path=str(artifact),issues=[],plan={'reason':'Fixture'})])
            originals={source:fingerprint(source) for source in page.records}
            page.sheet_cache=dict(page.prepared)
            page.passwords=dict.fromkeys(page.records,'fixture password')
            page.history=[dict(batch=page.batch,records=dict(page.records))]
            page.rebuild()
            page.resize(1000,700)
            page.show()
            def wait(predicate):
                deadline=time.monotonic()+10
                while time.monotonic()<deadline:
                    app.processEvents()
                    if predicate():
                        return
                    time.sleep(.01)
                self.fail(page.feedback.toolTip())
            try:
                page.save_state()
                self.assertTrue(page.clear_button.isVisible())
                page.combined_toggle.setChecked(True)
                wait(started.is_set)
                self.assertFalse(page.busy)
                page.search.setText('old search')
                page.filter.setCurrentText('Failed')
                page.clear_button.click()
                self.assertFalse(page.records)
                self.assertFalse(page.prepared)
                self.assertFalse(page.sheet_cache)
                self.assertFalse(page.passwords)
                self.assertFalse(page.history)
                self.assertEqual(page.preview_stack.currentIndex(),0)
                self.assertEqual(page.empty_title.text(),'Drop your Excel reports')
                self.assertEqual(page.preview_path,'')
                self.assertFalse(page.preview_nav.isVisible())
                self.assertFalse(page.report_sidebar.isVisible())
                self.assertFalse(page.clear_button.isVisible())
                self.assertFalse(page.prepare_button.isEnabled())
                self.assertFalse(page.export_button.isEnabled())
                self.assertEqual(page.search.text(),'')
                self.assertEqual(page.filter.currentText(),'All reports')
                self.assertEqual(page.thumbnails.count(),0)
                manifest=Path(str(host.settings.value('reports/last_manifest')))
                self.assertEqual(json.loads(manifest.read_text())['records'],{})
                release.set()
                wait(lambda:not page.jobs)
                self.assertEqual(page.preview_stack.currentIndex(),0)
                self.assertEqual(page.preview_path,'')
                self.assertIsNone(page.combined_preview)
                self.assertEqual(page.empty_title.text(),'Drop your Excel reports')
                restored=ReportPublisherPage(host)
                restored.restore_state('reports')
                self.assertEqual(restored.records,{})
                self.assertEqual(restored.preview_stack.currentIndex(),0)
                restored.deleteLater()
                self.assertEqual(fingerprint(artifact),pdf_hash)
                self.assertTrue(all(fingerprint(source)==digest for source,digest in originals.items()))
                # The next upload must be visible rather than retaining filters.
                source=next(iter(originals))
                page.job_event('report_item',dict(source=source,relative='A.xlsx',status='Not prepared',included=True,overrides={},sheets=[]))
                self.assertTrue(page.report_sidebar.isVisible())
                self.assertTrue(page.clear_button.isVisible())
                self.assertFalse(page.tree.topLevelItem(0).isHidden())
                page.tree.setCurrentItem(page.tree.topLevelItem(0))
                page.remove_selected()
                self.assertEqual(page.empty_title.text(),'Drop your Excel reports')
                self.assertEqual(page.preview_stack.currentIndex(),0)
            finally:
                release.set()
                wait(lambda:not page.jobs)
                page.close()
                page.deleteLater()
                app.processEvents()

    def test_removing_last_report_clears_individual_preview_and_saved_session(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import QSettings,QThreadPool
        from civtools.qt_report_publisher import ReportPublisherPage
        app=QApplication.instance() or QApplication([])
        class Host:
            def update_status(self):
                pass
        with tempfile.TemporaryDirectory() as temp:
            host=Host()
            host.settings=QSettings(str(Path(temp)/'settings.ini'),QSettings.Format.IniFormat)
            host.pool=QThreadPool()
            page=ReportPublisherPage(host)
            page.restore_state('reports')
            page.records={'source.xlsx':dict(source='source.xlsx',relative='source.xlsx',status='Not prepared',included=True,overrides={},
                sheets=[dict(name='Calculation',included=True,visible=True,empty=False,overrides={})])}
            page.rebuild()
            page.save_state()
            page.preview_path='previous.pdf'
            page.thumbnail_path='previous.pdf'
            page.tree.setCurrentItem(page.tree.topLevelItem(0).child(0))
            page.set_busy(True)
            page.remove_all()
            self.assertEqual(len(page.records),1)
            page.set_busy(False)
            page.remove_selected()
            self.assertEqual(page.records,{})
            self.assertEqual(page.preview_path,'')
            self.assertEqual(page.thumbnail_path,'')
            self.assertEqual(page.preview_stack.currentIndex(),0)
            self.assertEqual(page.empty_title.text(),'Drop your Excel reports')
            restored=ReportPublisherPage(host)
            restored.restore_state('reports')
            self.assertEqual(restored.records,{})
            page.undo_layout()
            self.assertEqual(page.records,{})
            page.deleteLater()
            restored.deleteLater()
            app.processEvents()

    def test_preview_page_field_tracks_scroll_and_accepts_committed_typing(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        import time
        import pymupdf as fitz
        from PySide6.QtWidgets import QApplication, QPushButton
        from PySide6.QtCore import QSettings, QThreadPool, Qt
        from PySide6.QtGui import QFontDatabase
        from PySide6.QtTest import QTest
        from civtools.qt_report_publisher import ReportPublisherPage
        from civtools.qt_style import style_for_scale
        app=QApplication.instance() or QApplication([])
        old_style=app.styleSheet()
        QFontDatabase.addApplicationFont(str(Path(__file__).resolve().parents[1]/'fonts'/'Poppins'/'Poppins-Regular.ttf'))
        class Host:
            def update_status(self):
                pass
        class Pdf:
            def call(self,operation,report,action,payload):
                return pdf_task(action,payload)
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            path=root/'twelve-pages.pdf'
            with fitz.open() as doc:
                for i in range(12):
                    w,h=(595,842) if i%2==0 else (842,595)
                    doc.new_page(width=w,height=h).insert_text((40,60),f'Page {i+1}')
                doc.save(path)
            host=Host()
            host.settings=QSettings(str(root/'state.ini'),QSettings.Format.IniFormat)
            host.pool=QThreadPool()
            host.pool.setMaxThreadCount(1)
            host.pdf=Pdf()
            page=ReportPublisherPage(host)
            page.resize(1000,700)
            page.show()
            def wait():
                deadline=time.monotonic()+10
                while time.monotonic()<deadline:
                    app.processEvents()
                    time.sleep(.01)
                    if not page.jobs and not page.preview.timer.isActive() and not page.preview_resize_timer.isActive() and not page.thumbnail_timer.isActive():
                        return
                self.fail(page.feedback.toolTip())
            try:
                for scale in (100,130):
                    with self.subTest(scale=scale):
                        app.setStyleSheet(style_for_scale(scale))
                        page.preview_path=str(path)
                        page.render_preview()
                        wait()
                        scroll=page.preview_scroll.verticalScrollBar()
                        scroll.setValue(scroll.maximum())
                        # The counter must update before rendering timers run.
                        self.assertEqual(page.page.value(),12)
                        self.assertEqual(page.page.text(),'12')
                        previous=next(b for b in page.preview_nav.findChildren(QPushButton) if b.accessibleName()=='Previous page')
                        previous.click()
                        self.assertEqual(page.page.value(),11)
                        self.assertEqual(scroll.value(),page.preview.rects[10].top()-18)
                        page.page.setValue(6)
                        wait()
                        editor=page.page
                        page.page.setFocus()
                        page.page.selectAll()
                        page.preview.update_current_page()
                        self.assertEqual(editor.selectedText(),'6')
                        QTest.keyClicks(editor,'1')
                        self.assertEqual(page.page.value(),6)
                        # Scrolling and image completion must preserve a draft.
                        scroll.setValue(page.preview.rects[7].top()-18)
                        page.preview.load_visible()
                        QTest.qWait(100)
                        self.assertEqual(editor.text(),'1')
                        QTest.keyClicks(editor,'2')
                        self.assertEqual(editor.text(),'12')
                        QTest.keyClick(editor,Qt.Key.Key_Return)
                        self.assertEqual(page.page.value(),12)
                        self.assertEqual(scroll.value(),page.preview.rects[11].top()-18)
                        page.page.selectAll()
                        QTest.keyClicks(editor,'3')
                        page.zoom.setFocus()
                        self.assertEqual(page.page.value(),3)
                        self.assertEqual(scroll.value(),page.preview.rects[2].top()-18)
                        self.assertGreater(editor.contentsRect().width(),editor.fontMetrics().horizontalAdvance('12')+4)
                        page.page.setFocus()
                        page.page.selectAll()
                        QTest.keyClicks(editor,'999')
                        QTest.keyClick(editor,Qt.Key.Key_Return)
                        self.assertEqual(page.page.value(),12)
                        self.assertEqual(editor.text(),'12')
                        page.page.selectAll()
                        QTest.keyClick(editor,Qt.Key.Key_Backspace)
                        QTest.keyClick(editor,Qt.Key.Key_Return)
                        self.assertEqual(editor.text(),'12')
                        page.page.selectAll()
                        QTest.keyClicks(editor,'4')
                        QTest.keyClick(editor,Qt.Key.Key_Escape)
                        self.assertEqual(page.page.value(),12)
                        self.assertEqual(editor.text(),'12')
                        wait()
            finally:
                wait()
                page.close()
                page.deleteLater()
                app.processEvents()
                app.setStyleSheet(old_style)

    def test_combined_preview_tracks_report_sheet_order_and_exclusions(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        import time
        import threading
        import pymupdf as fitz
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import QSettings,QThreadPool,Qt,QPointF
        from unittest.mock import patch
        from civtools.qt_report_publisher import ReportPublisherPage
        app=QApplication.instance() or QApplication([])
        class Host:
            def update_status(self):
                pass
        class Pdf:
            lock=threading.Lock()
            def call(self,operation,report,action,payload):
                with self.lock:
                    return pdf_task(action,payload)
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            host=Host()
            host.settings=QSettings(str(root/'settings.ini'),QSettings.Format.IniFormat)
            host.pool=QThreadPool()
            host.pool.setMaxThreadCount(1)
            host.pdf=Pdf()
            page=ReportPublisherPage(host)
            page.restore_state('reports')
            page.resize(1000,700)
            page.show()
            for name in ('A','B'):
                source=str(root/(name+'.xlsx'))
                Path(source).write_bytes(name.encode())
                sheets=[]
                rendered=[]
                for number in (1,2):
                    sheet=name+str(number)
                    path=root/(sheet+'.pdf')
                    with fitz.open() as doc:
                        doc.new_page().insert_text((40,60),sheet)
                        doc.save(path)
                    sheets.append(dict(name=sheet,included=True,visible=True,empty=False,overrides={}))
                    rendered.append(dict(name=sheet,path=str(path),issues=[],plan={'reason':'Test report'}))
                artifact=root/(name+'.pdf')
                mapping=pdf_task('assemble',dict(entries=rendered,output=str(artifact)))
                page.records[source]=dict(source=source,relative=name+'.xlsx',status='Ready',included=True,overrides={},sheets=sheets)
                page.prepared[source]=dict(source=source,relative=name+'.xlsx',fingerprint=fingerprint(source),artifact=str(artifact),rendered=rendered,issues=[],page_map=mapping)
            page.rebuild()
            page.output.setText(str(root/'Published'))
            def wait():
                deadline=time.monotonic()+15
                while time.monotonic()<deadline:
                    app.processEvents()
                    time.sleep(.01)
                    if not page.busy and not page.jobs and not page.combined_timer.isActive() and not page.preview_resize_timer.isActive() and not page.thumbnail_timer.isActive() and not page.preview.timer.isActive():
                        return
                self.fail(page.feedback.toolTip())
            def contents():
                self.assertIsNotNone(page.combined_preview)
                self.assertEqual(page.preview_path,page.combined_preview[1])
                with fitz.open(page.preview_path) as pdf:
                    return [p.get_text().strip() for p in pdf]
            page.combined_toggle.setChecked(True)
            wait()
            self.assertEqual(contents(),['A1','A2','B1','B2'])
            self.assertEqual(len(page.preview.rects),4)
            scroll=page.preview_scroll.verticalScrollBar()
            scroll.setValue(scroll.maximum())
            wait()
            self.assertEqual(page.page.value(),4)
            self.assertIn(3,page.preview.images)
            self.assertLessEqual(len(page.preview.images),3)
            page.page.setValue(2)
            wait()
            self.assertEqual(scroll.value(),page.preview.rects[1].top()-18)
            self.assertIn(1,page.preview.images)
            page.tree.setCurrentItem(page.tree.topLevelItem(1).child(0))
            wait()
            self.assertEqual(contents(),['A1','A2','B1','B2'])
            self.assertEqual(page.page.value(),2)
            page.cover.setText('Older cover')
            page.preview_combined(automatic=True)
            page.cover.setText('Latest cover')
            wait()
            self.assertEqual(contents()[0],'Latest cover')
            page.cover.clear()
            wait()
            page.tree.setCurrentItem(page.tree.topLevelItem(1))
            class Drop:
                accepted=False
                target=None
                def source(self):
                    return page.tree
                def position(self):
                    return QPointF(page.tree.visualItemRect(self.target or page.tree.topLevelItem(0)).topLeft())
                def setDropAction(self,action):
                    pass
                def accept(self):
                    self.accepted=True
                def ignore(self):
                    self.accepted=False
            drop=Drop()
            with patch('civtools.qt_report_workspace.QDrag') as drag:
                drag.return_value.exec.side_effect=lambda *args:page.tree.dropEvent(drop)
                page.tree.startDrag(Qt.DropAction.MoveAction)
                drag.return_value.exec.assert_called_once()
            self.assertTrue(drop.accepted)
            wait()
            self.assertEqual(page.tree.topLevelItemCount(),2)
            self.assertEqual([page.tree.topLevelItem(i).childCount() for i in range(2)],[2,2])
            self.assertEqual(contents(),['B1','B2','A1','A2'])
            page.tree.setCurrentItem(page.tree.topLevelItem(1).child(1))
            invalid_drop=Drop()
            invalid_drop.target=page.tree.topLevelItem(0).child(0)
            page.tree.dropEvent(invalid_drop)
            self.assertFalse(invalid_drop.accepted)
            sheet_drop=Drop()
            sheet_drop.target=page.tree.topLevelItem(1).child(0)
            with patch('civtools.qt_report_workspace.QDrag') as drag:
                drag.return_value.exec.side_effect=lambda *args:page.tree.dropEvent(sheet_drop)
                page.tree.startDrag(Qt.DropAction.MoveAction)
            self.assertTrue(sheet_drop.accepted)
            wait()
            self.assertEqual(page.tree.topLevelItem(1).childCount(),2)
            self.assertEqual(contents(),['B1','B2','A2','A1'])
            self.assertTrue(all(r['status']=='Ready' for r in page.records.values()))
            page.tree.topLevelItem(0).setCheckState(0,Qt.CheckState.Unchecked)
            wait()
            self.assertEqual(contents(),['A2','A1'])
            top=page.tree.topLevelItem(0)
            self.assertTrue(all(top.child(i).checkState(0)==Qt.CheckState.Unchecked for i in range(top.childCount())))
            self.assertTrue(all(not s['included'] for s in page.records[page.key(top)[0]]['sheets']))
            page.tree.topLevelItem(0).setCheckState(0,Qt.CheckState.Checked)
            wait()
            self.assertEqual(contents(),['B1','B2','A2','A1'])
            page.tree.topLevelItem(1).child(0).setCheckState(0,Qt.CheckState.Unchecked)
            wait()
            self.assertEqual(contents(),['B1','B2','A1'])
            self.assertEqual(page.tree.topLevelItem(1).checkState(0),Qt.CheckState.PartiallyChecked)
            page.tree.topLevelItem(1).child(0).setCheckState(0,Qt.CheckState.Checked)
            wait()
            self.assertEqual(contents(),['B1','B2','A2','A1'])
            # Deselecting every child must untick the workbook and preserve
            # cached pages for a later select-all.
            page.tree.topLevelItem(1).child(0).setCheckState(0,Qt.CheckState.Unchecked)
            wait()
            page.tree.topLevelItem(1).child(1).setCheckState(0,Qt.CheckState.Unchecked)
            wait()
            self.assertEqual(page.tree.topLevelItem(1).checkState(0),Qt.CheckState.Unchecked)
            self.assertEqual(contents(),['B1','B2'])
            page.tree.topLevelItem(1).setCheckState(0,Qt.CheckState.Checked)
            wait()
            self.assertEqual(contents(),['B1','B2','A2','A1'])
            page.cover.setText('Combined design report')
            page.contents.setChecked(True)
            wait()
            self.assertEqual(contents()[0],'Combined design report')
            self.assertEqual(contents()[2:],['B1','B2','A2','A1'])
            preview_hash=fingerprint(page.preview_path)
            page.export()
            wait()
            self.assertEqual(fingerprint(root/'Published'/'Combined report.pdf'),preview_hash)
            page.save_state()
            previous_page=page
            page=ReportPublisherPage(host)
            page.restore_state('reports')
            wait()
            self.assertTrue(page.combined_toggle.isChecked())
            self.assertEqual([Path(s).stem for s in page.records],['B','A'])
            self.assertEqual(contents()[2:],['B1','B2','A2','A1'])
            previous_page.deleteLater()
            page.deleteLater()
            app.processEvents()


@unittest.skipUnless(os.environ.get('CIVTOOLS_TEST_EXCEL')=='1','Set CIVTOOLS_TEST_EXCEL=1 for installed Excel integration')
class ExcelIntegration(unittest.TestCase):
    def test_native_sheet_layout_edit_reuses_unchanged_pdf(self):
        from unittest.mock import patch
        from openpyxl import Workbook
        import pymupdf as fitz
        from civtools.core.report_service import ExcelSession,prepare_report
        class Pdf:
            def call(self,operation,report,action,payload):
                return pdf_task(action,payload)
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            source=root/'two-sheets.xlsx'
            book=Workbook()
            book.active.title='Foundation'
            book.create_sheet('Loads')
            for sheet in book:
                sheet.append(['Design calculation','Value','Unit'])
                for number in range(1,30):
                    sheet.append([f'Design parameter {number}',number,'kN'])
                sheet.column_dimensions['A'].width=28
            book.save(source)
            original=fingerprint(source)
            record=dict(source=str(source),relative=source.name)
            layouts={name:resolve_settings({'orientation':'Portrait'}) for name in ('Foundation','Loads')}
            with ExcelSession(lambda *args:None,timeout=45) as session:
                first=prepare_report(record,layouts,resolve_settings(),root/'first',session,Pdf(),lambda *args:None)
                old_loads=first['rendered'][1]['path']
                old_hash=fingerprint(old_loads)
                edited=dict(layouts,Foundation=resolve_settings({'orientation':'Landscape'}))
                with patch.object(session,'call',wraps=session.call) as calls:
                    second=prepare_report(record,edited,resolve_settings(),root/'edited',session,Pdf(),lambda *args:None,cached=first)
                    self.assertEqual(calls.call_count,1)
                    self.assertEqual(list(calls.call_args.args[0]['selected']),['Foundation'])
                self.assertEqual(second['rendered'][1]['path'],old_loads)
                self.assertEqual(fingerprint(old_loads),old_hash)
                with fitz.open(second['artifact']) as doc:
                    self.assertGreater(doc[0].rect.width,doc[0].rect.height)
                    self.assertLess(doc[1].rect.width,doc[1].rect.height)
                    self.assertTrue(all('Design parameter 29' in page.get_text() for page in doc))
                with self.assertRaisesRegex(RuntimeError,'Source changed'):
                    session.call(dict(operation='inspect',source=str(source),expected_fingerprint='outdated'))
            self.assertEqual(fingerprint(source),original)

    def test_batch_reuses_excel_and_recycles_after_errors_and_stop(self):
        import threading
        from openpyxl import Workbook
        from civtools.core.report_service import ExcelSession
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'source.xlsx'
            book=Workbook()
            book.active.append(['Test load',42])
            book.save(source)
            stop=threading.Event()
            with ExcelSession(lambda *args:None,stop,timeout=30) as session:
                request=dict(operation='inspect',source=str(source))
                first=session.call(request)
                process=session.process
                second=session.call(request)
                self.assertIs(session.process,process)
                self.assertEqual(first,second)
                with self.assertRaises(RuntimeError):
                    session.call(dict(request,source=str(Path(temp)/'missing.xlsx')))
                self.assertIsNone(session.process)
                self.assertEqual(session.call(request),first)
                self.assertIsNot(session.process,process)
                stop.set()
                with self.assertRaises(InterruptedError):
                    session.call(request)
                self.assertIsNone(session.process)
                stop.clear()
                session.max_requests=1
                session.call(request)
                process=session.process
                session.call(request)
                self.assertIsNot(session.process,process)

    def test_workspace_folder_prepare_preview_and_export(self):
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        from PySide6.QtWidgets import QApplication
        from civtools.qt_app import CivToolsApp
        from openpyxl import Workbook
        import time
        app=QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            previous=os.environ.get('CIVTOOLS_CONFIG_DIR')
            os.environ['CIVTOOLS_CONFIG_DIR']=str(root/'config')
            window=None
            try:
                source=root/'Report.xlsx'
                wb=Workbook()
                wb.active.append(['Foundation design report',10,'kN'])
                wb.active.append(['Final design result',20,'kN'])
                wb.active.column_dimensions['A'].width=42
                wb.active.column_dimensions['B'].width=12
                wb.active.column_dimensions['C'].width=12
                wb.save(source)
                window=CivToolsApp()
                window.navigate('reports')
                page=window.pages['reports']
                page.output.setText(str(root/'Published'))
                def wait(predicate):
                    deadline=time.monotonic()+90
                    while not predicate() and time.monotonic()<deadline:
                        app.processEvents()
                        time.sleep(.01)
                    self.assertTrue(predicate(),page.feedback.text())
                page.intake([str(source)])
                wait(lambda:not page.busy)
                self.assertEqual(len(page.records),1)
                page.prepare()
                wait(lambda:not page.busy)
                self.assertEqual(page.records[str(source)]['status'],'Ready',str(page.prepared.get(str(source),{}).get('issues')))
                from unittest.mock import patch
                artifact=page.prepared[str(source)]['artifact']
                with patch('civtools.qt_report_publisher.ExcelSession._start',side_effect=AssertionError('Unchanged reports must use the prepared PDF')):
                    page.prepare()
                    wait(lambda:not page.busy)
                self.assertEqual(page.prepared[str(source)]['artifact'],artifact)
                page.tree.setCurrentItem(page.tree.topLevelItem(0))
                wait(lambda:page.preview.pixmap() is not None and not page.preview.pixmap().isNull())
                page.combined.setChecked(True)
                page.contents.setChecked(True)
                page.cover.setText('Foundation design calculations')
                page.per_sheet.setChecked(True)
                page.preview_combined()
                wait(lambda:not page.busy)
                self.assertIsNotNone(page.combined_preview,page.feedback.text())
                page.export()
                wait(lambda:not page.busy)
                self.assertTrue((root/'Published'/'Combined report.pdf').is_file(),page.feedback.text())
                self.assertEqual(fingerprint(root/'Published'/'Report.pdf'),fingerprint(page.prepared[str(source)]['artifact']))
                self.assertEqual(fingerprint(root/'Published'/'Combined report.pdf'),fingerprint(page.combined_preview[1]))
                wait(lambda:not page.jobs and not page.thumbnail_timer.isActive())
                self.assertFalse(page.thumbnails.item(0).icon().isNull())
            finally:
                if window:
                    window.close()
                if previous is None:
                    os.environ.pop('CIVTOOLS_CONFIG_DIR',None)
                else:
                    os.environ['CIVTOOLS_CONFIG_DIR']=previous

    def test_native_formats_images_formulas_and_section_breaks(self):
        import pythoncom
        import win32com.client
        import pymupdf as fitz
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image
        from PIL import Image as Raster
        from civtools.core.report_service import excel_call
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            image=root/'diagram.png'
            Raster.new('RGB',(100,50),'blue').save(image)
            source=root/'Template.xlsx'
            wb=Workbook()
            sheet=wb.active
            sheet.title='Calculation'
            for i in range(1,55):
                sheet.append([f'Parameter {i}',i,'kN'])
            sheet['B4']='=SUM(B2:B3)'
            sheet.merge_cells('A1:C1')
            sheet['A1']='Foundation design calculation'
            sheet['A1'].font=__import__('openpyxl').styles.Font(bold=True,size=12)
            sheet.column_dimensions['A'].width=32
            sheet.add_image(Image(str(image)),'E2')
            wb.save(source)
            pythoncom.CoInitialize()
            excel=win32com.client.DispatchEx('Excel.Application')
            excel.Visible=False
            excel.DisplayAlerts=False
            book=None
            paths=[]
            try:
                book=excel.Workbooks.Open(str(source))
                excel.CalculateFull()
                for suffix,format_id in [('xlsx',51),('xlsm',52),('xlsb',50),('xls',56)]:
                    path=root/f'Report.{suffix}'
                    book.SaveAs(str(path),FileFormat=format_id)
                    paths.append(path)
            finally:
                if book is not None:
                    book.Close(False)
                excel.Quit()
                pythoncom.CoUninitialize()
            for path in paths:
                with self.subTest(format=path.suffix):
                    output=root/path.suffix[1:]
                    output.mkdir()
                    before=fingerprint(path)
                    result=excel_call(dict(operation='prepare',source=str(path),selected={'Calculation':resolve_settings({'section_breaks':'25'})},staging=str(output)),lambda *args:None,timeout=45)
                    with fitz.open(result['rendered'][0]['path']) as pdf:
                        text=''.join(p.get_text() for p in pdf)
                        self.assertIn('Parameter 54',text)
                        self.assertIn('Foundation design calculation',text)
                        self.assertTrue(any(p.get_images() for p in pdf))
                    self.assertEqual(before,fingerprint(path))

    def test_real_excel_mixed_layout_and_original_preservation(self):
        from openpyxl import Workbook
        import pymupdf as fitz
        from civtools.core.report_service import excel_call
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            source=root/'Calculations.xlsx'
            wb=Workbook()
            sheet=wb.active
            sheet.title='Foundation'
            sheet.append(['Foundation calculation','Value','Unit'])
            for i in range(1,80):
                sheet.append([f'Design parameter {i}',i,'kN'])
            sheet.column_dimensions['A'].width=28
            sheet.column_dimensions['B'].width=15
            sheet.column_dimensions['C'].width=12
            wide=wb.create_sheet('Reactions')
            for i in range(40):
                wide.append([f'Case {i}']+[i+j for j in range(9)])
            for col in 'ABCDEFGHIJ':
                wide.column_dimensions[col].width=14
            wb.create_sheet('Hidden inputs').sheet_state='hidden'
            wb.save(source)
            before=fingerprint(source)
            inventory=excel_call(dict(operation='inspect',source=str(source)),lambda *args:None,timeout=45)
            self.assertEqual(len(inventory['sheets']),3)
            result=excel_call(dict(operation='prepare',source=str(source),selected={'Foundation':resolve_settings(), 'Reactions':resolve_settings()},staging=temp),lambda *args:None,timeout=45)
            self.assertEqual(len(result['rendered']),2)
            foundation,reactions=result['rendered']
            self.assertEqual(foundation['plan']['orientation'],'Portrait')
            self.assertEqual(reactions['plan']['orientation'],'Landscape',str(inventory['sheets'])+str(reactions['plan']))
            with fitz.open(foundation['path']) as pdf:
                text=''.join(p.get_text() for p in pdf)
                self.assertIn('Design parameter 79',text)
            self.assertEqual(fingerprint(source),before)


if __name__=='__main__':
    unittest.main()
