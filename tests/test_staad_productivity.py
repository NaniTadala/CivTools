import json
import math
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from civtools.core.staad_productivity import (ModelSnapshot, classify_members, id_list,
    group_commands, release_commands, renumber_plan, renumber_steps, combination_commands,
    tapered_commands, pedestal_forces, write_result_workbook)
from civtools.core.peb import PebConfig, build_peb_plan, geometry_commands
from civtools.core import staad_productivity_com as live


def model(path='example.std'):
    return ModelSnapshot(path, 'Metric',
        {1: (0, 0, 0), 2: (0, 6, 0), 3: (5, 6, 0), 4: (10, 6, 0),
         5: (10, 0, 0), 6: (5, 6, 5), 7: (5, 8, 0)},
        {10: (1, 2), 20: (2, 3), 30: (3, 4), 40: (4, 5), 50: (3, 6), 60: (2, 7)},
        [1, 5], [1, 2, 101])


class ProductivityPlansTests(unittest.TestCase):
    def test_snapshot_roundtrip_preserves_identity_and_rejects_missing_nodes(self):
        original = model()
        fingerprint = original.fingerprint
        original.source_signature = (1234, 987654321)
        restored = ModelSnapshot.from_json(original.to_json())
        self.assertEqual(restored.fingerprint, original.fingerprint)
        self.assertEqual(restored.source_signature, (1234, 987654321))
        self.assertEqual(original.fingerprint, fingerprint)
        restored.nodes[1] = (0, -1, 0)
        self.assertNotEqual(restored.fingerprint, original.fingerprint)
        data = json.loads(original.to_json())
        data['members']['10'] = [1, 999]
        with self.assertRaisesRegex(ValueError, 'invalid node'):
            ModelSnapshot.from_json(json.dumps(data))

    def test_nonfinite_geometry_and_invalid_ids_are_rejected(self):
        for bad in (float('nan'), float('inf')):
            with self.assertRaises(ValueError):
                ModelSnapshot('x', 'Metric', {1: (bad, 0, 0), 2: (1, 1, 1)}, {1: (1, 2)})
        for text in ('1 TO', '2 TO 1', '1 2 junk', '0', '1 TO 1000000'):
            with self.assertRaises(ValueError):
                id_list(text)
        self.assertEqual(id_list('1, 2 TO 4 2', range(1, 5)), [1, 2, 3, 4])

    def test_split_beam_between_columns_is_primary_and_ambiguous_slope_stays_unassigned(self):
        categories = classify_members(model())
        self.assertEqual(categories[10], 'Columns')
        self.assertEqual(categories[20], 'Primary beams X')
        self.assertEqual(categories[30], 'Primary beams X')
        self.assertEqual(categories[50], 'Secondary beams Z')
        self.assertEqual(categories[60], 'Sloping members')
        self.assertEqual(classify_members(model(), overrides={60: 'Rafters'})[60], 'Rafters')

    def test_primary_beams_require_support_connected_column_stacks(self):
        snapshot = ModelSnapshot('supported.std', 'Metric',
            {1: (0, 0, 0), 2: (0, 3, 0), 3: (0, 6, 0),
             4: (8, 0, 0), 5: (8, 3, 0), 6: (8, 6, 0), 7: (4, 6, 0),
             8: (0, 7, 4), 9: (0, 9, 4), 10: (8, 7, 4), 11: (8, 9, 4)},
            {1: (1, 2), 2: (2, 3), 3: (4, 5), 4: (5, 6),
             5: (3, 7), 6: (7, 6), 7: (8, 9), 8: (10, 11), 9: (9, 11)},
            [1, 4])
        categories = classify_members(snapshot)
        self.assertEqual(categories[5], 'Primary beams X')
        self.assertEqual(categories[6], 'Primary beams X')
        self.assertEqual(categories[9], 'Secondary beams X')
        snapshot.supports.clear()
        self.assertEqual(classify_members(snapshot)[5], 'Secondary beams X')

    def test_vertical_braces_in_xy_and_yz_planes_are_selected_separately_from_rafters(self):
        snapshot = ModelSnapshot('braced.std', 'Metric',
            {1: (0, 0, 0), 2: (0, 3, 0), 3: (0, 6, 0),
             4: (8, 0, 0), 5: (8, 3, 0), 6: (8, 6, 0),
             7: (0, 0, 8), 8: (0, 3, 8), 9: (0, 6, 8),
             10: (4, 8, 0), 11: (4, 1.5, 0)},
            {1: (1, 2), 2: (2, 3), 3: (4, 5), 4: (5, 6),
             5: (7, 8), 6: (8, 9), 7: (1, 11), 8: (11, 5),
             9: (1, 8), 10: (3, 10), 11: (10, 6), 12: (2, 6)}, [1, 4, 7])
        categories = classify_members(snapshot)
        self.assertEqual(categories[7], 'Vertical bracing')
        self.assertEqual(categories[8], 'Vertical bracing')
        self.assertEqual(categories[9], 'Vertical bracing')
        self.assertEqual(categories[12], 'Vertical bracing')
        self.assertEqual(categories[10], 'Rafters')
        self.assertEqual(categories[11], 'Rafters')

    def test_diagonal_in_identified_inclined_roof_plane_is_rafter_bracing(self):
        snapshot = ModelSnapshot('roof.std', 'Metric',
            {1: (0, 0, 0), 2: (0, 6, 0), 3: (8, 0, 0), 4: (8, 6, 0),
             5: (4, 8, 0), 6: (0, 0, 4), 7: (0, 6, 4),
             8: (8, 0, 4), 9: (8, 6, 4), 10: (4, 8, 4)},
            {1: (1, 2), 2: (3, 4), 3: (2, 5), 4: (5, 4),
             5: (6, 7), 6: (8, 9), 7: (7, 10), 8: (10, 9),
             9: (2, 7), 10: (5, 10), 11: (4, 9), 12: (2, 10)}, [1, 3, 6, 8])
        categories = classify_members(snapshot)
        self.assertEqual(categories[3], 'Rafters')
        self.assertEqual(categories[7], 'Rafters')
        self.assertEqual(categories[12], 'Rafter bracing')

    def test_groups_include_member_floor_and_support_sets(self):
        m = model()
        commands = group_commands(m, classify_members(m))
        self.assertIn('_CIV_PB_X 20 30', commands)
        self.assertIn('JOINT\n_CIV_SUPPORTS 1 5', commands)
        self.assertIn('FLOOR', commands)
        with self.assertRaises(ValueError):
            group_commands(m, classify_members(m), 'BAD\nLOAD')

    def test_release_axes_end_and_long_continuations(self):
        commands = release_commands(range(1, 300), 'BOTH', ['MY', 'MZ'])
        self.assertIn('LOCAL', commands)
        self.assertIn('START MY MZ', commands)
        self.assertIn('END MY MZ', commands)
        self.assertTrue(all(len(line) <= 78 for line in commands.splitlines()))
        with self.assertRaises(ValueError):
            release_commands([1], 'END', ['FX', 'FY', 'FZ', 'MX', 'MY', 'MZ'])

    def test_dense_renumbering_has_no_collisions_even_for_swaps(self):
        for mapping in ({1: 2, 2: 1}, {10: 1, 20: 2, 30: 3}, {1: 1, 2: 2}):
            active = {old: old for old in mapping}
            for old, new in renumber_steps(mapping):
                self.assertNotIn(new, active)
                active[new] = active.pop(old)
            self.assertEqual({old: new for new, old in active.items()}, mapping)
        mapping = renumber_plan(model(), 100, classify_members(model()))
        self.assertEqual(set(mapping.values()), set(range(100, 106)))

    def test_combinations_validate_duplicates_nested_cases_and_injection(self):
        rows = [[101, 'ULS', 1, 1.5], [101, 'ULS', 2, -1.5]]
        self.assertIn('1 1.5 2 -1.5', combination_commands(rows, [1, 2]))
        for invalid in (rows + [rows[0]], [[101, 'X\nFINISH', 1, 1]],
                        [[101, 'ULS', 101, 1]], [[101, 'ULS', 1, float('nan')]]):
            with self.assertRaises(ValueError):
                combination_commands(invalid)

    def test_tapered_dimensions_and_unit_command(self):
        self.assertIn('1 2 TAPERED 0.8 0.01 0.6', tapered_commands([1, 2], [.8, .01, .6, .3, .02, .3, .02]))
        for dimensions in ([.4, .01, .6, .3, .02, .3, .02], [.03, .01, .03, .3, .02, .3, .02]):
            with self.assertRaises(ValueError):
                tapered_commands([1], dimensions)

    def test_pedestal_global_y_transfer_retains_torsion(self):
        actual = pedestal_forces((10, 100, 20, 30, 40, 50), 2, .5, .5, 24)
        self.assertEqual(actual, (10, 112, 20, 70, 40, 30))

    def test_workbook_units_component_cases_and_existing_output_on_failure(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'results.xlsx'
            rows = [(1, 1, (2, 10, 0, 0, 4, 0)), (1, 2, (-5, 3, 0, 0, -2, 0))]
            write_result_workbook(path, model(), 'Reactions', rows)
            wb = load_workbook(path, read_only=True)
            try:
                self.assertEqual(wb['Results']['C1'].value, 'FX (kN)')
                self.assertEqual(list(wb['Envelopes'].values)[1], (1, 'FX', 'kN', -5, 2, 2, 1))
            finally:
                wb.close()
            saved = path.read_bytes()
            with self.assertRaises(ValueError):
                write_result_workbook(path, model(), 'Reactions', [(1, 1, (float('nan'),) * 6)])
            self.assertEqual(saved, path.read_bytes())
            english = model()
            english.base_unit = 'English'
            write_result_workbook(path, english, 'Displacements', rows)
            wb = load_workbook(path, read_only=True)
            try:
                self.assertEqual(wb['Results']['C1'].value, 'DX (in)')
                self.assertEqual(wb['Results']['F1'].value, 'RX (rad)')
            finally:
                wb.close()


class PebPlanTests(unittest.TestCase):
    def test_all_variants_are_connected_unique_and_purlin_spacing_is_bounded(self):
        for variant, columns in [('A', 2), ('B', 3), ('C', 4), ('D', 2)]:
            config = PebConfig(variant=variant, inner_left=4.3, inner_right=3.7)
            plan = build_peb_plan(config)
            with self.subTest(variant=variant):
                self.assertEqual(len(plan.foundations), config.frames * columns)
                self.assertEqual(len(plan.nodes), len(set(tuple(round(v, 10) for v in xyz) for xyz in plan.nodes)))
                self.assertEqual(len(plan.beams), len({tuple(sorted((a, b))) for a, b, _ in plan.beams}))
                self.assertTrue(all(a != b for a, b, _ in plan.beams))
                adjacent = {i: set() for i in range(len(plan.nodes))}
                for a, b, kind in plan.beams:
                    adjacent[a].add(b)
                    adjacent[b].add(a)
                    if kind == 'rafter':
                        self.assertLessEqual(math.dist(plan.nodes[a], plan.nodes[b]), config.purlin_spacing + 1e-8)
                visited, pending = set(), [0]
                while pending:
                    node = pending.pop()
                    if node not in visited:
                        visited.add(node)
                        pending.extend(adjacent[node] - visited)
                self.assertEqual(len(visited), len(plan.nodes))
                self.assertIn('UNIT METER KN', geometry_commands(plan))
                self.assertTrue(geometry_commands(plan).endswith('FINISH\n'))

    def test_extreme_spacing_bad_geometry_and_invalid_variants_fail_before_allocation(self):
        for kwargs in ({'purlin_spacing': 1e-20}, {'tie_spacing': 1e-20},
                       {'ridge_height': 2}, {'variant': 'AB'}, {'frames': 1},
                       {'variant': 'C', 'inner_left': 11}, {'width': float('inf')}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                build_peb_plan(PebConfig(**kwargs))


class StaadBoundaryTests(unittest.TestCase):
    def test_failed_save_prevents_backup_and_model_changes(self):
        staad = MagicMock()
        staad.SaveModel.return_value = False
        with patch.object(live.shutil, 'copy2') as copy:
            with self.assertRaisesRegex(RuntimeError, 'could not save'):
                live._backup_model(staad, model(), lambda *args: None)
            copy.assert_not_called()

    @staticmethod
    def selection_staad():
        staad = MagicMock()
        expected = model()
        staad.GetSTAADFile.return_value = expected.path
        staad.Geometry.GetNodeCount.return_value = len(expected.nodes)
        staad.Geometry.GetMemberCount.return_value = len(expected.members)
        members, nodes = set(), set()
        geometry = staad.Geometry
        geometry._geometry.ClearMemberSelection.side_effect = members.clear
        geometry._geometry.ClearNodeSelection.side_effect = nodes.clear
        geometry.SelectMultipleBeams.side_effect = members.update
        geometry.SelectMultipleNodes.side_effect = lambda ids: nodes.update(ids) or True
        geometry.SelectBeam.side_effect = lambda mid: members.add(mid) or True
        geometry.SelectNode.side_effect = lambda node: nodes.add(node) or True
        geometry.GetNoOfSelectedBeams.side_effect = lambda: len(members)
        geometry.GetNoOfSelectedNodes.side_effect = lambda: len(nodes)
        geometry.GetSelectedBeams.side_effect = lambda: sorted(members)
        geometry.GetSelectedNodes.side_effect = lambda: sorted(nodes)
        def inverse(entity):
            if entity == 1:
                nodes.symmetric_difference_update(expected.nodes)
            else:
                members.symmetric_difference_update(expected.members)
        staad.View._view.SelectInverse.side_effect = inverse
        return staad

    def test_selection_accepts_void_clear_return_without_broken_wrapper(self):
        # openstaadpy 0.2's public clear methods do `retVal < 0` on None.
        for nodes, ids in ((False, [30, 10, 30]), (True, [5, 1, 5])):
            with self.subTest(nodes=nodes):
                expected, staad = model(), self.selection_staad()
                geometry = staad.Geometry
                geometry.ClearMemberSelection.side_effect = TypeError("'<' not supported between instances of 'NoneType' and 'int'")
                geometry.ClearNodeSelection.side_effect = geometry.ClearMemberSelection.side_effect
                geometry._geometry.ClearMemberSelection.return_value = None
                geometry._geometry.ClearNodeSelection.return_value = None
                geometry.SelectMultipleBeams.return_value = None
                geometry.SelectMultipleNodes.return_value = True
                with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
                    result = live.select_members(expected, ids, lambda *args: None, nodes=nodes)
                clear = geometry._geometry.ClearNodeSelection if nodes else geometry._geometry.ClearMemberSelection
                select = geometry.SelectMultipleNodes if nodes else geometry.SelectMultipleBeams
                clear.assert_called_once_with()
                select.assert_called_once_with(sorted(set(ids)))
                geometry.ClearMemberSelection.assert_not_called()
                geometry.ClearNodeSelection.assert_not_called()
                self.assertIn('Selected 2', str(result))
                staad.View.RefreshView.assert_called_once_with()

    def test_silent_bulk_failure_retries_scalar_selection_and_verifies(self):
        for nodes, ids in ((False, [10, 30]), (True, [1, 5])):
            with self.subTest(nodes=nodes):
                expected, staad = model(), self.selection_staad()
                bulk = staad.Geometry.SelectMultipleNodes if nodes else staad.Geometry.SelectMultipleBeams
                bulk.side_effect = None
                bulk.return_value = None
                with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
                    live.select_members(expected, ids, lambda *args: None, nodes=nodes)
                getter = staad.Geometry.GetSelectedNodes if nodes else staad.Geometry.GetSelectedBeams
                self.assertEqual(getter(), ids)
                single = staad.Geometry.SelectNode if nodes else staad.Geometry.SelectBeam
                self.assertEqual([call.args[0] for call in single.call_args_list], ids)

    def test_silent_selection_failure_reports_mismatched_readback(self):
        expected, staad = model(), self.selection_staad()
        staad.Geometry.SelectMultipleBeams.side_effect = None
        staad.Geometry.SelectBeam.side_effect = None
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
            with self.assertRaisesRegex(RuntimeError, 'did not match'):
                live.select_members(expected, [10], lambda *args: None)

    def test_combined_selection_and_deselecting_last_group(self):
        expected, staad = model(), self.selection_staad()
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture') as capture:
            for members, nodes in (([10, 30], [1, 5]), ([30], [1, 5]), ([], [1, 5]), ([], [])):
                live.set_selection(expected, members, nodes, lambda *args: None)
                self.assertEqual(staad.Geometry.GetSelectedBeams(), members)
                self.assertEqual(staad.Geometry.GetSelectedNodes(), nodes)
            capture.assert_not_called()
        staad.Geometry.GetNodeCoordinates.assert_not_called()
        staad.Geometry.GetMemberIncidence.assert_not_called()

    def test_selection_cache_rejects_changed_target_counts_or_saved_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'model.std'
            path.write_text('STAAD SPACE')
            expected = model(str(path))
            stat = path.stat()
            expected.source_signature = (stat.st_size, stat.st_mtime_ns)
            for change in ('path', 'nodes', 'members', 'file'):
                with self.subTest(change=change):
                    staad = self.selection_staad()
                    staad.GetSTAADFile.return_value = str(path)
                    if change == 'path':
                        staad.GetSTAADFile.return_value = 'other.std'
                    elif change == 'nodes':
                        staad.Geometry.GetNodeCount.return_value += 1
                    elif change == 'members':
                        staad.Geometry.GetMemberCount.return_value += 1
                    else:
                        path.write_text('STAAD SPACE\n* saved model changed')
                    with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture') as capture:
                        with self.assertRaisesRegex(RuntimeError, 'Process members again'):
                            live.set_selection(expected, [10], [], lambda *args: None)
                        capture.assert_not_called()
                    staad.Geometry.SelectMultipleBeams.assert_not_called()

    def test_support_deselect_clears_nodes_without_bulk_inverse(self):
        expected, staad = model(), self.selection_staad()
        with patch.object(live, 'connect', return_value=staad):
            live.set_selection(expected, [10], [1, 5], lambda *args: None)
            staad.Geometry.SelectMultipleNodes.reset_mock()
            staad.View._view.SelectInverse.reset_mock()
            staad.Geometry._geometry.ClearNodeSelection.reset_mock()
            live.set_selection(expected, [10], [], lambda *args: None)
        self.assertEqual(staad.Geometry.GetSelectedNodes(), [])
        self.assertEqual(staad.Geometry.GetSelectedBeams(), [10])
        staad.Geometry.SelectMultipleNodes.assert_not_called()
        staad.View._view.SelectInverse.assert_not_called()
        staad.Geometry._geometry.ClearNodeSelection.assert_called_once_with()

    def test_fast_deselect_uses_bulk_inverse_with_clear_fallback(self):
        for inverse_works in (True, False):
            with self.subTest(inverse_works=inverse_works):
                expected, staad = model(), self.selection_staad()
                staad.Geometry.SelectMultipleBeams([10, 30])
                if not inverse_works:
                    staad.View._view.SelectInverse.side_effect = None
                with patch.object(live, 'connect', return_value=staad):
                    live.set_selection(expected, [], [], lambda *args: None)
                self.assertEqual(staad.Geometry.GetSelectedBeams(), [])
                staad.View._view.SelectInverse.assert_called_once_with(2)
                if inverse_works:
                    staad.Geometry._geometry.ClearMemberSelection.assert_not_called()
                else:
                    staad.Geometry._geometry.ClearMemberSelection.assert_called_once_with()

    def test_redraw_is_restored_on_failure_without_showing_hidden_children(self):
        from types import SimpleNamespace
        staad = MagicMock()
        staad.GetMainWindowHandle.return_value = 10
        gui = MagicMock()
        gui.IsWindow.return_value = True
        gui.IsWindowVisible.side_effect = lambda handle: handle != 12
        gui.EnumChildWindows.side_effect = lambda handle, callback, result: [
            callback(child, result) for child in (11, 12)]
        gui.error = RuntimeError
        constants = SimpleNamespace(WM_SETREDRAW=11, RDW_INVALIDATE=1, RDW_ALLCHILDREN=128)
        with patch.dict(sys.modules, {'win32gui': gui, 'win32con': constants}):
            with self.assertRaisesRegex(RuntimeError, 'selection failed'):
                with live.selection_redraw_batch(staad):
                    raise RuntimeError('selection failed')
        self.assertEqual([call.args for call in gui.SendMessage.call_args_list],
                         [(10, 11, 0, 0), (11, 11, 0, 0), (11, 11, 1, 0), (10, 11, 1, 0)])
        gui.RedrawWindow.assert_called_once_with(10, None, None, 129)

    def test_failed_combined_selection_restores_previous_staad_selection(self):
        expected, staad = model(), self.selection_staad()
        staad.Geometry.SelectMultipleBeams([10])
        staad.Geometry.SelectMultipleNodes([1])
        select_nodes = staad.Geometry.SelectMultipleNodes.side_effect
        def reject_new_nodes(ids):
            if ids == [5]:
                raise RuntimeError('STAAD rejected node selection')
            return select_nodes(ids)
        staad.Geometry.SelectMultipleNodes.side_effect = reject_new_nodes
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
            with self.assertRaisesRegex(RuntimeError, 'rejected node selection'):
                live.set_selection(expected, [20, 30], [5], lambda *args: None)
        self.assertEqual(staad.Geometry.GetSelectedBeams(), [10])
        self.assertEqual(staad.Geometry.GetSelectedNodes(), [1])

    def test_invalid_selection_count_cannot_silently_report_success(self):
        expected, staad = model(), self.selection_staad()
        staad.Geometry.GetNoOfSelectedBeams.side_effect = None
        staad.Geometry.GetNoOfSelectedBeams.return_value = None
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
            with self.assertRaisesRegex(RuntimeError, 'valid selection count'):
                live.select_members(expected, [10], lambda *args: None)

    def test_clear_failure_prevents_selection(self):
        expected, staad = model(), MagicMock()
        staad.Geometry._geometry.ClearMemberSelection.return_value = -1
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
            with self.assertRaisesRegex(RuntimeError, 'could not clear'):
                live.select_members(expected, [10], lambda *args: None)
        staad.Geometry.SelectMultipleBeams.assert_not_called()

    def test_node_selection_failure_is_reported(self):
        expected, staad = model(), self.selection_staad()
        staad.Geometry.SelectMultipleNodes.side_effect = None
        staad.Geometry.SelectMultipleNodes.return_value = False
        staad.Geometry.SelectNode.side_effect = None
        staad.Geometry.SelectNode.return_value = False
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
            with self.assertRaisesRegex(RuntimeError, 'could not select'):
                live.select_members(expected, [1], lambda *args: None, nodes=True)

    def test_changed_model_prevents_selection(self):
        expected = model()
        current = model()
        current.nodes[1] = (1, 0, 0)
        staad = MagicMock()
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=current):
            with self.assertRaisesRegex(RuntimeError, 'changed'):
                live.select_members(expected, [10], lambda *args: None)
        staad.Geometry.SelectMultipleBeams.assert_not_called()

    def test_no_results_never_writes_workbook(self):
        expected, staad = model(), MagicMock()
        staad.Output.AreResultsAvailable.return_value = False
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected), patch.object(live, 'write_result_workbook') as write:
            with self.assertRaisesRegex(RuntimeError, 'Run analysis'):
                live.export_results(expected, 'Reactions', [1], [1], 'x.xlsx', lambda *args: None)
        write.assert_not_called()

    def test_failed_result_read_does_not_publish_partial_workbook(self):
        expected, staad = model(), MagicMock()
        staad.Output.AreResultsAvailable.return_value = True
        staad.Output.GetSupportReactions.side_effect = [(1, 2, 3, 4, 5, 6), RuntimeError('No result')]
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected), patch.object(live, 'write_result_workbook') as write:
            with self.assertRaisesRegex(RuntimeError, 'No result'):
                live.export_results(expected, 'Reactions', [1], [1, 2], 'x.xlsx', lambda *args: None)
        write.assert_not_called()

    def test_renumber_creates_backup_and_rejects_partial_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'model.std'
            path.write_text('original model')
            expected, staad = model(str(path)), MagicMock()
            for name in ('GetLastPlateNo', 'GetLastSolidNo', 'GetLastPhysicalMemberNo'):
                getattr(staad.Geometry, name).return_value = 0
            staad.GetSTAADFile.return_value = str(path)
            staad.Geometry.RenumberBeam.side_effect = [True, False]
            mapping = renumber_plan(expected)
            with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
                with self.assertRaisesRegex(RuntimeError, 'partial changes'):
                    live.apply_renumber(expected, mapping, lambda *args: None)
            backups = list(Path(temp).glob('model.civtools-*.std'))
            self.assertEqual(len(backups), 1)
            self.assertEqual(backups[0].read_text(), 'original model')
            self.assertTrue(backups[0].with_suffix('.renumber.json').is_file())
            self.assertEqual(staad.SaveModel.call_count, 1)

    def test_backup_failure_prevents_any_renumber(self):
        staad, expected = MagicMock(), model()
        for name in ('GetLastPlateNo', 'GetLastSolidNo', 'GetLastPhysicalMemberNo'):
            getattr(staad.Geometry, name).return_value = 0
        with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected), patch.object(live.shutil, 'copy2', side_effect=PermissionError('No write')):
            with self.assertRaises(PermissionError):
                live.apply_renumber(expected, renumber_plan(expected), lambda *args: None)
        staad.Geometry.RenumberBeam.assert_not_called()

    def test_successful_renumber_verifies_member_connectivity_and_saves(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'model.std'
            path.write_text('test model')
            staad, expected = MagicMock(), model(str(path))
            mapping = renumber_plan(expected)
            for name in ('GetLastPlateNo', 'GetLastSolidNo', 'GetLastPhysicalMemberNo'):
                getattr(staad.Geometry, name).return_value = 0
            staad.GetSTAADFile.return_value = str(path)
            staad.Geometry.RenumberBeam.return_value = True
            staad.Geometry.GetBeamList.return_value = list(mapping.values())
            reverse = {new: old for old, new in mapping.items()}
            staad.Geometry.GetMemberIncidence.side_effect = lambda new: expected.members[reverse[new]]
            with patch.object(live, 'connect', return_value=staad), patch.object(live, 'capture', return_value=expected):
                result = live.apply_renumber(expected, mapping, lambda *args: None)
            self.assertIn('Re-run analysis', str(result))
            self.assertTrue(Path(result.output_path).is_file())
            self.assertEqual(staad.SaveModel.call_count, 2)


if __name__ == '__main__':
    unittest.main()
